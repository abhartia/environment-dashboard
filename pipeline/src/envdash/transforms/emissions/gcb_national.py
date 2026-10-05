"""Global Carbon Budget 2025, national fossil carbon dioxide: territorial and consumption-based emissions by country,
from 1990.

Input: National_Fossil_Carbon_Emissions_2025_v1.0.xlsx (ICOS PID 11676/loCXyssaalv6DPdO6Qdj90qQ). Sheets "Territorial
Emissions" (1850-2024) and "Consumption Emissions" (1990-2024; country values end in 2023) each have a block of notes,
an "MtC/yr" line, a row of upper-case column names, a row of display names, then one row per year with one column per
country, then regions, international shipping and aviation, a statistical difference and the world. The two sheets must
have the same columns in the same order. "Regions" lists the countries of each region. The territorial sheet's stated
dimension runs to row 1,048,576, so the reader resets it and reads the rows the file actually holds.

Accounting. Territorial emissions are those released inside a country's borders. Consumption-based emissions are
allocated to the country where the goods and services are finally consumed ("computed as in Peters et al. (2011),
using 'Territorial Emissions' as reference"). Fossil emissions here are gross of the cement carbonation sink and
national values exclude bunker fuels ("National estimates include emissions from fossil fuel combustion and oxidation
and cement production and excludes emissions from bunker fuels. World totals include emissions from bunker fuels.").

The "Emissions Transfers" sheet is not published. The Global Carbon Budget 2025 paper describes it as "Consumption minus
territorial emissions", but in this file a column's transfer equals territorial minus consumption-based emissions (in
4,202 of the 4,205 cells of countries, EU27, international transport and the world where all three sheets have a value,
to within 0.000001 MtC), the opposite sign; three cells and the producer's regional columns match neither. Until the
producer says which sign is meant, the two series it is computed from are published and the difference is left to the
reader.

Years. The indicator covers the years of the consumption sheet (1990-2024): territorial values before 1990 are not
compared here (the same national series from 1750 is emissions.gcp-2025.fossil-co2-by-country). Empty cells are not
published and never read as zero; most countries have no consumption-based value at all, and country values stop in
2023 while the world and international transport go to 2024.

Entities. Display names (the second header row) map to entity codes by the entity table's name column, except for the
names in ALIASES, declared from this file. Columns in NOT_ENTITIES (the producer's own regions and the statistical
difference) are not published. EU27 is published only after checking that the "Regions" sheet's EU27 list resolves to
exactly the 27 member states of envdash.geo.EU27_MEMBERS. A column header saying "(INCLUDING ...)" (France with Monaco,
Italy with San Marino) is copied to the observations as a note.

Units. "All values in million tonnes of carbon per year. For values in million tonnes of CO2 per year, multiply the
values below by 3.664": values are multiplied by 3.664 with exact decimal arithmetic on the stored numbers.

Version. The file carries no version label; PINNED maps the pinned ICOS object to its version and submission date.
"""

from __future__ import annotations

import io
from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from pathlib import Path

import openpyxl

from envdash import geo
from envdash.models import Dimension, DimensionValue, Display, Observation, Scope, Unit
from envdash.paths import Paths
from envdash.transform import Input, InputFile, Result, Spec, Transform, Validation

SOURCE = "gcb-2025-national"
NATIONAL = Input(SOURCE, "national-fossil")
MTC_TO_MTCO2 = Decimal("3.664")

PINNED: dict[str, tuple[str, date]] = {
    # meta.icos-cp.eu/objects/loCXyssaalv6DPdO6Qdj90qQ (read 2026-10-04):
    # National_Fossil_Carbon_Emissions_2025_v1.0.xlsx, submitted 2026-05-17, latestVersion = itself.
    "https://data.icos-cp.eu/objects/loCXyssaalv6DPdO6Qdj90qQ": ("2025 v1.0", date(2026, 5, 17)),
}

TERRITORIAL = "Territorial Emissions"
CONSUMPTION = "Consumption Emissions"
REGIONS = "Regions"

ACCOUNTING: tuple[tuple[str, str, str], ...] = (
    # (dimension value id, label, sheet)
    ("territorial", "Territorial (where the emissions are released)", TERRITORIAL),
    ("consumption", "Consumption-based (where the goods and services are consumed)", CONSUMPTION),
)

UNIT_NOTE = (
    "All values in million tonnes of carbon per year. For values in million tonnes of CO2 per year, multiply the "
    "values below by 3.664"
)
REQUIRED: dict[str, tuple[str, ...]] = {
    TERRITORIAL: (
        UNIT_NOTE,
        "(1) National estimates include emissions from fossil fuel combustion and oxidation and cement production and "
        "excludes emissions from bunker fuels. World totals include emissions from bunker fuels.",
        "(4) The statistical difference presented on column HX is the difference between the world emissions and the "
        "sum of the emissions for each countries and for the bunker fuels.",
    ),
    CONSUMPTION: (
        UNIT_NOTE,
        "Consumption emissions are computed as in Peters et al. (2011), using 'Territorial Emissions' as reference",
    ),
}

# Display names (second header row) that are not the entity table's name for the same place. Each was read from the
# file (sha256 968097cacb1a…); the target is our entity code.
ALIASES: dict[str, str] = {
    "Bahamas": "BHS",
    "Bonaire, Saint Eustatius and Saba": "BES",
    "Brunei Darussalam": "BRN",
    "Cape Verde": "CPV",
    "Congo": "COG",
    "Côte d'Ivoire": "CIV",
    "Faeroe Islands": "FRO",
    "Micronesia (Federated States of)": "FSM",
    "Hong Kong": "HKG",
    "Macao": "MAC",
    "State of Palestine": "PSE",
    "Sint Maarten (Dutch part)": "SXM",
    "Sao Tome and Principe": "STP",
    "Serbia": "SRB",
    "Eswatini": "SWZ",
    "Timor-Leste": "TLS",
    "Türkiye": "TUR",
    "Tanzania": "TZA",
    "USA": "USA",
    "Viet Nam": "VNM",
    "Wallis and Futuna Islands": "WLF",
    "EU27": "EU27",
    "International Shipping": "INTL_SEA",
    "International Aviation": "INTL_AIR",
    "World": "WLD",
}
# The producer's own groupings and its statistical difference: not entities here.
NOT_ENTITIES = frozenset(
    {
        "KP Annex B", "Non KP Annex B", "OECD", "Non-OECD", "Africa", "Asia", "Central America", "Europe",
        "Middle East", "North America", "Oceania", "South America", "Statistical Difference",
    }
)  # fmt: skip


class GcbNationalFormatError(ValueError):
    pass


@dataclass(frozen=True)
class Column:
    upper: str
    name: str
    entity: str


@dataclass(frozen=True)
class Sheet:
    columns: tuple[Column | None, ...]
    """One per data column (None for NOT_ENTITIES)."""
    years: tuple[int, ...]
    values: dict[tuple[str, int], Decimal]
    """(entity, year) -> MtC/yr as stored; empty cells absent."""


def _resolve(name: str, entities: geo.EntityTable | None) -> str:
    if name in ALIASES:
        return geo.resolve(ALIASES[name], "iso3", entities=entities)
    return geo.resolve(name, "name", entities=entities)


def _decimal(v: object, where: str) -> Decimal:
    if isinstance(v, bool) or not isinstance(v, int | float):
        raise GcbNationalFormatError(f"{where}: expected a number, found {v!r}")
    return Decimal(repr(v))


def _rows(wb, sheet: str) -> list[tuple]:
    if sheet not in wb.sheetnames:
        raise GcbNationalFormatError(f"no sheet {sheet!r} (sheets: {wb.sheetnames})")
    ws = wb[sheet]
    ws.reset_dimensions()  # the stated dimension (to row 1,048,576 and column AMJ) is not what the sheet holds
    rows = [tuple(r) for r in ws.iter_rows(values_only=True)]
    filled = [i for i, r in enumerate(rows) if any(c is not None for c in r)]
    rows = rows[: filled[-1] + 1] if filled else []
    width = max((len(r) for r in rows), default=0)
    return [r + (None,) * (width - len(r)) for r in rows]


def read_sheet(rows: list[tuple], sheet: str, entities: geo.EntityTable | None) -> Sheet:
    unit_at = next((i for i, r in enumerate(rows) if r and r[0] == "MtC/yr"), None)
    if unit_at is None:
        raise GcbNationalFormatError(f"{sheet}: no 'MtC/yr' line")
    notes = " ".join(str(c) for r in rows[:unit_at] for c in r if c is not None)
    for s in REQUIRED[sheet]:
        if s not in notes:
            raise GcbNationalFormatError(f"{sheet} no longer says {s!r}; re-read the file before trusting these rules")
    upper, names = rows[unit_at + 1], rows[unit_at + 2]
    if upper[0] is not None or names[0] is not None:
        raise GcbNationalFormatError(f"{sheet}: the header rows do not start with an empty cell")
    width = max(i for i, c in enumerate(names) if c is not None) + 1
    if any(c is None for c in names[1:width]) or any(c is None for c in upper[1:width]):
        raise GcbNationalFormatError(f"{sheet}: a column has no name")
    columns: list[Column | None] = []
    for u, n in zip(upper[1:width], names[1:width], strict=True):
        columns.append(None if n in NOT_ENTITIES else Column(str(u), str(n), _resolve(str(n), entities)))
    codes = [c.entity for c in columns if c is not None]
    if len(set(codes)) != len(codes):
        raise GcbNationalFormatError(f"{sheet}: two columns resolve to the same entity")
    years: list[int] = []
    values: dict[tuple[str, int], Decimal] = {}
    for n, r in enumerate(rows[unit_at + 3 :], start=unit_at + 4):
        if all(c is None for c in r):
            continue
        if any(c is not None for c in r[width:]):
            raise GcbNationalFormatError(f"{sheet} row {n} has values beyond column {width}")
        year = r[0]
        if isinstance(year, bool) or not isinstance(year, int):
            raise GcbNationalFormatError(f"{sheet} row {n}: year {year!r} is not an integer")
        if years and year != years[-1] + 1:
            raise GcbNationalFormatError(f"{sheet} row {n}: year {year} does not follow {years[-1]}")
        years.append(year)
        for col, v in zip(columns, r[1:width], strict=True):
            if col is not None and v is not None:
                values[(col.entity, year)] = _decimal(v, f"{sheet} row {n} {col.name}")
    if not years:
        raise GcbNationalFormatError(f"{sheet}: no data rows")
    return Sheet(tuple(columns), tuple(years), values)


def read_eu27(rows: list[tuple], entities: geo.EntityTable | None) -> frozenset[str]:
    found = [r for r in rows if r and r[0] == "EU27"]
    if len(found) != 1 or not isinstance(found[0][1], str):
        raise GcbNationalFormatError("the Regions sheet has no single EU27 row")
    return frozenset(_resolve(n.strip(), entities) for n in found[0][1].split(","))


def read_national(raw: bytes, *, entities: geo.EntityTable | None = None) -> dict[str, Sheet]:
    wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
    try:
        sheets = {s: read_sheet(_rows(wb, s), s, entities) for _, _, s in ACCOUNTING}
        eu27 = read_eu27(_rows(wb, REGIONS), entities)
    finally:
        wb.close()
    cols = {s: [(c.upper, c.name) if c else None for c in sh.columns] for s, sh in sheets.items()}
    if cols[TERRITORIAL] != cols[CONSUMPTION]:
        raise GcbNationalFormatError("the territorial and consumption sheets do not have the same columns")
    if eu27 != geo.EU27_MEMBERS:
        raise GcbNationalFormatError(
            f"the Regions sheet's EU27 differs from the 27 member states: {sorted(eu27 ^ geo.EU27_MEMBERS)}"
        )
    if not set(sheets[CONSUMPTION].years) <= set(sheets[TERRITORIAL].years):
        raise GcbNationalFormatError("the consumption sheet has years the territorial sheet lacks")
    return sheets


def _pinned(f: InputFile) -> tuple[str, date]:
    url = str(f.snapshot.url) if f.snapshot.url else None
    if url not in PINNED:
        raise GcbNationalFormatError(f"{url!r} is not a pinned Global Carbon Budget object; add its version to PINNED")
    return PINNED[url]


def observations(sheets: dict[str, Sheet]) -> list[Observation]:
    years = sheets[CONSUMPTION].years
    columns = [c for c in sheets[TERRITORIAL].columns if c is not None]
    obs: list[Observation] = []
    for acc_id, _, sheet in ACCOUNTING:
        values = sheets[sheet].values
        for col in columns:
            note = f"The file's column header reads '{col.upper}'." if "(INCLUDING" in col.upper else None
            for y in years:
                v = values.get((col.entity, y))
                if v is None:
                    continue
                obs.append(
                    Observation(
                        entity=col.entity,
                        period=f"{y:04d}",
                        value=float(v * MTC_TO_MTCO2),
                        note=note,
                        dims={"accounting": acc_id},
                    )
                )
    return obs


def _run(files: dict[str, InputFile]) -> Result:
    f = files[NATIONAL.key]
    version, published = _pinned(f)
    sheets = read_national(f.path.read_bytes())
    years = sheets[CONSUMPTION].years
    return Result(
        observations=observations(sheets),
        vintage=version,
        date_published=published.isoformat(),
        steps=[
            f"Read the sheets '{TERRITORIAL}' and '{CONSUMPTION}' of "
            f"National_Fossil_Carbon_Emissions_2025_v1.0.xlsx (Global Carbon Budget {version}, ICOS object "
            "loCXyssaalv6DPdO6Qdj90qQ), in million tonnes of carbon per year, and kept the years of the consumption "
            f"sheet, {years[0]}–{years[-1]}.",
            "Mapped column names to entity codes by the entity table's names and an explicit alias table; left out the "
            "producer's regions (Kyoto Protocol Annex B and non-Annex B, OECD and non-OECD, and eight continental "
            "regions) and the statistical difference column. EU27 is kept after checking that the Regions sheet lists "
            "exactly the 27 member states.",
            "Converted from million tonnes of carbon to million tonnes of carbon dioxide by multiplying by 3.664, the "
            'factor the sheets state ("multiply the values below by 3.664"), with exact decimal arithmetic on the '
            "stored values.",
            "Cells left empty in the file are not published and are never read as zero.",
        ],
        changes="converted from million tonnes of carbon to million tonnes of carbon dioxide (× 3.664); years before "
        "1990 left out.",
    )


def transforms(paths: Paths) -> list[Transform]:
    return [
        Transform(
            spec=Spec(
                id="emissions.gcb-2025.territorial-vs-consumption",
                title="Territorial and consumption-based fossil carbon dioxide emissions by country",
                description="Each country's fossil carbon dioxide emissions counted two ways since 1990: where they "
                "are released (territorial), and where the goods and services they went into are finally consumed "
                "(consumption-based). Where consumption-based emissions are higher, the country's imports carried "
                "more emissions than its exports.",
                kind="series",
                unit=Unit(code="MtCO2/yr", label="million tonnes of carbon dioxide per year", short="Mt CO₂/yr"),
                display=Display(decimals=1),
                scope=Scope(
                    geography="Countries and territories, the European Union (27), international aviation, "
                    "international shipping and the world",
                    lulucf="excluded",
                    bunkers="excluded",
                    basis="Carbon dioxide only, from fossil fuels and cement production, before the cement "
                    "carbonation sink is subtracted "
                    "(gross). National values exclude international aviation and shipping bunker fuels, which are "
                    "separate entities; the world value includes them. Consumption-based accounting follows Peters et "
                    "al. (2011).",
                ),
                geo_coverage="mixed",
                headline_entity="WLD",
                dimensions=(
                    Dimension(
                        id="accounting",
                        label="Accounting",
                        values=[DimensionValue(id=i, label=label) for i, label, _ in ACCOUNTING],
                    ),
                ),
                headline_dims=(("accounting", "territorial"),),
            ),
            inputs=(NATIONAL,),
            run=_run,
            module_file=Path(__file__),
            validation=Validation(min_rows=8_000, value_range=(-10.0, 50_000.0)),
        )
    ]
