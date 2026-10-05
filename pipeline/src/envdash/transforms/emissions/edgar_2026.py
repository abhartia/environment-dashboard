"""EDGAR 2026 (EDGAR_2026_GHG, JRC): greenhouse gas emissions of every country, as EDGAR publishes them.

Four indicators, all in this one module so that its sha256 (recorded in every processing step and build key) covers
every line of code that decides them:

- ghg.edgar-2026.total-by-country: total greenhouse gas emissions excluding land use, land-use change and forestry
  (LULUCF), in million tonnes of CO2-equivalent (IPCC AR5 GWP-100) per year, 1970-2025, for every country EDGAR
  reports, international aviation, international shipping, the EU27 and the world;
- ghg-per-capita.edgar-2026.by-country: the same per person, in tonnes of CO2-equivalent per person per year;
- ghg-gas-share.edgar-2026.global: the report's own shares of fossil CO2, CH4, N2O and F-gases in the 2025 world
  total, in percent;
- ghg-sector-change.edgar-2026.global: the report's own change in world emissions by sector, 2025 against 1990, 2005
  and 2024, in percent.

Licence. Every total includes IEA-EDGAR CO2, which is CC BY-NC-ND 4.0 (class no-derivatives, see
pipeline/sources/edgar-2026-ghg.yaml). So nothing is computed here: every published number is one cell of the booklet
workbook, unchanged (the double stored in the cell, no rounding), or a number printed in the report PDF. No share,
sum, per-person value, ratio or unit conversion of ours. The exports go to data-private only (the build routes them by
class).

Inputs. EDGAR_2026_GHG_booklet_2026.xlsx (artifact ghg-booklet), sheets GHG_totals_by_country and
GHG_per_capita_by_country; and the report "GHG emissions of all world countries - 2026 Report" (artifact report-pdf,
301 pages). The workbook's "info" sheet must still state the units and gases quoted in INFO_STATEMENTS, and its
"Citations and references" sheet the version; the PDF pages named below must still carry the sentences quoted. If
anything differs, the transform stops: a person must re-read the files. Page numbers are 1-based pages of the PDF file
(the printed page number is two lower).

Scope. The report states that the totals exclude LULUCF (PDF page 11: "global GHG emissions (excluding LULUCF)
reached a new record high of 54.1 Gt CO₂eq in 2025") and that the world total includes international shipping and
aviation (PDF page 74, Annex 5: "Global totals for all countries, including international shipping and aviation").
Country rows exclude them: international aviation and shipping are their own rows (EDGAR codes AIR and SEA).

Preliminary years. The report says (PDF page 50) "the emissions for the Fast-Track years (2024-2025) reported in this
booklet will be updated in subsequent editions of this booklet". Values for 2024 and 2025 are published with status
preliminary and that sentence as their note.

Entities. EDGAR codes are ISO 3166-1 alpha-3 except the ones declared in ALIASES (aggregates, international
transport, and two codes that are not current ISO codes); every other code goes through geo.resolve(code, "iso3"),
which raises for anything unknown. Six rows cover a country together with a neighbour, as their names in the file say
(COMBINED, e.g. ESP "Spain and Andorra"): they are published under the code EDGAR gives them, with the producer's row
name as a note on every value. SCG "Serbia and Montenegro" is one row for two countries (EDGAR has no separate Serbia
or Montenegro rows) and is published as SRB_MNE, Serbia and Montenegro as reported together. GUF "French Guiana" and
REU "Réunion" are published under their ISO codes (envdash/geo.py EXTRA_TERRITORIES). ANT is labelled "Curaçao" in
the workbook and on its report profile page (PDF page 128, population 168.000k in 2025), so it is published as CUW.

Publisher cross-checks. The report's profile pages print each total and per-person value for 1990, 2005, 2015 and 2025
to three decimals (World on PDF page 75, China 119, the United States 278); those rows are the checks below, for this
same vintage.

Gas shares. Quoted from the main findings (PDF page 11). The World profile (PDF page 75) prints the same four
shares around its pie chart; the transform reads them there too and stops unless they are identical.

Sector change. Read from the table of the World profile (PDF page 75): the sector labels and column headings must be
exactly as in SECTORS and SECTOR_COLUMNS, in that order. The values are whole percentages as printed ("0%" is
printed for changes that round to zero).
"""

from __future__ import annotations

import io
import logging
import re
from dataclasses import dataclass
from pathlib import Path

import openpyxl

from envdash import geo, textmatch
from envdash.models import Dimension, DimensionValue, Display, Observation, PublishedValueRef, Scope, Unit
from envdash.paths import Paths
from envdash.transform import Input, InputFile, PublisherCheck, Result, Spec, Transform, Validation

SOURCE = "edgar-2026-ghg"
BOOKLET = Input(SOURCE, "ghg-booklet")
REPORT = Input(SOURCE, "report-pdf")
VINTAGE = "EDGAR_2026_GHG"
BOOKLET_URL = "https://edgar.jrc.ec.europa.eu/booklet/EDGAR_2026_GHG_booklet_2026.xlsx"
REPORT_URL = "https://edgar.jrc.ec.europa.eu/booklet/GHG_emissions_of_all_world_countries_booklet_2026report.pdf"
REPORT_PAGES = 301
FIRST_YEAR, LAST_YEAR = 1970, 2025
FAST_TRACK_YEARS = frozenset({2024, 2025})

TOTALS_SHEET = "GHG_totals_by_country"
PER_CAPITA_SHEET = "GHG_per_capita_by_country"
HEADER = ("EDGAR Country Code", "Country")

INFO_STATEMENTS = (
    "GHG emissions include CO2 (fossil only), CH4, N2O and F-gases. They are aggregated using Global Warming Potential "
    "values from IPCC AR5 (GWP-100 AR5).",
    "values in GHG_totals_by_country sheet are expressed in Mt CO2eq/yr",
    "values in GHG_per_capita_by_countr sheet are expressed in t CO2eq/cap/yr",
)
CITATION_STATEMENT = "IEA-EDGAR CO2, EDGAR CH4, EDGAR N2O and EDGAR F-gases version EDGAR_2026_GHG (2026)."


@dataclass(frozen=True)
class Quote:
    page: int
    """1-based page of the report PDF."""
    text: str


LULUCF_QUOTE = Quote(
    11,
    "The latest EDGAR estimates indicate that global GHG emissions (excluding LULUCF) reached a new record high of "
    "54.1 Gt CO₂eq in 2025",
)
BUNKERS_QUOTE = Quote(74, "Global totals for all countries, including international shipping and aviation")
FAST_TRACK_QUOTE = Quote(
    50,
    "the emissions for the Fast-Track years (2024-2025) reported in this booklet will be updated in subsequent "
    "editions of this booklet",
)
GAS_QUOTE = Quote(
    11,
    "In 2025, fossil CO₂ remained the dominant component of global GHG emissions, accounting for 73.8% of the total. "
    "CH₄ contributed 17.6%, N₂O 5.3% and F-gases 3.3%.",
)
GAS_LOCATOR = "Executive summary, Main findings, printed page 9 (PDF page 11)"
WORLD_PROFILE_PAGE = 75
SECTOR_LOCATOR = "Annex 5, World profile, table of changes by sector, printed page 73 (PDF page 75)"

FAST_TRACK_NOTE = (
    "Fast-Track estimate: the report says the emissions for the Fast-Track years (2024-2025) will be updated in "
    "subsequent editions."
)

# EDGAR code -> (row name in the workbook, our entity code).
ALIASES: dict[str, tuple[str, str]] = {
    "GLOBAL TOTAL": ("GLOBAL TOTAL", "WLD"),
    "EU27": ("EU27", "EU27"),
    "AIR": ("International Aviation", "INTL_AIR"),
    "SEA": ("International Shipping", "INTL_SEA"),
    # Old ISO code of the Netherlands Antilles; EDGAR names the row and its profile page (PDF page 128) Curaçao.
    "ANT": ("Curaçao", "CUW"),
    # One row for two countries, under a code that is neither; EDGAR has no separate SRB or MNE rows.
    "SCG": ("Serbia and Montenegro", "SRB_MNE"),
}
# Rows covering a country with a neighbour, by their names in the workbook; published under the code EDGAR gives.
COMBINED: dict[str, str] = {
    "CHE": "Switzerland and Liechtenstein",
    "ESP": "Spain and Andorra",
    "FRA": "France and Monaco",
    "ISR": "Israel and Palestine, State of",
    "ITA": "Italy, San Marino and the Holy See",
    "SDN": "Sudan and South Sudan",
}

SECTORS: tuple[tuple[str, str], ...] = (
    ("power-industry", "Power Industry"),
    ("industrial-combustion-and-processes", "Industrial Combustion and Processes"),
    ("buildings", "Buildings"),
    ("transport", "Transport"),
    ("fuel-production", "Fuel Production"),
    ("agriculture", "Agriculture"),
    ("waste", "Waste"),
    ("all-sectors", "All sectors"),
)
SECTOR_COLUMNS: tuple[tuple[str, str], ...] = (
    ("1990", "2025 vs 1990"),
    ("2005", "2025 vs 2005"),
    ("2024", "2025 vs 2024"),
)
GASES: tuple[tuple[str, str, str], ...] = (
    # (dimension value id, label, name on the World profile's pie chart)
    ("fossil-co2", "Fossil carbon dioxide (CO₂)", "CO2"),
    ("ch4", "Methane (CH₄)", "CH4"),
    ("n2o", "Nitrous oxide (N₂O)", "N2O"),
    ("f-gases", "Fluorinated gases (F-gases)", "F-gases"),
)

MT_CO2E_YR = Unit(code="MtCO2e/yr", label="million tonnes of carbon dioxide equivalent per year", short="Mt CO₂e/yr")
T_CO2E_PERSON_YR = Unit(
    code="tCO2e/person/yr",
    label="tonnes of carbon dioxide equivalent per person per year",
    short="t CO₂e/person/yr",
)
PERCENT_OF_TOTAL = Unit(code="percent", label="percent of world greenhouse gas emissions", short="%")
PERCENT_CHANGE = Unit(code="percent", label="percent change", short="%")

GAS_BASIS = (
    "CO₂ (fossil only: fuels and industrial processes), CH₄, N₂O and F-gases, added up with IPCC AR5 100-year global "
    "warming potentials."
)


class EdgarFormatError(ValueError):
    pass


# --- workbook ----------------------------------------------------------------------------------------------------


@dataclass(frozen=True)
class Row:
    code: str
    name: str
    values: tuple[float, ...]
    """One per year FIRST_YEAR..LAST_YEAR, as stored in the cells."""


def _cells(raw: bytes, sheets: tuple[str, ...]) -> dict[str, list[tuple]]:
    wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
    try:
        missing = [s for s in (*sheets, "info", "Citations and references") if s not in wb.sheetnames]
        if missing:
            raise EdgarFormatError(f"no sheet(s) {missing} (sheets: {wb.sheetnames})")
        return {
            s: [tuple(r) for r in wb[s].iter_rows(values_only=True)]
            for s in {*sheets, "info", "Citations and references"}
        }
    finally:
        wb.close()


def check_statements(sheets: dict[str, list[tuple]]) -> None:
    info = {str(c).strip() for r in sheets["info"] for c in r if c is not None}
    for s in INFO_STATEMENTS:
        if s not in info:
            raise EdgarFormatError(f"the info sheet no longer says {s!r}; re-read the file before trusting these rules")
    citations = " ".join(str(c) for r in sheets["Citations and references"] for c in r if c is not None)
    if CITATION_STATEMENT not in citations:
        raise EdgarFormatError(f"the citation sheet no longer names {CITATION_STATEMENT!r}: another version?")


def read_sheet(rows: list[tuple], sheet: str) -> list[Row]:
    """Every coded row of a by-country sheet: the countries, then (after a blank row) EU27, then GLOBAL TOTAL."""
    years = tuple(range(FIRST_YEAR, LAST_YEAR + 1))
    header = rows[0]
    if tuple(header[:2]) != HEADER or tuple(header[2:]) != years:
        raise EdgarFormatError(
            f"{sheet}: header {header[:3]}...{header[-1:]} != {HEADER} + years {years[0]}-{years[-1]}"
        )
    out: list[Row] = []
    for n, r in enumerate(rows[1:], start=2):
        if all(c is None for c in r):
            continue
        code, name, values = r[0], r[1], r[2:]
        if not isinstance(code, str) or not isinstance(name, str):
            raise EdgarFormatError(f"{sheet} row {n}: code {code!r} or name {name!r} is not text")
        for y, v in zip(years, values, strict=True):
            if isinstance(v, bool) or not isinstance(v, int | float):
                raise EdgarFormatError(f"{sheet} row {n} ({code}) {y}: {v!r} is not a number")
        out.append(Row(code, name, tuple(float(v) for v in values)))
    codes = [r.code for r in out]
    if len(set(codes)) != len(codes):
        raise EdgarFormatError(f"{sheet}: a code appears twice")
    if codes[-2:] != ["EU27", "GLOBAL TOTAL"]:
        raise EdgarFormatError(f"{sheet}: the last rows are {codes[-2:]}, not EU27 then GLOBAL TOTAL")
    return out


def entity_for(row: Row) -> tuple[str, str | None]:
    """(our entity code, note) for an EDGAR row. Raises for an undeclared code."""
    if row.code in ALIASES:
        name, ours = ALIASES[row.code]
        if row.name != name:
            raise EdgarFormatError(f"row {row.code} is named {row.name!r}, not {name!r}")
        return geo.resolve(ours, "iso3"), None
    if row.code in COMBINED:
        if row.name != COMBINED[row.code]:
            raise EdgarFormatError(f"row {row.code} is named {row.name!r}, not {COMBINED[row.code]!r}")
        return geo.resolve(row.code, "iso3"), f'EDGAR reports this row as "{row.name}": the value covers them together.'
    code = geo.resolve(row.code, "iso3")
    kind = geo.entity(code).kind
    if kind not in ("country", "territory"):
        raise EdgarFormatError(f"row {row.code} ({row.name!r}) resolves to {code}, a {kind}; declare it in ALIASES")
    return code, None


def observations(rows: list[Row]) -> list[Observation]:
    """Observations for every row."""
    obs: list[Observation] = []
    for row in rows:
        code, note = entity_for(row)
        for year, v in zip(range(FIRST_YEAR, LAST_YEAR + 1), row.values, strict=True):
            fast = year in FAST_TRACK_YEARS
            notes = " ".join(n for n in (note, FAST_TRACK_NOTE if fast else None) if n) or None
            obs.append(
                Observation(
                    entity=code,
                    period=f"{year:04d}",
                    value=v,
                    status="preliminary" if fast else "final",
                    note=notes,
                )
            )
    return obs


# --- report PDF --------------------------------------------------------------------------------------------------


def report_pages(pdf: bytes, numbers: set[int]) -> dict[int, str]:
    """The text pypdf extracts from the named pages (1-based) of the report, after checking its page count."""
    from pypdf import PdfReader

    # As in textmatch.pdf_pages_text: pypdf logs font-encoding warnings that do not affect the text.
    logging.getLogger("pypdf").setLevel(logging.ERROR)
    reader = PdfReader(io.BytesIO(pdf))
    if len(reader.pages) != REPORT_PAGES:
        raise EdgarFormatError(f"the report has {len(reader.pages)} pages, not {REPORT_PAGES}: another edition?")
    return {n: reader.pages[n - 1].extract_text() or "" for n in sorted(numbers)}


def verify(pages: dict[int, str], *quotes: Quote) -> None:
    for q in quotes:
        if not textmatch.contains(pages[q.page], q.text):
            raise EdgarFormatError(f"PDF page {q.page} no longer says {q.text[:80]!r}...; re-read the report")


def _collapsed(text: str) -> str:
    return " ".join(text.split())


def gas_shares(pages: dict[int, str]) -> dict[str, float]:
    """The four shares of the 2025 world total, quoted from PDF page 11 and confirmed on the World profile."""
    verify(pages, GAS_QUOTE)
    # The quote names the gases in this order: fossil CO₂, CH₄, N₂O, F-gases.
    m = re.fullmatch(
        r"In 2025, fossil CO₂ remained the dominant component of global GHG emissions, accounting for "
        r"(\d+\.\d)% of the total\. CH₄ contributed (\d+\.\d)%, N₂O (\d+\.\d)% and F-gases (\d+\.\d)%\.",
        GAS_QUOTE.text,
    )
    assert m is not None
    quoted = dict(zip((gid for gid, _, _ in GASES), m.groups(), strict=True))
    profile = _collapsed(pages[WORLD_PROFILE_PAGE])
    if not profile.startswith("WORLD GHG emissions by sector GHG % in 2025"):
        raise EdgarFormatError(f"PDF page {WORLD_PROFILE_PAGE} is not the World profile: {profile[:60]!r}")
    for gid, _, chart_name in GASES:
        hits = re.findall(rf"(?<![\w-]){re.escape(chart_name)} (\d+\.\d)(?= )", profile)
        if hits != [quoted[gid]]:
            raise EdgarFormatError(
                f"the World profile shows {chart_name} {hits}, the main findings {quoted[gid]}%: they must agree"
            )
    return {gid: float(v) for gid, v in quoted.items()}


def sector_changes(pages: dict[int, str]) -> tuple[dict[tuple[str, str], int], str]:
    """{(sector, since year): whole percent} from the World profile's table, and the table's text as printed."""
    profile = _collapsed(pages[WORLD_PROFILE_PAGE])
    if not profile.startswith("WORLD GHG emissions by sector"):
        raise EdgarFormatError(f"PDF page {WORLD_PROFILE_PAGE} is not the World profile: {profile[:60]!r}")
    pattern = " ".join(label for _, label in SECTOR_COLUMNS) + "".join(
        rf" {re.escape(label)}" + r" ([+-]?\d+)%" * len(SECTOR_COLUMNS) for _, label in SECTORS
    )
    found = list(re.finditer(pattern, profile))
    if len(found) != 1:
        raise EdgarFormatError(f"the World profile's sector table was found {len(found)} times, not once")
    m = found[0]
    values = iter(int(g) for g in m.groups())
    out = {(sid, since): next(values) for sid, _ in SECTORS for since, _ in SECTOR_COLUMNS}
    return out, m.group(0)


# --- results -----------------------------------------------------------------------------------------------------


def _check_urls(files: dict[str, InputFile], *inputs: Input) -> None:
    expected = {BOOKLET.key: BOOKLET_URL, REPORT.key: REPORT_URL}
    for i in inputs:
        url = str(files[i.key].snapshot.url) if files[i.key].snapshot.url else None
        if url != expected[i.key]:
            raise EdgarFormatError(f"{i.key} came from {url!r}, not the EDGAR 2026 file {expected[i.key]}")


def _scope_pages(files: dict[str, InputFile], extra: set[int]) -> dict[int, str]:
    pages = report_pages(files[REPORT.key].path.read_bytes(), {11, 50, 74} | extra)
    verify(pages, LULUCF_QUOTE, BUNKERS_QUOTE, FAST_TRACK_QUOTE)
    return pages


def _sheet_result(files: dict[str, InputFile], sheet: str, unit_words: str) -> Result:
    _check_urls(files, BOOKLET, REPORT)
    sheets = _cells(files[BOOKLET.key].path.read_bytes(), (sheet,))
    check_statements(sheets)
    rows = read_sheet(sheets[sheet], sheet)
    _scope_pages(files, set())
    obs = observations(rows)
    booklet, report = files[BOOKLET.key].snapshot, files[REPORT.key].snapshot
    present = {r.code for r in rows}
    combined = ", ".join(f'{c} "{n}"' for c, n in COMBINED.items() if c in present)
    aliased = ", ".join(f'{c} "{n}" is {ours}' for c, (n, ours) in ALIASES.items() if c in present)
    return Result(
        observations=obs,
        vintage=VINTAGE,
        steps=[
            f"Read the sheet '{sheet}' of EDGAR_2026_GHG_booklet_2026.xlsx (sha256 {booklet.sha256[:12]}…): one row "
            f"per country or aggregate, one column per year {FIRST_YEAR}–{LAST_YEAR}, in {unit_words} as the "
            "workbook's info sheet states. Each value is published exactly as stored in its cell: nothing is "
            "rounded, added up, divided or converted, because the licence of the fossil CO₂ part (IEA-EDGAR CO2, "
            "CC BY-NC-ND 4.0) allows no derivatives.",
            f"Matched EDGAR's codes to ours: ISO 3166-1 alpha-3 codes directly, except {aliased} (ANT is the old "
            "code of the Netherlands Antilles; EDGAR names the row Curaçao, here and on its report profile; SCG is "
            "one row for Serbia and Montenegro together). Rows that cover a country together with a neighbour "
            f"({combined}) are published under the code EDGAR gives them, with EDGAR's row name as a note on every "
            "value.",
            f"Checked against the report PDF (sha256 {report.sha256[:12]}…): the totals exclude land use, land-use "
            f"change and forestry (page {LULUCF_QUOTE.page}: {LULUCF_QUOTE.text!r}), and the world total includes "
            f"international shipping and aviation (page {BUNKERS_QUOTE.page}: {BUNKERS_QUOTE.text!r}).",
            f"Values for {min(FAST_TRACK_YEARS)} and {max(FAST_TRACK_YEARS)} are marked preliminary, because the "
            f"report says (page {FAST_TRACK_QUOTE.page}) {FAST_TRACK_QUOTE.text!r}.",
        ],
    )


def _run_totals(files: dict[str, InputFile]) -> Result:
    return _sheet_result(files, TOTALS_SHEET, "million tonnes of CO₂-equivalent per year (Mt CO2eq/yr)")


def _run_per_capita(files: dict[str, InputFile]) -> Result:
    return _sheet_result(files, PER_CAPITA_SHEET, "tonnes of CO₂-equivalent per person per year (t CO2eq/cap/yr)")


def _run_gas_shares(files: dict[str, InputFile]) -> Result:
    _check_urls(files, REPORT)
    pages = _scope_pages(files, {WORLD_PROFILE_PAGE})
    shares = gas_shares(pages)
    sha = files[REPORT.key].snapshot.sha256
    return Result(
        observations=[
            Observation(
                entity="WLD",
                period=f"{LAST_YEAR}",
                value=shares[gid],
                status="preliminary",
                note=FAST_TRACK_NOTE,
                dims={"gas": gid},
            )
            for gid, _, _ in GASES
        ],
        vintage=VINTAGE,
        steps=[
            f"Quoted from {GAS_LOCATOR} of the report. The quote was found in the text of page {GAS_QUOTE.page} of "
            f"the PDF snapshot (sha256 {sha[:12]}…) before publishing.",
            # No values in the step: this class's numbers stay out of the public catalogue.
            "Value: each percentage as printed, one per gas ("
            + ", ".join(label for _, label, _ in GASES)
            + "). They are shares of the world total excluding land use, land-use change and forestry (page "
            f"{LULUCF_QUOTE.page}).",
            f"The World profile in Annex 5 (page {WORLD_PROFILE_PAGE}) prints the same four shares beside its chart; "
            "the build stops unless they are identical.",
            f"Marked preliminary: {LAST_YEAR} is a Fast-Track year (page {FAST_TRACK_QUOTE.page}).",
        ],
        published_value=PublishedValueRef(document=SOURCE, locator=GAS_LOCATOR, quote=GAS_QUOTE.text),
    )


def _run_sector_changes(files: dict[str, InputFile]) -> Result:
    _check_urls(files, REPORT)
    pages = _scope_pages(files, {WORLD_PROFILE_PAGE})
    changes, table = sector_changes(pages)
    sha = files[REPORT.key].snapshot.sha256
    return Result(
        observations=[
            Observation(
                entity="WLD",
                period=f"{LAST_YEAR}",
                value=float(changes[(sid, since)]),
                status="preliminary",
                note=FAST_TRACK_NOTE,
                dims={"sector": sid, "since": since},
            )
            for sid, _ in SECTORS
            for since, _ in SECTOR_COLUMNS
        ],
        vintage=VINTAGE,
        steps=[
            f"Read from {SECTOR_LOCATOR} of the report, in the text of page {WORLD_PROFILE_PAGE} of the PDF snapshot "
            f"(sha256 {sha[:12]}…): the columns '2025 vs 1990', '2025 vs 2005' and '2025 vs 2024' for seven sectors "
            "and all sectors, exactly as labelled in SECTORS. Each value is the whole percentage printed, including "
            "where the report prints a change that rounds to zero.",
            "Sectors are the report's: 'Industrial Combustion and Processes' is one row in this table. The changes are "
            f"in world emissions excluding land use, land-use change and forestry (page {LULUCF_QUOTE.page}), "
            f"including international shipping and aviation (page {BUNKERS_QUOTE.page}).",
            f"Marked preliminary: {LAST_YEAR} is a Fast-Track year (page {FAST_TRACK_QUOTE.page}).",
        ],
        published_value=PublishedValueRef(document=SOURCE, locator=SECTOR_LOCATOR, quote=table),
    )


# --- publisher cross-checks (the report's own profile pages, same vintage) ----------------------------------------


def _profile_check(entity: str, period: str, stated: str, row: str) -> PublisherCheck:
    return PublisherCheck(
        source_id=SOURCE, vintage=VINTAGE, entity=entity, period=period, stated=stated, quote=row, url=REPORT_URL
    )


# (entity, period, total Mt CO2eq/yr, per person t CO2eq/cap/yr, the profile row as printed, PDF page)
PROFILE_ROWS: tuple[tuple[str, str, str, str, str, int], ...] = (
    ("WLD", "2025", "54149.174", "6.616", "2025 54149.174 6.616 0.301 8.185G", 75),
    ("WLD", "1990", "32137.148", "6.030", "1990 32137.148 6.030 0.535 5.330G", 75),
    ("CHN", "2025", "15980.974", "11.107", "2025 15980.974 11.107 0.453 1.439G", 119),
    ("USA", "2025", "6018.301", "17.533", "2025 6018.301 17.533 0.230 343.255M", 278),
)


# --- transforms --------------------------------------------------------------------------------------------------


TOTALS_SCOPE = Scope(
    geography="Every country EDGAR reports, international aviation and shipping, the EU27 and the world",
    gwp="AR5-GWP100",
    lulucf="excluded",
    bunkers="included",
    basis=GAS_BASIS + " Country values exclude international aviation and shipping, which are their own entities; "
    "the world total includes them. Some rows cover a country together with a neighbour, as noted on their values.",
)
PER_CAPITA_SCOPE = Scope(
    geography="Every country EDGAR reports, the EU27 and the world",
    gwp="AR5-GWP100",
    lulucf="excluded",
    bunkers="included",
    basis=GAS_BASIS + " Country values exclude international aviation and shipping; the world value includes them. "
    "Some rows cover a country together with a neighbour, as noted on their values.",
)


def transforms(paths: Paths) -> list[Transform]:
    here = Path(__file__)
    gas_dim = Dimension(id="gas", label="Gas", values=[DimensionValue(id=i, label=label) for i, label, _ in GASES])
    sector_dim = Dimension(
        id="sector", label="Sector", values=[DimensionValue(id=i, label=label) for i, label in SECTORS]
    )
    since_dim = Dimension(
        id="since",
        label="Compared with",
        values=[DimensionValue(id=i, label=f"{LAST_YEAR} compared with {i}") for i, _ in SECTOR_COLUMNS],
    )
    years = LAST_YEAR - FIRST_YEAR + 1
    return [
        Transform(
            spec=Spec(
                id="ghg.edgar-2026.total-by-country",
                title="Greenhouse gas emissions by country (EDGAR)",
                description="Greenhouse gases released each year since 1970 by every country, by international "
                "aviation and shipping, by the EU27 and by the world, in carbon dioxide equivalent, as published in "
                "EDGAR 2026. Land use, land-use change and forestry are not included. Values for 2024 and 2025 are "
                "EDGAR's Fast-Track estimates and will be revised.",
                kind="series",
                unit=MT_CO2E_YR,
                display=Display(decimals=1),
                scope=TOTALS_SCOPE,
                geo_coverage="mixed",
                headline_entity="WLD",
            ),
            inputs=(BOOKLET, REPORT),
            run=_run_totals,
            module_file=here,
            validation=Validation(min_rows=200 * years, value_range=(0.0, 60000.0)),
            checks=tuple(_profile_check(e, p, total, row) for e, p, total, _, row, _ in PROFILE_ROWS),
        ),
        Transform(
            spec=Spec(
                id="ghg-per-capita.edgar-2026.by-country",
                title="Greenhouse gas emissions per person by country (EDGAR)",
                description="Greenhouse gases released each year since 1970 per person living in each country, the "
                "EU27 and the world, in carbon dioxide equivalent, as published in EDGAR 2026. Land use, land-use "
                "change and forestry are not included. Values for 2024 and 2025 are EDGAR's Fast-Track estimates "
                "and will be revised.",
                kind="series",
                unit=T_CO2E_PERSON_YR,
                display=Display(decimals=2),
                scope=PER_CAPITA_SCOPE,
                geo_coverage="mixed",
                headline_entity="WLD",
            ),
            inputs=(BOOKLET, REPORT),
            run=_run_per_capita,
            module_file=here,
            validation=Validation(min_rows=200 * years, value_range=(0.0, 400.0)),
            checks=tuple(_profile_check(e, p, pc, row) for e, p, _, pc, row, _ in PROFILE_ROWS),
        ),
        Transform(
            spec=Spec(
                id="ghg-gas-share.edgar-2026.global",
                title="Share of each gas in world greenhouse gas emissions, 2025 (EDGAR)",
                description="How much of the world's greenhouse gas emissions in 2025 was fossil carbon dioxide, "
                "methane, nitrous oxide and fluorinated gases, as stated in the EDGAR 2026 report, measured in "
                "carbon dioxide equivalent. Land use, land-use change and forestry are not included. 2025 is a "
                "Fast-Track estimate.",
                kind="published-value",
                unit=PERCENT_OF_TOTAL,
                display=Display(decimals=1),
                scope=Scope(
                    geography="World",
                    gwp="AR5-GWP100",
                    lulucf="excluded",
                    bunkers="included",
                    basis=GAS_BASIS,
                ),
                geo_coverage="global-only",
                headline_entity="WLD",
                dimensions=(gas_dim,),
                headline_dims=(("gas", "fossil-co2"),),
            ),
            inputs=(REPORT,),
            run=_run_gas_shares,
            module_file=here,
            validation=Validation(min_rows=len(GASES), value_range=(0.0, 100.0)),
        ),
        Transform(
            spec=Spec(
                id="ghg-sector-change.edgar-2026.global",
                title="Change in world greenhouse gas emissions by sector (EDGAR)",
                description="How much world greenhouse gas emissions from each sector changed by 2025 compared with "
                "1990, 2005 and 2024, as printed in the EDGAR 2026 report's world profile, in whole percent. Land "
                "use, land-use change and forestry are not included. 2025 is a Fast-Track estimate.",
                kind="published-value",
                unit=PERCENT_CHANGE,
                display=Display(decimals=0),
                scope=Scope(
                    geography="World",
                    baseline="The same sector's emissions in 1990, 2005 or 2024",
                    gwp="AR5-GWP100",
                    lulucf="excluded",
                    bunkers="included",
                    basis=GAS_BASIS + " Sectors as defined in the report; international aviation and shipping are "
                    "counted in transport.",
                ),
                geo_coverage="global-only",
                headline_entity="WLD",
                dimensions=(sector_dim, since_dim),
                headline_dims=(("sector", "all-sectors"), ("since", "1990")),
            ),
            inputs=(REPORT,),
            run=_run_sector_changes,
            module_file=here,
            validation=Validation(min_rows=len(SECTORS) * len(SECTOR_COLUMNS), value_range=(-100.0, 1000.0)),
        ),
    ]
