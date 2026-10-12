"""Lancet Countdown 2025 report, world values: heat-related deaths (1.1.5), potential work hours lost to heat
(1.1.3), heatwave days per person observed and attributable to climate change (1.1.1), heatwave days of infants and
adults over 65 (1.1.1), deaths from wildfire smoke (1.2.1), the share of land in extreme drought (1.2.2), extreme
precipitation events over land (1.2.3) and deaths from outdoor air pollution by sector and fuel (3.2.1). And by
country: heatwave days observed and attributable to climate change (1.1.1) and work hours lost per worker (1.1.3).

Inputs: the indicator workbooks of source lancet-countdown-2025. Each has a DATA GUIDANCE sheet (report, indicator,
the report's key finding, column units, licence and suggested citation) and one sheet per geography; the Global sheet
is read (1.2.3 has two global sheets; only its event counts are read), and for 1.1.1 attributable and 1.1.3 also the
country sheet. Country rows carry ISO3 codes, each resolved through the entity crosswalk; the producer's country
names, WHO regions and HDI levels are not read.

Vintage. The guidance sheet names the report ("The 2025 report of the Lancet Countdown on health and climate
change"). The vintage is "<year> report". A report year other than REPORT_YEAR stops the build: each report is its
own registry entry (indicator methods change between reports, so vintages are never spliced). The suggested
citation printed in the workbook must equal the registry's citation text, and the report's publication date is
read from it ("published online Oct 29").

Units. Each value column's unit is read from the guidance sheet's column table and must be what this module
assumes: AN, Mortality (1.2.1) and the Deaths columns (3.2.1) deaths; the work-hours-lost columns "Hours in 1000s";
Observed, Counterfactual, Attributable_to_CC and the average heatwave days of infants and over-65s "days"; the drought
shares "%"; the 1.2.3 event count "Count". A changed unit, column header or sheet layout stops the build until a
person re-reads the workbook.

Values. Deaths, heatwave days, drought shares and event counts are published as in the Global sheets; no missing
value is filled. Work hours lost are divided by one million (exact decimal arithmetic on the stored digits) to give
billions of hours, the unit of the report's own statement. The deaths from fossil fuels are the one sum: coal plus
liquids and gas over the seven sectors of the 3.2.1 Global sheet, in whole deaths, as the report's own key finding
adds them. Identities of the files are checked rather than assumed: TotalSunAgCon equals the sum of the four sector
columns; Observed equals Counterfactual plus Attributable_to_CC; and the land share with at least one month of
extreme drought is never below the share with three months, nor that below six.

Country gaps. The 1.1.1 country sheet must be complete (every country, every year of the Global sheet, each row's
identity). In the 1.1.3 country sheet a value per worker may be a "#DIV/0!" (no one employed in the four sectors) or
blank (no data); both are published as missing with that reason, never filled.

Left out, with the reason. 1.2.2's 12-month drought share reads NA in 15 of its 74 years, so it is not published.
The 1.1.1 vulnerable-populations Global sheet's over-75 columns are not described in the guidance column table and
read 0 in 1980–1999, when the over-65 columns do not; the guidance says a 0 is a value, so they cannot be read as
missing either, and they are not published. 1.2.3's land-share sheet is not published: its rolling windows are
irregular (nine or ten years, the latest "2015-2023" where the key finding says 2015–24) and its "%" columns hold
fractions.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import openpyxl

from envdash import geo
from envdash.models import Dimension, DimensionValue, Display, Observation, Scope, Unit
from envdash.paths import Paths
from envdash.transform import Input, InputFile, PublisherCheck, Result, Spec, Transform, Validation

SOURCE = "lancet-countdown-2025"
REPORT_YEAR = 2025
VINTAGE = f"{REPORT_YEAR} report"
MORTALITY = Input(SOURCE, "heat-mortality-1-1-5")
WORK_HOURS = Input(SOURCE, "work-hours-lost-1-1-3")
HEATWAVE = Input(SOURCE, "heatwave-days-attributable-1-1-1")
VULNERABLE = Input(SOURCE, "heatwave-exposure-vulnerable-1-1-1")
SMOKE = Input(SOURCE, "wildfire-smoke-deaths-1-2-1")
DROUGHT = Input(SOURCE, "drought-land-share-1-2-2")
RAIN = Input(SOURCE, "extreme-precipitation-1-2-3")
AIR = Input(SOURCE, "air-pollution-deaths-3-2-1")

REPORT = re.compile(r"The (\d{4}) report of the Lancet Countdown on health and climate change")
CITATION = (
    "Romanello M, Walawender M, Hsu S-C, et al. The 2025 report of the Lancet Countdown on health and climate "
    "change. Lancet 2025; published online Oct 29. https://doi.org/10.1016/S0140-6736(25)01919-1."
)
PUBLISHED_ONLINE = re.compile(r"Lancet (\d{4}); published online ([A-Z][a-z]{2} \d{1,2})\.")


class LancetFormatError(ValueError):
    pass


# --- reading the workbooks ----------------------------------------------------------------------------------------


@dataclass(frozen=True)
class Guidance:
    report_year: int
    indicator: str
    key_finding: str
    citation: str
    published: date
    units: dict[str, str]
    """Column header (as printed in the guidance table, stripped) -> its Units cell."""


def _cells(row: tuple) -> list:
    return [c for c in row if c is not None and not (isinstance(c, str) and not c.strip())]


def _after(rows: list[list], heading: str) -> list:
    for i, r in enumerate(rows):
        if r and isinstance(r[0], str) and r[0].strip() == heading:
            for nxt in rows[i + 1 :]:
                if nxt:
                    return nxt
    raise LancetFormatError(f"DATA GUIDANCE has no {heading!r} section")


def read_guidance(wb: openpyxl.Workbook) -> Guidance:
    if "DATA GUIDANCE" not in wb.sheetnames:
        raise LancetFormatError(f"no DATA GUIDANCE sheet in {wb.sheetnames}")
    rows = [_cells(r) for r in wb["DATA GUIDANCE"].iter_rows(values_only=True)]
    fields = {r[0].strip(): r[1] for r in rows if len(r) == 2 and isinstance(r[0], str)}
    m = REPORT.fullmatch(str(fields.get("Report", "")).strip())
    if not m:
        raise LancetFormatError(f"DATA GUIDANCE Report is {fields.get('Report')!r}; expected {REPORT.pattern!r}")
    key_finding = _after(rows, "1. Report Key Findings")
    citation = _after(rows, "3. Suggested Citation")
    if len(key_finding) != 1 or len(citation) != 1:
        raise LancetFormatError("DATA GUIDANCE key finding and suggested citation should each be one cell")
    cite = str(citation[0]).strip()
    pm = PUBLISHED_ONLINE.search(cite)
    if not pm:
        raise LancetFormatError(f"cannot read the publication date from the suggested citation {cite!r}")
    published = datetime.strptime(f"{pm.group(2)} {pm.group(1)}", "%b %d %Y").date()
    units: dict[str, str] = {}
    for r in rows:
        # A column letter, or a range of them ("H-K", "L - O" in the 3.2.1 guidance).
        if len(r) == 4 and isinstance(r[0], str) and re.fullmatch(r"[A-Z](\s*-\s*[A-Z])?", r[0].strip()):
            units[str(r[1]).strip()] = str(r[2]).strip()
    return Guidance(
        report_year=int(m.group(1)),
        indicator=str(fields.get("Indicator", "")).strip(),
        key_finding=str(key_finding[0]).strip(),
        citation=cite,
        published=published,
        units=units,
    )


def _open(raw: bytes) -> openpyxl.Workbook:
    return openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)


def _check_guidance(g: Guidance, indicator: str, units: dict[str, str], name: str) -> None:
    if g.report_year != REPORT_YEAR:
        raise LancetFormatError(
            f"{name}: workbook is from the {g.report_year} report; {SOURCE} holds the {REPORT_YEAR} report. A new "
            "report is a new registry entry."
        )
    if g.citation != CITATION:
        raise LancetFormatError(f"{name}: suggested citation changed to {g.citation!r}")
    if g.indicator != indicator:
        raise LancetFormatError(f"{name}: workbook is {g.indicator!r}, expected {indicator!r}")
    for col, unit in units.items():
        if g.units.get(col) != unit:
            raise LancetFormatError(f"{name}: guidance gives column {col!r} units {g.units.get(col)!r}, not {unit!r}")


def read_table(wb: openpyxl.Workbook, sheet: str, columns: tuple[str, ...], name: str) -> list[tuple]:
    if sheet not in wb.sheetnames:
        raise LancetFormatError(f"{name}: no sheet {sheet!r} in {wb.sheetnames}")
    rows = list(wb[sheet].iter_rows(values_only=True))
    if not rows or tuple(rows[0]) != columns:
        raise LancetFormatError(f"{name}: {sheet!r} columns {rows[0] if rows else None} != expected {columns}")
    out = []
    for r in rows[1:]:
        if any(c is None for c in r):
            raise LancetFormatError(f"{name}: {sheet!r} row {r} has an empty cell")
        if not isinstance(r[0], int) or isinstance(r[0], bool):
            raise LancetFormatError(f"{name}: {sheet!r} Year {r[0]!r} is not a whole number")
        if any(not isinstance(c, int | float) or isinstance(c, bool) for c in r[1:]):
            raise LancetFormatError(f"{name}: {sheet!r} row {r} has a non-numeric value")
        out.append(tuple(r))
    years = [r[0] for r in out]
    if years != sorted(set(years)):
        raise LancetFormatError(f"{name}: {sheet!r} years are not unique and increasing")
    return out


def dec(x: int | float) -> Decimal:
    """The number as stored in the workbook (shortest repr of the double), as an exact decimal."""
    return Decimal(repr(x)) if isinstance(x, float) else Decimal(x)


# --- 1.1.5 heat-related deaths ------------------------------------------------------------------------------------

MORTALITY_SHEET = "2025 Report Data_Global"
MORTALITY_COLUMNS = ("Year", "AF", "AN")


def parse_mortality(raw: bytes) -> tuple[list[Observation], Guidance]:
    wb = _open(raw)
    g = read_guidance(wb)
    _check_guidance(g, "Indicator 1.1.5: Heat Related Mortality", {"AN": "Deaths", "AF": "%"}, "1.1.5")
    rows = read_table(wb, MORTALITY_SHEET, MORTALITY_COLUMNS, "1.1.5")
    obs = []
    for year, _af, an in rows:
        if not isinstance(an, int):
            raise LancetFormatError(f"1.1.5: AN for {year} is {an!r}, not a whole number of deaths")
        obs.append(Observation(entity="WLD", period=f"{year:04d}", value=float(an)))
    return obs, g


def _mortality(files: dict[str, InputFile]) -> Result:
    obs, g = parse_mortality(files[MORTALITY.key].path.read_bytes())
    return Result(
        observations=obs,
        vintage=VINTAGE,
        year=str(g.report_year),
        date_published=g.published.isoformat(),
        steps=[
            f"Read the AN column (number of deaths attributable to heat) of the sheet {MORTALITY_SHEET!r} of the "
            f"indicator 1.1.5 workbook of the Lancet Countdown {g.report_year} report, one value a year "
            f"{obs[0].period}–{obs[-1].period}. Values are published as in the workbook.",
            "The workbook's guidance sheet gives the unit of AN as deaths. The share of all deaths (AF) and the "
            "regional sheets are not published here.",
        ],
    )


# --- 1.1.3 potential work hours lost ------------------------------------------------------------------------------

WORK_SHEET = "2025 Report Data_Global"
WORK_COLUMNS = (
    "Year",
    "EmplPop15+",
    "WHL200Serv",
    "WHL300Manuf",
    "WHL400sunAgr",
    "WHL400sunConstr",
    "TotalSunAgCon",
    "TotalSunWHLpp",
)
# (dimension value, label, Global-sheet column, column header as printed in the guidance table)
SECTORS = (
    ("total", "All four sectors", "TotalSunAgCon", "TotalsunAgCon"),
    ("services", "Services (indoors or in shade)", "WHL200Serv", "WHL200Serv"),
    ("manufacturing", "Manufacturing (indoors or in shade)", "WHL300Manuf", "WHL300Manuf"),
    ("agriculture", "Agriculture (in the sun)", "WHL400sunAgr", "WHL400sunAgr"),
    ("construction", "Construction (in the sun)", "WHL400sunConstr", "WHL400sunConstr"),
)
THOUSANDS_PER_BILLION = Decimal(1_000_000)


def parse_work_hours(raw: bytes) -> tuple[list[Observation], Guidance]:
    wb = _open(raw)
    g = read_guidance(wb)
    _check_guidance(
        g,
        "Indicator 1.1.3: Change in Labour Capacity",
        {guide: "Hours in 1000s" for *_, guide in SECTORS},
        "1.1.3",
    )
    rows = read_table(wb, WORK_SHEET, WORK_COLUMNS, "1.1.3")
    idx = {c: i for i, c in enumerate(WORK_COLUMNS)}
    by_sector: dict[str, list[Observation]] = {s: [] for s, *_ in SECTORS}
    for r in rows:
        thousands = {col: dec(r[idx[col]]) for _, _, col, _ in SECTORS}
        parts = sum((thousands[c] for c in ("WHL200Serv", "WHL300Manuf", "WHL400sunAgr", "WHL400sunConstr")), Decimal())
        if abs(parts - thousands["TotalSunAgCon"]) > thousands["TotalSunAgCon"] * Decimal("1e-9"):
            raise LancetFormatError(
                f"1.1.3: {r[0]} TotalSunAgCon {thousands['TotalSunAgCon']} is not the sum of the four sectors {parts}"
            )
        for sector, _, col, _ in SECTORS:
            by_sector[sector].append(
                Observation(
                    entity="WLD",
                    period=f"{r[0]:04d}",
                    value=float(thousands[col] / THOUSANDS_PER_BILLION),
                    dims={"sector": sector},
                )
            )
    return [o for s, *_ in SECTORS for o in by_sector[s]], g


def _work_hours(files: dict[str, InputFile]) -> Result:
    obs, g = parse_work_hours(files[WORK_HOURS.key].path.read_bytes())
    years = sorted({o.period for o in obs})
    return Result(
        observations=obs,
        vintage=VINTAGE,
        year=str(g.report_year),
        date_published=g.published.isoformat(),
        steps=[
            f"Read the sheet {WORK_SHEET!r} of the indicator 1.1.3 workbook of the Lancet Countdown {g.report_year} "
            f"report, {years[0]}–{years[-1]}: potential work hours lost in services (WHL200Serv), manufacturing "
            "(WHL300Manuf), agriculture (WHL400sunAgr), construction (WHL400sunConstr) and in total (TotalSunAgCon), "
            "in thousands of hours as the guidance sheet states.",
            "Checked that the total equals the sum of the four sectors in every year (to one part in a billion).",
            "Divided every value by one million to give billions of hours (exact decimal arithmetic on the values "
            "as stored in the workbook). Employment and hours lost per worker are not published here.",
        ],
        changes="work hours lost converted from thousands of hours to billions of hours.",
    )


# --- 1.1.1 heatwave days attributable to climate change -----------------------------------------------------------

HEATWAVE_SHEET = "2025_Report _Data_Global"
HEATWAVE_COLUMNS = ("Year", "Observed", "Counterfactual", "Attributable_to_CC")
SCENARIOS = (
    ("observed", "Observed", "Observed"),
    ("counterfactual", "Expected without human-caused climate change", "Counterfactual"),
    ("attributable", "Attributable to climate change", "Attributable_to_CC"),
)


def parse_heatwave(raw: bytes) -> tuple[list[Observation], Guidance]:
    wb = _open(raw)
    g = read_guidance(wb)
    _check_guidance(
        g,
        "Indicator 1.1.1: Exposure of Vulnerable Populations to Heatwaves -- attributable heatwave days",
        {"Observed": "days", "Counterfactual": "days", "Attributable_to_CC": "days"},
        "1.1.1",
    )
    rows = read_table(wb, HEATWAVE_SHEET, HEATWAVE_COLUMNS, "1.1.1")
    by: dict[str, list[Observation]] = {s: [] for s, *_ in SCENARIOS}
    for year, observed, counterfactual, attributable in rows:
        o, c, a = dec(observed), dec(counterfactual), dec(attributable)
        if abs(o - c - a) > Decimal("1e-9"):
            raise LancetFormatError(f"1.1.1: {year} Observed {o} != Counterfactual {c} + Attributable_to_CC {a}")
        for scenario, value in (("observed", o), ("counterfactual", c), ("attributable", a)):
            by[scenario].append(
                Observation(entity="WLD", period=f"{year:04d}", value=float(value), dims={"scenario": scenario})
            )
    return [o for s, *_ in SCENARIOS for o in by[s]], g


def _heatwave(files: dict[str, InputFile]) -> Result:
    obs, g = parse_heatwave(files[HEATWAVE.key].path.read_bytes())
    years = sorted({o.period for o in obs})
    return Result(
        observations=obs,
        vintage=VINTAGE,
        year=str(g.report_year),
        date_published=g.published.isoformat(),
        steps=[
            f"Read the sheet {HEATWAVE_SHEET!r} of the indicator 1.1.1 (attributable heatwave days) workbook of the "
            f"Lancet Countdown {g.report_year} report, {years[0]}–{years[-1]}: heatwave days observed, expected "
            "without climate change (Counterfactual) and attributable to climate change (Attributable_to_CC), "
            "published as in the workbook.",
            "Checked that observed days equal expected days plus attributable days in every year.",
        ],
    )


# --- the country sheets -------------------------------------------------------------------------------------------

COUNTRY_KEYS = ("Country", "ISO3", "WHO Region", "HDI Level", "Year")
"""The five columns before the values in each country sheet (the guidance column table's spelling; 1.1.1's sheet
writes 'WHO region' and 'HDI level'). Only ISO3 and Year are read: the producer's names, regions and HDI levels are
not published."""
MISSING = {
    None: "The producer's sheet leaves this cell blank, which its guidance says means no data was available.",
    "#DIV/0!": "The producer's sheet holds a division by zero here: it gives no one employed in the four sectors, so "
    "there is no value per worker.",
}
"""The empty cells a country sheet may hold, with what each means; any other text stops the build. The guidance's
"N/A" (could not be calculated) does not occur in the value columns read here."""


def read_countries(wb: openpyxl.Workbook, sheet: str, columns: tuple, name: str) -> list[tuple[str, int, tuple]]:
    """(our entity code, year, the row's cells after Year) for each row of a country sheet. ISO3 must be in the entity
    crosswalk (resolved, never guessed); each country appears in one block of rows, years unique and increasing."""
    if sheet not in wb.sheetnames:
        raise LancetFormatError(f"{name}: no sheet {sheet!r} in {wb.sheetnames}")
    rows = list(wb[sheet].iter_rows(values_only=True))
    if not rows or tuple(rows[0]) != columns:
        raise LancetFormatError(f"{name}: {sheet!r} columns {rows[0] if rows else None} != expected {columns}")
    out: list[tuple[str, int, tuple]] = []
    codes: dict[str, str] = {}
    last: tuple[str, int] | None = None
    for r in rows[1:]:
        iso, year = r[1], r[4]
        if not isinstance(iso, str) or not isinstance(year, int) or isinstance(year, bool):
            raise LancetFormatError(f"{name}: {sheet!r} row {r} has no ISO3 code or whole-number Year")
        if last is None or iso != last[0]:
            if iso in codes:
                raise LancetFormatError(f"{name}: {sheet!r} rows of {iso} are not together")
            codes[iso] = geo.resolve(iso, "iso3")
        elif year <= last[1]:
            raise LancetFormatError(f"{name}: {sheet!r} {iso} years are not unique and increasing at {year}")
        last = (iso, year)
        out.append((codes[iso], year, tuple(r[5:])))
    if len(set(codes.values())) != len(codes):
        raise LancetFormatError(f"{name}: two ISO3 codes of {sheet!r} name the same entity")
    return out


def _number(v, name: str, where: str) -> Decimal | None:
    """A value cell as an exact decimal, or None for one of the MISSING cells (never filled)."""
    if isinstance(v, int | float) and not isinstance(v, bool):
        return dec(v)
    if v in MISSING:
        return None
    raise LancetFormatError(f"{name}: {where} holds {v!r}, neither a number nor one of {list(MISSING)}")


# --- 1.1.1 heatwave days attributable to climate change, by country -----------------------------------------------

HEATWAVE_COUNTRY_SHEET = "2025_Report_Data_Country"
HEATWAVE_COUNTRY_COLUMNS = ("Country", "ISO3", "WHO region", "HDI level", "Year", *HEATWAVE_COLUMNS[1:])


def parse_heatwave_countries(raw: bytes) -> tuple[list[Observation], Guidance]:
    """The World (Global sheet) and each country's heatwave days, observed, expected without climate change and
    attributable to it. Every country must have exactly the Global sheet's years and every value; a negative
    attributable value (a country with fewer heatwave days than its modelled counterfactual) is published as it is."""
    world, g = parse_heatwave(raw)
    wb = _open(raw)
    rows = read_countries(wb, HEATWAVE_COUNTRY_SHEET, HEATWAVE_COUNTRY_COLUMNS, "1.1.1 countries")
    global_years = sorted({int(o.period) for o in world if o.period})
    by: dict[str, list[Observation]] = {s: [o for o in world if o.dims["scenario"] == s] for s, *_ in SCENARIOS}
    seen: dict[str, list[int]] = {}
    for entity, year, cells in rows:
        where = f"{entity} {year}"
        o, c, a = (_number(v, "1.1.1 countries", where) for v in cells)
        if o is None or c is None or a is None:
            raise LancetFormatError(f"1.1.1 countries: {where} has an empty value; the sheet is expected complete")
        if abs(o - c - a) > Decimal("1e-9"):
            raise LancetFormatError(f"1.1.1 countries: {where} Observed {o} != Counterfactual {c} + Attributable {a}")
        seen.setdefault(entity, []).append(year)
        for scenario, value in (("observed", o), ("counterfactual", c), ("attributable", a)):
            by[scenario].append(
                Observation(entity=entity, period=f"{year:04d}", value=float(value), dims={"scenario": scenario})
            )
    short = {e: ys for e, ys in seen.items() if ys != global_years}
    if short:
        raise LancetFormatError(f"1.1.1 countries: years differ from the Global sheet's {global_years}: {short}")
    return [o for s, *_ in SCENARIOS for o in by[s]], g


def _heatwave_countries(files: dict[str, InputFile]) -> Result:
    obs, g = parse_heatwave_countries(files[HEATWAVE.key].path.read_bytes())
    years = sorted({o.period for o in obs})
    countries = sorted({o.entity for o in obs} - {"WLD"})
    return Result(
        observations=obs,
        vintage=VINTAGE,
        year=str(g.report_year),
        date_published=g.published.isoformat(),
        steps=[
            f"Read the sheets {HEATWAVE_SHEET!r} (the world) and {HEATWAVE_COUNTRY_SHEET!r} ({len(countries)} "
            "countries and territories by ISO3 code) of the indicator 1.1.1 (attributable heatwave days) workbook of "
            f"the Lancet Countdown {g.report_year} report, {years[0]}–{years[-1]}: heatwave days observed, expected "
            "without climate change (Counterfactual) and attributable to climate change (Attributable_to_CC), "
            "published as in the workbook. The producer's country names, WHO regions and HDI levels are not "
            "published.",
            "Mapped each ISO3 code to the site's entity of that code (pipeline/geo/entities.csv); every code is in it.",
            "Checked that observed days equal expected days plus attributable days in every row, and that every "
            "country has a value for each year of the world sheet. A country with fewer heatwave days than its "
            "modelled counterfactual has a negative attributable value, published as it is.",
        ],
    )


# --- 1.1.3 potential work hours lost per worker, world and by country ---------------------------------------------

WORK_COUNTRY_SHEET = "2025 Report Data_Country"
WORK_COUNTRY_COLUMNS = (
    *COUNTRY_KEYS,
    "TotPopn",
    "Pop 15+",
    " EmplServ",
    " EmplManuf",
    " EmplAgr",
    " EmplConstr",
    "EmplPop 15+",
    "WHLpp200",
    "WHLpp300",
    "WHLpp400shade",
    "WHLpp400sun",
    "WHL200Serv",
    "WHL300Manuf",
    "WHL400sunAgr",
    "WHL400sunConstr",
    "TotalsunAgCon",
    "TotalSunWHLpp",
)
PER_WORKER = "TotalSunWHLpp"
"""Mean potential work hours lost per worker (services and manufacturing in shade, agriculture and construction in
the sun), in hours: the total hours lost over the people employed in the four sectors."""


def _per_worker(employed: Decimal | None, thousands: Decimal | None, per: Decimal | None, where: str) -> None:
    """The per-worker value is the total over the employed: checked where all three are given; a division by zero
    must come from no one employed."""
    if per is None:
        return
    if employed is None or thousands is None or employed == 0:
        raise LancetFormatError(f"1.1.3: {where} has hours lost per worker {per} without total hours and employment")
    if abs(thousands * 1000 / employed - per) > per * Decimal("1e-9"):
        raise LancetFormatError(f"1.1.3: {where} hours per worker {per} != {thousands} thousand / {employed} employed")


def parse_work_per_worker(raw: bytes) -> tuple[list[Observation], Guidance, dict[str, list[str]]]:
    """Hours lost per worker for the World (Global sheet) and each country, and the country-years left out, by why."""
    wb = _open(raw)
    g = read_guidance(wb)
    _check_guidance(
        g,
        "Indicator 1.1.3: Change in Labour Capacity",
        {"TotalSunWHLpp": "Hours", "TotalsunAgCon": "Hours in 1000s", "EmplPop 15+": "Persons"},
        "1.1.3",
    )
    idx = {c: i for i, c in enumerate(WORK_COLUMNS)}
    world = []
    for r in read_table(wb, WORK_SHEET, WORK_COLUMNS, "1.1.3"):
        per = dec(r[idx[PER_WORKER]])
        _per_worker(dec(r[idx["EmplPop15+"]]), dec(r[idx["TotalSunAgCon"]]), per, f"WLD {r[0]}")
        world.append(Observation(entity="WLD", period=f"{r[0]:04d}", value=float(per)))
    cols = {c: i - len(COUNTRY_KEYS) for i, c in enumerate(WORK_COUNTRY_COLUMNS)}
    obs: list[Observation] = []
    left_out: dict[str, list[str]] = {}
    for entity, year, cells in read_countries(wb, WORK_COUNTRY_SHEET, WORK_COUNTRY_COLUMNS, "1.1.3 countries"):
        where = f"{entity} {year}"
        employed, thousands, per = (
            _number(cells[cols[c]], "1.1.3 countries", f"{where} {c}")
            for c in ("EmplPop 15+", "TotalsunAgCon", PER_WORKER)
        )
        _per_worker(employed, thousands, per, where)
        if per is None:
            cell = cells[cols[PER_WORKER]]
            if cell == "#DIV/0!" and employed != 0:
                raise LancetFormatError(f"1.1.3 countries: {where} divides by zero with {employed} people employed")
            left_out.setdefault(str(cell), []).append(where)
            obs.append(Observation(entity=entity, period=f"{year:04d}", value=None, missing_reason=MISSING[cell]))
            continue
        obs.append(Observation(entity=entity, period=f"{year:04d}", value=float(per)))
    return world + obs, g, left_out


def _work_per_worker(files: dict[str, InputFile]) -> Result:
    obs, g, left_out = parse_work_per_worker(files[WORK_HOURS.key].path.read_bytes())
    years = sorted({o.period for o in obs})
    countries = sorted({o.entity for o in obs} - {"WLD"})
    gaps = "; ".join(f"{len(v)} {'blank' if why == 'None' else why!r}" for why, v in sorted(left_out.items()))
    return Result(
        observations=obs,
        vintage=VINTAGE,
        year=str(g.report_year),
        date_published=g.published.isoformat(),
        steps=[
            f"Read the column {PER_WORKER} (mean potential work hours lost per worker, in hours as the guidance sheet "
            f"states) of the sheets {WORK_SHEET!r} (the world) and {WORK_COUNTRY_SHEET!r} ({len(countries)} countries "
            f"and territories by ISO3 code) of the indicator 1.1.3 workbook of the Lancet Countdown {g.report_year} "
            f"report, {years[0]}–{years[-1]}, published as in the workbook.",
            "Mapped each ISO3 code to the site's entity of that code (pipeline/geo/entities.csv); every code is in it.",
            "Checked in every row that the value equals the total hours lost (TotalsunAgCon, thousands of hours) over "
            "the people employed in the four sectors (EmplPop 15+).",
            f"Published the country-years with no value per worker as missing, with the reason, never filled: {gaps} "
            "cells. A division by zero is accepted only where the sheet gives no one employed in the four sectors.",
        ],
    )


# --- 1.1.1 heatwave days of infants and adults over 65 ------------------------------------------------------------

VULNERABLE_SHEET = "2025 Report Data_Global"
VULNERABLE_COLUMNS = (
    "Year",
    "exposure_total_infants",
    "exposure_average_infants",
    "exposure_total_65",
    "exposure_average_65",
    "exposure_total_75",
    "exposure_average_75",
)
# (dimension value, label, Global-sheet column, column header as printed in the guidance table)
GROUPS = (
    ("over-65", "Adults over 65", "exposure_average_65", "exposures_average_65"),
    ("infants", "Infants under 1 year old", "exposure_average_infants", "exposure_average_infant"),
)


def parse_vulnerable(raw: bytes) -> tuple[list[Observation], Guidance]:
    wb = _open(raw)
    g = read_guidance(wb)
    _check_guidance(
        g,
        "1.1.1 exposure of vulnerable populations to heatwaves",
        {guide: "days" for *_, guide in GROUPS},
        "1.1.1 vulnerable",
    )
    rows = read_table(wb, VULNERABLE_SHEET, VULNERABLE_COLUMNS, "1.1.1 vulnerable")
    idx = {c: i for i, c in enumerate(VULNERABLE_COLUMNS)}
    obs = [
        Observation(entity="WLD", period=f"{r[0]:04d}", value=float(r[idx[col]]), dims={"group": group})
        for group, _, col, _ in GROUPS
        for r in rows
    ]
    return obs, g


def _vulnerable(files: dict[str, InputFile]) -> Result:
    obs, g = parse_vulnerable(files[VULNERABLE.key].path.read_bytes())
    years = sorted({o.period for o in obs})
    return Result(
        observations=obs,
        vintage=VINTAGE,
        year=str(g.report_year),
        date_published=g.published.isoformat(),
        steps=[
            f"Read the sheet {VULNERABLE_SHEET!r} of the indicator 1.1.1 (vulnerable populations) workbook of the "
            f"Lancet Countdown {g.report_year} report, {years[0]}–{years[-1]}: the average heatwave days per person of "
            "adults over 65 (exposure_average_65) and of infants under one year old (exposure_average_infants), in "
            "days as the guidance sheet states, published as in the workbook.",
            "Total person-days are not published. The over-75 columns are not published either: the guidance "
            "sheet's column table does not describe them, and they read 0 in every year before 2000, when the "
            "over-65 values do not.",
        ],
    )


# --- 1.2.1 deaths from wildfire smoke -----------------------------------------------------------------------------

SMOKE_SHEET = "2025 Report Data_Global"
SMOKE_COLUMNS = ("Year", "Mortality")
SMOKE_URL = (
    "https://lancetcountdown.org/wp-content/uploads/2025/10/"
    "Indicator-1.2.1_Mortality_Data-Download_2025-Lancet-Countdown-Report-1.xlsx"
)
SMOKE_KEY_FINDING = (
    "Deaths from wildfire-derived PM2.5 air pollution reached a record-high 154,000 in 2024, up by 36% from the "
    "2003-2012 average."
)


def parse_smoke(raw: bytes) -> tuple[list[Observation], Guidance]:
    wb = _open(raw)
    g = read_guidance(wb)
    _check_guidance(g, "Indicator 1.2.1: Fire smoke", {"Mortality": "Deaths"}, "1.2.1")
    rows = read_table(wb, SMOKE_SHEET, SMOKE_COLUMNS, "1.2.1")
    return [Observation(entity="WLD", period=f"{year:04d}", value=float(v)) for year, v in rows], g


def _smoke(files: dict[str, InputFile]) -> Result:
    obs, g = parse_smoke(files[SMOKE.key].path.read_bytes())
    return Result(
        observations=obs,
        vintage=VINTAGE,
        year=str(g.report_year),
        date_published=g.published.isoformat(),
        steps=[
            f"Read the Mortality column (deaths due to fire-originated PM2.5) of the sheet {SMOKE_SHEET!r} of the "
            f"indicator 1.2.1 (fire smoke, mortality) workbook of the Lancet Countdown {g.report_year} report, one "
            f"value a year {obs[0].period}–{obs[-1].period}. Values are published as in the workbook (modelled "
            "deaths, not whole numbers).",
        ],
    )


# --- 1.2.2 share of land in extreme drought -----------------------------------------------------------------------

DROUGHT_SHEET = "2025 Report Data_Global"
DROUGHT_COLUMNS = ("Year", "Months", "Percent_in_drought")
# (dimension value = the sheet's Months key, label, column header of the same share in the guidance table)
DURATIONS = (
    ("1m", "At least 1 month", "perc_area_1m"),
    ("3m", "At least 3 months", "perc_area_3m"),
    ("6m", "At least 6 months", "perc_area_6m"),
)
LEFT_OUT_DURATION = "12m"
"""In the Global sheet but NA in 15 of its 74 years (and in no column of the guidance table): not published."""
DROUGHT_URL = (
    "https://lancetcountdown.org/wp-content/uploads/2025/10/"
    "Indicator-1.2.2_Data-Download_2025-Lancet-Countdown-Report.xlsx"
)
DROUGHT_KEY_FINDING = (
    "The percentage of the global land area affected by at least one month of extreme drought reached a "
    "record-breaking 60.7% in 2024, 299% above the 1951-1960 average"
)


def parse_drought(raw: bytes) -> tuple[list[Observation], Guidance]:
    wb = _open(raw)
    g = read_guidance(wb)
    _check_guidance(g, "1.2.2: drought", {guide: "%" for *_, guide in DURATIONS}, "1.2.2")
    if DROUGHT_SHEET not in wb.sheetnames:
        raise LancetFormatError(f"1.2.2: no sheet {DROUGHT_SHEET!r} in {wb.sheetnames}")
    rows = list(wb[DROUGHT_SHEET].iter_rows(values_only=True))
    if not rows or tuple(rows[0]) != DROUGHT_COLUMNS:
        raise LancetFormatError(f"1.2.2: {DROUGHT_SHEET!r} columns {rows[0] if rows else None} != {DROUGHT_COLUMNS}")
    keys = [k for k, *_ in DURATIONS]
    months_of: dict[int, list[str]] = {}
    shares: dict[int, dict[str, Decimal]] = {}
    for year, months, value in rows[1:]:
        if not isinstance(year, int) or isinstance(year, bool):
            raise LancetFormatError(f"1.2.2: Year {year!r} is not a whole number")
        months_of.setdefault(year, []).append(months)
        if months == LEFT_OUT_DURATION:
            continue
        if months not in keys:
            raise LancetFormatError(f"1.2.2: {year} has Months {months!r}, not one of {[*keys, LEFT_OUT_DURATION]}")
        if not isinstance(value, int | float) or isinstance(value, bool):
            raise LancetFormatError(f"1.2.2: {year} {months} is {value!r}, not a number")
        shares.setdefault(year, {})[months] = dec(value)
    years = list(months_of)
    if years != sorted(years):
        raise LancetFormatError("1.2.2: years are not increasing")
    for year in years:
        if sorted(months_of[year]) != sorted([*keys, LEFT_OUT_DURATION]):
            raise LancetFormatError(f"1.2.2: {year} has Months {months_of[year]}, not each of 1m, 3m, 6m, 12m once")
        s = shares[year]
        if not s["1m"] >= s["3m"] >= s["6m"]:
            raise LancetFormatError(f"1.2.2: {year} shares do not fall with duration: {s}")
    obs = [
        Observation(entity="WLD", period=f"{year:04d}", value=float(shares[year][k]), dims={"months": k})
        for k in keys
        for year in years
    ]
    return obs, g


def _drought(files: dict[str, InputFile]) -> Result:
    obs, g = parse_drought(files[DROUGHT.key].path.read_bytes())
    years = sorted({o.period for o in obs})
    return Result(
        observations=obs,
        vintage=VINTAGE,
        year=str(g.report_year),
        date_published=g.published.isoformat(),
        steps=[
            f"Read the sheet {DROUGHT_SHEET!r} of the indicator 1.2.2 (drought) workbook of the Lancet Countdown "
            f"{g.report_year} report, {years[0]}–{years[-1]}: the percent of the world's land area with at least 1, "
            "3 and 6 months of extreme drought in each year (Percent_in_drought by Months), published as in the "
            "workbook.",
            "Checked that every year has each duration once and that the share falls (or stays level) as the "
            "duration lengthens. The 12-month share is not published: it reads NA in 15 of the 74 years.",
        ],
    )


# --- 1.2.3 extreme precipitation events ---------------------------------------------------------------------------

RAIN_SHEET = "2025 Report Data_Event_Global"
RAIN_COLUMNS = ("Year", "Avg count extreme events")
RAIN_PERIODS = ((1961, 1990), (1995, 2024))
"""The baseline and recent 30-year periods the guidance sheet says the event counts cover; 1991–1994 have no row."""
RAIN_URL = (
    "https://lancetcountdown.org/wp-content/uploads/2025/10/"
    "Indicator-1.2.3_Data-Download_2025-Lancet-Countdown-Report-1.xlsx"
)
RAIN_QUOTE = "in 2024 the average annual number of >99th-percentile events per 79 km² reached a record high of five"
"""The second half of the workbook's key finding; the first half is about the land-share sheet, not published."""


def parse_rain(raw: bytes) -> tuple[list[Observation], Guidance]:
    wb = _open(raw)
    g = read_guidance(wb)
    _check_guidance(g, "1.2.3 Extreme Precipitation", {"Avg count extreme events": "Count"}, "1.2.3")
    rows = read_table(wb, RAIN_SHEET, RAIN_COLUMNS, "1.2.3")
    expected = [y for first, last in RAIN_PERIODS for y in range(first, last + 1)]
    if [r[0] for r in rows] != expected:
        raise LancetFormatError(f"1.2.3: {RAIN_SHEET!r} years are not exactly {RAIN_PERIODS}")
    return [Observation(entity="WLD", period=f"{year:04d}", value=float(v)) for year, v in rows], g


def _rain(files: dict[str, InputFile]) -> Result:
    obs, g = parse_rain(files[RAIN.key].path.read_bytes())
    return Result(
        observations=obs,
        vintage=VINTAGE,
        year=str(g.report_year),
        date_published=g.published.isoformat(),
        steps=[
            f"Read the column 'Avg count extreme events' of the sheet {RAIN_SHEET!r} of the indicator 1.2.3 "
            f"(extreme precipitation) workbook of the Lancet Countdown {g.report_year} report: the land-area-weighted "
            "average count of daily precipitation events above the 1961–1990 99th percentile, for each year of "
            "1961–1990 and 1995–2024, published as in the workbook.",
            "The workbook has no values for 1991–1994, and none are published. The sheet of land shares with more or "
            "fewer events is not published (irregular nine- or ten-year windows, values stored as fractions under a "
            "'%' heading).",
        ],
    )


# --- 3.2.1 deaths from outdoor air pollution by sector and fuel ---------------------------------------------------

AIR_SHEET = "Global"
_FUEL_HEADS = ("Coal", "Liquids and Gas", "Biomass", "Other")
AIR_HEADER = (
    ("Year", "Sector", "Population", "Deaths", None, None, None, "Mortality Rate", None, None, None, "PM2.5")
    + (None,) * 3,
    (None, None, None) + _FUEL_HEADS * 3,
)
"""The Global sheet's two header rows: the fuel names sit in row 2 under the merged Deaths, Mortality Rate and PM2.5
headings. Only the four Deaths columns are read."""
# (dimension value, the Global sheet's Sector label)
AIR_SECTORS = (
    ("power-plants", "Power plants"),
    ("industry", "Industry"),
    ("transport", "Transport"),
    ("households", "Households"),
    ("waste", "Waste"),
    ("agriculture", "Agriculture"),
    ("other", "Other"),
)
# (dimension value, the Global sheet's fuel label, column index of its deaths)
AIR_FUELS = (
    ("coal", "Coal", 3),
    ("liquids-gas", "Liquids and Gas", 4),
    ("biomass", "Biomass", 5),
    ("other", "Other", 6),
)
FOSSIL_FUELS = ("coal", "liquids-gas")
# (dimension value, label) of the fossil-fuel deaths
FOSSIL_DIMS = (
    ("total", "Coal, liquids and gas"),
    ("coal", "Coal"),
    ("liquids-gas", "Liquids and Gas"),
)
AIR_URL = (
    "https://lancetcountdown.org/wp-content/uploads/2025/10/"
    "Indicator-3.2.1_Data-Download_2025-Lancet-Countdown-Report-1.xlsx"
)
AIR_KEY_FINDING = (
    "Deaths attributable to ambient PM2.5 from fossil fuel combustion decreased by 5.8 % from 2.68 million in 2010 "
    "to 2.52 million in 2022."
)
AIR_BASIS = (
    "Modelled premature deaths: emissions by sector and fuel in IIASA's GAINS model, atmospheric chemistry and "
    "dispersion at about 7–10 km (meteorology fixed at 2015), population from UN World Population Prospects 2017 and "
    "exposure–response functions from a meta-analysis of cohort studies. 'Other' is other human-caused sources, as the "
    "producer puts it. Energy use comes from IEA World Energy Balances 2024 as a model input; no IEA values are "
    "published. The producer notes the model and the shape of the exposure–response relationship are uncertain, and "
    "that the grid may miss strong local pollution."
)


def read_air(raw: bytes) -> tuple[list[int], dict[tuple[int, str], dict[str, int]], Guidance]:
    """Years, and the deaths of each (year, sector dimension value) by fuel dimension value."""
    wb = _open(raw)
    g = read_guidance(wb)
    _check_guidance(g, "3.2.1: Mortality from ambient air pollution by sector and fuel", {"Deaths": "deaths"}, "3.2.1")
    if AIR_SHEET not in wb.sheetnames:
        raise LancetFormatError(f"3.2.1: no sheet {AIR_SHEET!r} in {wb.sheetnames}")
    rows = list(wb[AIR_SHEET].iter_rows(values_only=True))
    if tuple(tuple(r) for r in rows[:2]) != AIR_HEADER:
        raise LancetFormatError(f"3.2.1: {AIR_SHEET!r} header rows {rows[:2]} != expected {AIR_HEADER}")
    sector_of = {label: s for s, label in AIR_SECTORS}
    cells: dict[tuple[int, str], dict[str, int]] = {}
    years: list[int] = []
    for r in rows[2:]:
        year, label = r[0], r[1]
        if not isinstance(year, int) or isinstance(year, bool):
            raise LancetFormatError(f"3.2.1: Year {year!r} is not a whole number")
        if label not in sector_of:
            raise LancetFormatError(f"3.2.1: {year} has Sector {label!r}, not one of {list(sector_of)}")
        if (year, sector_of[label]) in cells:
            raise LancetFormatError(f"3.2.1: {year} {label} appears twice")
        deaths = {}
        for fuel, fuel_label, col in AIR_FUELS:
            v = r[col]
            if not isinstance(v, int) or isinstance(v, bool):
                raise LancetFormatError(f"3.2.1: {year} {label} {fuel_label} deaths {v!r} is not a whole number")
            deaths[fuel] = v
        cells[(year, sector_of[label])] = deaths
        if not years or years[-1] != year:
            years.append(year)
    if years != sorted(set(years)):
        raise LancetFormatError(f"3.2.1: {AIR_SHEET!r} years are not grouped and increasing")
    for year in years:
        missing = [s for s, _ in AIR_SECTORS if (year, s) not in cells]
        if missing:
            raise LancetFormatError(f"3.2.1: {year} has no row for {missing}")
    return years, cells, g


def parse_deaths_by_sector_fuel(raw: bytes) -> tuple[list[Observation], Guidance]:
    years, cells, g = read_air(raw)
    obs = [
        Observation(
            entity="WLD",
            period=f"{year:04d}",
            value=float(cells[(year, sector)][fuel]),
            dims={"sector": sector, "fuel": fuel},
        )
        for sector, _ in AIR_SECTORS
        for fuel, *_ in AIR_FUELS
        for year in years
    ]
    return obs, g


def parse_fossil_deaths(raw: bytes) -> tuple[list[Observation], Guidance]:
    years, cells, g = read_air(raw)
    by: dict[str, list[Observation]] = {d: [] for d, _ in FOSSIL_DIMS}
    for year in years:
        per_fuel = {f: sum(cells[(year, s)][f] for s, _ in AIR_SECTORS) for f in FOSSIL_FUELS}
        per_fuel["total"] = sum(per_fuel[f] for f in FOSSIL_FUELS)
        for d, _ in FOSSIL_DIMS:
            by[d].append(Observation(entity="WLD", period=f"{year:04d}", value=float(per_fuel[d]), dims={"fuel": d}))
    return [o for d, _ in FOSSIL_DIMS for o in by[d]], g


def _air_step(g: Guidance, years: list[str]) -> str:
    return (
        f"Read the sheet {AIR_SHEET!r} of the indicator 3.2.1 (mortality from ambient air pollution by sector and "
        f"fuel) workbook of the Lancet Countdown {g.report_year} report, {years[0]}–{years[-1]}: for each year and "
        "each of the seven emitting sectors (power plants, industry, transport, households, waste, agriculture, "
        "other), the deaths attributable to PM2.5 from coal, liquids and gas, biomass and other sources, in deaths as "
        "the guidance sheet states."
    )


def _deaths_by_sector_fuel(files: dict[str, InputFile]) -> Result:
    obs, g = parse_deaths_by_sector_fuel(files[AIR.key].path.read_bytes())
    years = sorted({o.period for o in obs if o.period})
    return Result(
        observations=obs,
        vintage=VINTAGE,
        year=str(g.report_year),
        date_published=g.published.isoformat(),
        steps=[
            _air_step(g, years),
            "Published every sector and fuel as in the workbook, zeros included (the guidance sheet says a 0 is a "
            "value). Death rates, PM2.5 concentrations and the regional sheets are not published.",
        ],
    )


def _fossil_deaths(files: dict[str, InputFile]) -> Result:
    obs, g = parse_fossil_deaths(files[AIR.key].path.read_bytes())
    years = sorted({o.period for o in obs if o.period})
    return Result(
        observations=obs,
        vintage=VINTAGE,
        year=str(g.report_year),
        date_published=g.published.isoformat(),
        steps=[
            _air_step(g, years),
            "Added the coal deaths of the seven sectors, and separately their liquids-and-gas deaths, for each year "
            "(whole deaths, exact), and added the two for deaths from fossil fuels, as the report's key finding "
            "does. Biomass and other sources are left out.",
        ],
        changes="deaths from coal and from liquids and gas added over the seven emitting sectors, and the two "
        "fuels added together.",
    )


# --- transforms ---------------------------------------------------------------------------------------------------


def transforms(paths: Paths) -> list[Transform]:
    here = Path(__file__)
    return [
        Transform(
            spec=Spec(
                id="heat.lancet-2025.deaths-global",
                title="Heat-related deaths, world",
                description="Deaths attributable to heat each year worldwide, 1990–2021, as estimated by the Lancet "
                "Countdown (indicator 1.1.5). The numbers are modelled from temperature–mortality relationships, "
                "not counted.",
                kind="series",
                unit=Unit(code="deaths", label="deaths", short="deaths"),
                display=Display(decimals=0),
                scope=Scope(
                    geography="World",
                    basis="Modelled in three stages (Zhao et al. 2021): the association between temperature and "
                    "mortality estimated in each location, pooled by multilevel meta-regression, and predicted for "
                    "all countries. Inputs include ERA5 temperatures, UN population data, the World Mortality "
                    "Dataset and the Global Burden of Disease Study 2021.",
                ),
                geo_coverage="global-only",
                headline_entity="WLD",
            ),
            inputs=(MORTALITY,),
            run=_mortality,
            module_file=here,
            validation=Validation(min_rows=32, value_range=(1.0, 5_000_000.0)),
        ),
        Transform(
            spec=Spec(
                id="heat.lancet-2025.labour-hours-global",
                title="Potential work hours lost to heat, world",
                description="Working hours that heat stress could have cost each year worldwide, 1990–2024, in "
                "services, manufacturing, agriculture and construction, as estimated by the Lancet Countdown "
                "(indicator 1.1.3).",
                kind="series",
                unit=Unit(code="bn-hours", label="billion hours", short="bn hours"),
                display=Display(decimals=1),
                dimensions=(
                    Dimension(
                        id="sector",
                        label="Sector",
                        values=[DimensionValue(id=s, label=label) for s, label, *_ in SECTORS],
                    ),
                ),
                scope=Scope(
                    geography="World",
                    basis="Potential hours lost: wet bulb globe temperature from ERA5 linked to the metabolic rate "
                    "of typical work (services 200 W and manufacturing 300 W in shade, agriculture and construction "
                    "400 W in the sun), applied to employed people aged 15 and over in each sector. Formal "
                    "employment only. The producer notes that from 2020 grid cells shared between countries are "
                    "handled differently, which may show as a step between 2019 and 2020.",
                ),
                geo_coverage="global-only",
                headline_entity="WLD",
                headline_dims=(("sector", "total"),),
            ),
            inputs=(WORK_HOURS,),
            run=_work_hours,
            module_file=here,
            validation=Validation(min_rows=5 * 35, value_range=(0.0, 5000.0)),
            checks=(
                PublisherCheck(
                    source_id=SOURCE,
                    vintage=VINTAGE,
                    entity="WLD",
                    period="2024",
                    stated="640",
                    quote="A record-high 640 billion potential work hours were lost in 2024",
                    url="https://lancetcountdown.org/wp-content/uploads/2025/10/"
                    "Indicator-1.1.3_PWHL_Data-Download_2025-Lancet-Countdown-Report_v2-1.xlsx",
                    dims=(("sector", "total"),),
                ),
            ),
        ),
        Transform(
            spec=Spec(
                id="heat.lancet-2025.labour-hours-per-worker",
                title="Potential work hours lost to heat per worker, world and by country",
                description="Working hours that heat stress could have cost each worker on average each year, "
                "1990–2024, in the world and in each of 219 countries and territories, across services, manufacturing, "
                "agriculture and construction, as estimated by the Lancet Countdown (indicator 1.1.3). Heat of any "
                "cause: the hours are not attributed to climate change. Where the producer's sheet has no one "
                "employed in the four sectors, or no value, the year is published as missing.",
                kind="series",
                unit=Unit(code="hours", label="hours per worker", short="hours"),
                display=Display(decimals=1),
                scope=Scope(
                    geography="World, and countries and territories as the producer draws them (UN territories)",
                    basis="Potential hours lost: wet bulb globe temperature from ERA5 linked to the metabolic rate "
                    "of typical work (services 200 W and manufacturing 300 W in shade, agriculture and construction "
                    "400 W in the sun), applied to employed people aged 15 and over in each sector, and divided by the "
                    "people employed in the four sectors. Formal employment only. The producer notes that from 2020 "
                    "grid cells shared between countries are handled differently, which may show as a step between "
                    "2019 and 2020.",
                ),
                geo_coverage="mixed",
                headline_entity="WLD",
            ),
            inputs=(WORK_HOURS,),
            run=_work_per_worker,
            module_file=here,
            validation=Validation(min_rows=35 * 200, value_range=(0.0, 3000.0)),
        ),
        Transform(
            spec=Spec(
                id="heat.lancet-2025.heatwave-days-global",
                title="Heatwave days per person, observed and attributable to climate change, world",
                description="Heatwave days that people were exposed to on average each year, 2020–2024, as "
                "observed, as expected in a modelled climate without human-caused warming, and the difference "
                "attributable to climate change, from the Lancet Countdown (indicator 1.1.1).",
                kind="series",
                unit=Unit(code="days", label="days per person", short="days"),
                display=Display(decimals=1),
                dimensions=(
                    Dimension(
                        id="scenario",
                        label="Scenario",
                        values=[DimensionValue(id=s, label=label) for s, label, _ in SCENARIOS],
                    ),
                ),
                scope=Scope(
                    geography="World",
                    basis="A heatwave is at least two consecutive days when both minimum and maximum temperature "
                    "exceed the local 1986–2005 95th percentile. Observed days from ERA5; days expected without "
                    "human-caused warming from 24 paired CMIP6 climate model simulations.",
                ),
                geo_coverage="global-only",
                headline_entity="WLD",
                headline_dims=(("scenario", "observed"),),
            ),
            inputs=(HEATWAVE,),
            run=_heatwave,
            module_file=here,
            validation=Validation(min_rows=3 * 5, value_range=(0.0, 366.0)),
        ),
        Transform(
            spec=Spec(
                id="heat.lancet-2025.heatwave-days-countries",
                title="Heatwave days per person, observed and attributable to climate change, world and by country",
                description="Heatwave days that people were exposed to on average each year, 2020–2024, in the world "
                "and in each of 220 countries and territories: as observed, as expected in a modelled climate without "
                "human-caused warming, and the difference attributable to climate change, from the Lancet Countdown "
                "(indicator 1.1.1). Estimates for small countries are more uncertain, as the producer notes; where a "
                "country had fewer heatwave days than its modelled counterfactual, the attributable value is "
                "negative.",
                kind="series",
                unit=Unit(code="days", label="days per person", short="days"),
                display=Display(decimals=1),
                dimensions=(
                    Dimension(
                        id="scenario",
                        label="Scenario",
                        values=[DimensionValue(id=s, label=label) for s, label, _ in SCENARIOS],
                    ),
                ),
                scope=Scope(
                    geography="World, and countries and territories as the producer draws them (UN territories)",
                    basis="A heatwave is at least two consecutive days when both minimum and maximum temperature "
                    "exceed the local 1986–2005 95th percentile. Observed days from ERA5; days expected without "
                    "human-caused warming from 24 paired CMIP6 climate model simulations. Each country's value is its "
                    "people's average.",
                ),
                geo_coverage="mixed",
                headline_entity="WLD",
                headline_dims=(("scenario", "attributable"),),
            ),
            inputs=(HEATWAVE,),
            run=_heatwave_countries,
            module_file=here,
            validation=Validation(min_rows=3 * 5 * 200, value_range=(-366.0, 366.0)),
        ),
        Transform(
            spec=Spec(
                id="heat.lancet-2025.heatwave-days-vulnerable",
                title="Heatwave days per person for adults over 65 and infants, world",
                description="Heatwave days that adults over 65 and infants under one year old were exposed to on "
                "average each year worldwide, 1980–2024, from the Lancet Countdown (indicator 1.1.1). These are "
                "heatwaves of any cause as observed in ERA5; the days are not attributed to climate change.",
                kind="series",
                unit=Unit(code="days", label="days per person", short="days"),
                display=Display(decimals=1),
                dimensions=(
                    Dimension(
                        id="group",
                        label="Group",
                        values=[DimensionValue(id=s, label=label) for s, label, *_ in GROUPS],
                    ),
                ),
                scope=Scope(
                    geography="World",
                    basis="A heatwave is at least two consecutive days when both minimum and maximum temperature "
                    "exceed the local 1986–2005 95th percentile; the report's key finding compares each year with "
                    "the 1986–2005 average. Each value is the group's heatwave person-days over its population, "
                    "from ERA5 temperatures, hybrid gridded demographic data and WorldPop age structure. The "
                    "producer warns that the two population sources may make values before and after 2000 "
                    "inconsistent.",
                ),
                geo_coverage="global-only",
                headline_entity="WLD",
                headline_dims=(("group", "over-65"),),
            ),
            inputs=(VULNERABLE,),
            run=_vulnerable,
            module_file=here,
            validation=Validation(min_rows=2 * 45, value_range=(0.0, 366.0)),
        ),
        Transform(
            spec=Spec(
                id="smoke.lancet-2025.wildfire-deaths",
                title="Deaths from wildfire smoke, world",
                description="Deaths each year worldwide from breathing fine particles (PM2.5) in the smoke of "
                "fires, 2003–2024, as modelled by the Lancet Countdown (indicator 1.2.1). It counts all fire "
                "smoke: the share caused by climate change is not attributed.",
                kind="series",
                unit=Unit(code="deaths", label="deaths", short="deaths"),
                display=Display(decimals=0),
                scope=Scope(
                    geography="World",
                    basis="Modelled, not counted: daily fire PM2.5 from the SILAM chemistry transport model driven "
                    "by MODIS fire radiative power (a fire hidden by cloud, or too small, is missed), and deaths from "
                    "the FUSION exposure–response model as the difference between results with and without fires "
                    "(code by A. Pozzer), with Global Burden of Disease baseline mortality and GPWv4 population.",
                ),
                geo_coverage="global-only",
                headline_entity="WLD",
            ),
            inputs=(SMOKE,),
            run=_smoke,
            module_file=here,
            validation=Validation(min_rows=22, value_range=(1.0, 1_000_000.0)),
            checks=(
                PublisherCheck(
                    source_id=SOURCE,
                    vintage=VINTAGE,
                    entity="WLD",
                    period="2024",
                    # Printed "154,000": written in deaths with its three significant digits, so the tolerance is
                    # half the last printed digit (±500), as for any stated value.
                    stated="154E3",
                    quote=SMOKE_KEY_FINDING,
                    url=SMOKE_URL,
                ),
            ),
        ),
        Transform(
            spec=Spec(
                id="drought.lancet-2025.land-share",
                title="Share of land in extreme drought, world",
                description="Percent of the world's land with at least one, three or six months of extreme "
                "meteorological drought in each year, 1951–2024, estimated by the Lancet Countdown (indicator "
                "1.2.2) from a drought index (SPEI-6, against 1986–2005). The index weighs precipitation against "
                "the evaporative demand of the air, so warmer air itself pushes it towards drought. It is not "
                "drought caused by climate change, and it does not measure water or crop shortages or the people "
                "affected.",
                kind="series",
                unit=Unit(code="percent", label="percent of land area", short="%"),
                display=Display(decimals=1),
                dimensions=(
                    Dimension(
                        id="months",
                        label="Months of extreme drought in the year",
                        values=[DimensionValue(id=k, label=label) for k, label, _ in DURATIONS],
                    ),
                ),
                scope=Scope(
                    geography="World land",
                    basis="Extreme meteorological drought by the 6-month Standardised Precipitation-"
                    "Evapotranspiration Index (SPEI, CSIC; the workbook does not name the version), an estimated "
                    "index from gridded station data, compared with its 1986–2005 baseline. Meteorological drought "
                    "only: the producer notes that it does not capture hydrological or agricultural drought, nor "
                    "the population living in, or depending on, drought-affected areas. The 12-month share is not "
                    "published (NA in 15 of 74 years).",
                ),
                geo_coverage="global-only",
                headline_entity="WLD",
                headline_dims=(("months", "1m"),),
            ),
            inputs=(DROUGHT,),
            run=_drought,
            module_file=here,
            validation=Validation(min_rows=3 * 74, value_range=(0.0, 100.0)),
            checks=(
                PublisherCheck(
                    source_id=SOURCE,
                    vintage=VINTAGE,
                    entity="WLD",
                    period="2024",
                    stated="60.7",
                    quote=DROUGHT_KEY_FINDING,
                    url=DROUGHT_URL,
                    dims=(("months", "1m"),),
                ),
            ),
        ),
        Transform(
            spec=Spec(
                id="rain.lancet-2025.extreme-events",
                title="Days of extreme rainfall per year over land, world",
                description="Average number of days a year on which daily precipitation at a place on land "
                "exceeded its local 1961–1990 99th percentile, for each year of 1961–1990 and 1995–2024, from the "
                "Lancet Countdown (indicator 1.2.3) using the ERA5-Land reanalysis. It counts days of heavy "
                "rainfall, not floods, and is not attributed to climate change. The workbook has no values for "
                "1991–1994.",
                kind="series",
                unit=Unit(
                    code="days",
                    label="days per year above the local 1961–1990 99th percentile of daily precipitation, average "
                    "over land",
                    short="days",
                ),
                display=Display(decimals=2),
                scope=Scope(
                    geography="World land",
                    basis="Daily precipitation from the ERA5-Land reanalysis (grid cells of about 79 square "
                    "kilometres), which does not directly assimilate rain-gauge data. An event is a day above the "
                    "cell's 1961–1990 99th percentile; counts are averaged over land, weighted by area. The producer "
                    "publishes only the baseline years 1961–1990 and the 30 recent years 1995–2024.",
                ),
                geo_coverage="global-only",
                headline_entity="WLD",
            ),
            inputs=(RAIN,),
            run=_rain,
            module_file=here,
            validation=Validation(min_rows=60, value_range=(0.0, 366.0)),
            checks=(
                PublisherCheck(
                    source_id=SOURCE,
                    vintage=VINTAGE,
                    entity="WLD",
                    period="2024",
                    stated="5",  # printed as the word "five"
                    quote=RAIN_QUOTE,
                    url=RAIN_URL,
                ),
            ),
        ),
        Transform(
            spec=Spec(
                id="air-pollution.lancet-2025.deaths-by-sector-fuel",
                title="Deaths from outdoor air pollution by sector and fuel, world",
                description="Deaths each year worldwide attributable to outdoor fine particles (PM2.5) from each "
                "emitting sector and fuel, 2007–2022, as modelled by the Lancet Countdown with IIASA's GAINS model "
                "(indicator 3.2.1). Deaths from burning coal, oil and gas are a harm of the burning itself, not of "
                "the warming.",
                kind="series",
                unit=Unit(code="deaths", label="deaths", short="deaths"),
                display=Display(decimals=0),
                dimensions=(
                    Dimension(
                        id="sector",
                        label="Emitting sector",
                        values=[DimensionValue(id=s, label=label) for s, label in AIR_SECTORS],
                    ),
                    Dimension(
                        id="fuel",
                        label="Fuel",
                        values=[DimensionValue(id=f, label=label) for f, label, _ in AIR_FUELS],
                    ),
                ),
                scope=Scope(
                    geography="World",
                    basis=AIR_BASIS,
                ),
                geo_coverage="global-only",
                headline_entity="WLD",
                headline_dims=(("sector", "power-plants"), ("fuel", "coal")),
            ),
            inputs=(AIR,),
            run=_deaths_by_sector_fuel,
            module_file=here,
            validation=Validation(min_rows=7 * 4 * 16, value_range=(0.0, 5_000_000.0)),
        ),
        Transform(
            spec=Spec(
                id="air-pollution.lancet-2025.fossil-deaths",
                title="Deaths from fossil-fuel air pollution, world",
                description="Deaths each year worldwide from breathing outdoor fine particles (PM2.5) from burning "
                "coal, oil and gas, 2007–2022: a harm of the burning itself, not of the warming. Calculated by "
                "adding the Lancet Countdown's deaths from coal and from liquids and gas over all seven emitting "
                "sectors (indicator 3.2.1), modelled with IIASA's GAINS model.",
                kind="derived",
                unit=Unit(code="deaths", label="deaths", short="deaths"),
                display=Display(decimals=0),
                dimensions=(
                    Dimension(
                        id="fuel",
                        label="Fuel",
                        values=[DimensionValue(id=f, label=label) for f, label in FOSSIL_DIMS],
                    ),
                ),
                scope=Scope(
                    geography="World",
                    basis=AIR_BASIS,
                ),
                geo_coverage="global-only",
                headline_entity="WLD",
                headline_dims=(("fuel", "total"),),
            ),
            inputs=(AIR,),
            run=_fossil_deaths,
            module_file=here,
            validation=Validation(min_rows=3 * 16, value_range=(0.0, 10_000_000.0)),
            # Printed "2.68 million" and "2.52 million": written in deaths with their three significant digits, so the
            # tolerance is half the last printed digit (±5,000).
            checks=tuple(
                PublisherCheck(
                    source_id=SOURCE,
                    vintage=VINTAGE,
                    entity="WLD",
                    period=period,
                    stated=stated,
                    quote=AIR_KEY_FINDING,
                    url=AIR_URL,
                    dims=(("fuel", "total"),),
                )
                for period, stated in (("2010", "2.68E6"), ("2022", "2.52E6"))
            ),
        ),
    ]
