"""Entities: the codes observations carry (models.Observation.entity), and how producers' codes map to them.

pipeline/geo/entities.csv is the explicit crosswalk, one row per entity:

  iso3          our entity code: ISO 3166-1 alpha-3 where one exists, else a code of ours (aggregates, stations), or
                a CCKP region code as CCKP writes it (IND.11510)
  name          English name (Natural Earth ADMIN for its polygons, GEOUNIT for its tiny-country points, CCKP's name
                for a region)
  kind          country | territory | aggregate | station | region
  ne_adm0_a3    Natural Earth ADM0_A3 of the 1:50m admin-0 polygon drawn for it; empty if none
  ne_iso_a3_eh  Natural Earth ISO_A3_EH of that polygon (or of its tiny-country point); empty if none
  member_of     aggregates it belongs to, ';'-separated (only memberships stated below; WLD is implied); for a region,
                the one country or territory it lies in

It is generated, never edited by hand: `uv run envdash geo build` reads the current natural-earth snapshots
(admin-0-countries-50m and admin-0-tiny-countries-50m, public domain) and the wb-cckp snapshot subnational-names (CC BY
4.0 with the World Bank Dataset Terms), and writes the CSV and its provenance sidecar (entities.provenance.json: the
snapshot sha256s it was built from).

Rules, all deterministic:
- Every 1:50m admin-0 polygon is one row. Its code is ISO_A3_EH, because Natural Earth sets ISO_A3 to -99 for France
  and Norway; where ISO_A3_EH is -99 (Kosovo, Somaliland, Northern Cyprus, Siachen Glacier) it is ADM0_A3. Where
  several polygons share one ISO_A3_EH (Australia, Indian Ocean Territories and Ashmore and Cartier Islands are all
  AUS), the polygon whose ADM0_A3 equals that code keeps it and the others take their ADM0_A3. A code still shared
  after that stops the build.
- kind: country when Natural Earth's ADMIN equals its SOVEREIGNT and its TYPE is Sovereign country, Country,
  Sovereignty or Disputed; otherwise territory. This is Natural Earth's default de facto point of view (so, for
  example, Palestine and Western Sahara, TYPE Indeterminate, are territories; Kosovo and Somaliland are countries),
  which the map states.
- A tiny-country point whose ISO_A3_EH (not -99) belongs to no polygon row is a row too (San Marino, Monaco, the
  Vatican, Gibraltar, Guadeloupe, Martinique, Tokelau, the Caribbean Netherlands), with ne_adm0_a3 empty.
- EXTRA_TERRITORIES, AGGREGATES and STATIONS below are added as declared. UN M49 regions are not entities here: M49
  is registered as display-only terms.
- Regions are the states, provinces and other first-level units the World Bank's Climate Change Knowledge Portal
  (CCKP) publishes values for, read from its names workbook (code, name): one row per CCKP region code, kind region,
  member_of the country or territory its prefix names (KSV, CCKP's code for Kosovo, is KOS through
  SOURCE_SCHEMES["wb-cckp"]). Left out, by declared rules that stop the build when the file no longer fits them: the
  DISPUTED_UNITS units coded DIS.* (areas the World Bank assigns to no country, named only by their code), and the
  three units CCKP lists under China for Hong Kong, Macao and Taiwan in some of its files (CHINA_LISTED), which are
  entities of their own here and in CCKP's other files. A region named only by a placeholder (PLACEHOLDER_NAMES, with
  their exact counts) or by its own code (NAMED_BY_CODE) is named "Unnamed region of {country} ({code})"; a name is
  never borrowed from another file. Region names must be unique within their country.
- Regions are looked up by code only. EntityTable.by indexes the name and Natural Earth columns over the other rows,
  so a region that shares a country's name (Georgia, the US state; Luxembourg, a district of Luxembourg) never makes a
  name ambiguous.

`resolve(code_or_name, scheme)` maps a producer's code or name to our code and raises UnknownEntity for anything not
in the table, never guessing (no fuzzy names, no prefix matches, no substitute for a missing entity). Schemes are the
table's own columns (iso3, ne_adm0_a3, ne_iso_a3_eh, name) or a source id from SOURCE_SCHEMES, whose explicit alias
table is tried before its base column.
"""

from __future__ import annotations

import csv
import io
import re
import struct
import zipfile
from collections import Counter
from collections.abc import Mapping
from dataclasses import dataclass, field
from functools import cache
from pathlib import Path
from types import MappingProxyType
from typing import Literal

import openpyxl

from envdash import canonical, snapshots
from envdash.models import ENTITY
from envdash.paths import Paths

COLUMNS = ("iso3", "name", "kind", "ne_adm0_a3", "ne_iso_a3_eh", "member_of")
EntityKind = Literal["country", "territory", "aggregate", "station", "region"]
NE_SOURCE = "natural-earth"
NE_COUNTRIES = "admin-0-countries-50m"
NE_TINY = "admin-0-tiny-countries-50m"
CCKP_SOURCE = "wb-cckp"
CCKP_NAMES = "subnational-names"
COUNTRY_TYPES = frozenset({"Sovereign country", "Country", "Sovereignty", "Disputed"})

# The 27 member states listed on https://european-union.europa.eu/principles-countries-history/eu-countries_en
# (read 2026-10-04: "Pages (27)", Austria to Sweden), by ISO 3166-1 alpha-3.
EU27_MEMBERS = frozenset(
    {
        "AUT", "BEL", "BGR", "HRV", "CYP", "CZE", "DNK", "EST", "FIN", "FRA", "DEU", "GRC", "HUN", "IRL",
        "ITA", "LVA", "LTU", "LUX", "MLT", "NLD", "POL", "PRT", "ROU", "SVK", "SVN", "ESP", "SWE",
    }
)  # fmt: skip

# Territories with an ISO 3166-1 code that a registered source reports separately but Natural Earth's 1:50m layers
# do not draw separately. No polygon: a map shows them as "no shape at this scale".
EXTRA_TERRITORIES: dict[str, tuple[str, str]] = {
    # gcp-fossil-co2-2025 mtco2-flat has "Christmas Island","CXR"; Natural Earth draws it inside Indian Ocean
    # Territories (ADM0_A3 IOA, ISO_A3_EH AUS) together with the Cocos (Keeling) Islands.
    "CXR": ("Christmas Island", "gcp-fossil-co2-2025"),
    # Overseas regions and dependencies that producers report separately but Natural Earth draws inside another
    # country's polygon (FRA, NOR) or omits. Reported separately by: Ember, EIA, EDGAR, FAOSTAT, FRA, CCKP.
    "GUF": ("French Guiana", "ember-yearly"),
    "REU": ("Réunion", "ember-yearly"),
    "MYT": ("Mayotte", "faostat"),
    "SJM": ("Svalbard and Jan Mayen", "faostat"),
    "BVT": ("Bouvet Island", "wb-cckp"),
    "CCK": ("Cocos (Keeling) Islands", "wb-cckp"),
    "UMI": ("United States Minor Outlying Islands", "wb-cckp"),
    # The World Bank's WDI reports the Channel Islands (Jersey and Guernsey) as one economy, code CHI.
    "CHI": ("Channel Islands", "wb-wdi"),
}

# Aggregates and producer pseudo-entities: no polygon. Codes are ours (ENTITY: [A-Z0-9_]{2,12}).
AGGREGATES: dict[str, str] = {
    "WLD": "World",
    "EU27": "European Union (27 member states since 1 February 2020)",
    "INTL_AIR": "International aviation",
    "INTL_SEA": "International shipping",
    "INTL_BUNKERS": "International transport (aviation and shipping bunker fuels)",
    # Historical pseudo-entities with their own rows in gcp-fossil-co2-2025 and jones-2025-national-contributions.
    "KWT_OILFIRES": "Kuwaiti oil fires (1991)",
    "PAC_ISLANDS": "Pacific Islands Trust Territory (Palau), as reported before Palau's independence",
    "RYUKYU": "Ryukyu Islands, as reported while under US administration",
    "NH": "Northern Hemisphere",
    "SH": "Southern Hemisphere",
    # Serbia and Montenegro reported as one entity in their historical years: EDGAR's and EIA's row SCG and FAOSTAT
    # area 186.
    "SRB_MNE": "Serbia and Montenegro (as reported together)",
    # GWIS's own continental totals of burned area. They are GWIS's published aggregates (its grouping of countries
    # into continents), not M49 membership lists, which stay unregistered.
    "UN_AFR": "Africa (as grouped by GWIS)",
    "UN_AME": "Americas (as grouped by GWIS)",
    "UN_ASI": "Asia (as grouped by GWIS)",
    "UN_EUR": "Europe (as grouped by GWIS)",
    "UN_OCE": "Oceania (as grouped by GWIS)",
    # EFFIS's total for the countries in the EU Civil Protection Mechanism, as EFFIS lists them.
    "UCPM": "Countries in the EU Civil Protection Mechanism (as listed by EFFIS)",
    # UNEP's Adaptation Gap Report: the developing countries whose adaptation costs and finance needs it estimates,
    # and the recipients of the international public adaptation finance it tracks (non-Annex I Parties to the
    # UNFCCC, per its online annex 4.B). UNEP publishes no member list, so no memberships are declared.
    "UNEP_AGR_DEV": "Developing countries (as grouped by UNEP's Adaptation Gap Report)",
    # Studies that cost something "in Europe" without a list of countries: Arnaiz del Pozo & Cloete (2022) price their
    # ammonia plants at "European energy prices" on a Western Europe cost basis, and Sacchi et al. (2023) cost
    # synthetic jet fuel with European electricity for flights from the EU-27 and EFTA. Each indicator's scope says
    # which Europe it is; this is not a membership list.
    "EUR_STUDY": "Europe (as the study cited defines it)",
    # desnz-subnational-electricity reports Great Britain (code K03000001: England, Scotland and Wales, without
    # Northern Ireland) and the London region (E12000007, the 33 London local authorities), not the United Kingdom.
    "GBR_GB": "Great Britain (England, Scotland and Wales)",
    "GBR_LONDON": "London (the London region of England)",
}

STATIONS: dict[str, str] = {
    "MLO": "Mauna Loa Observatory, Hawaii",
    "LAWDOME": "Law Dome ice cores, Antarctica",
    "ALOHA": "Station ALOHA, North Pacific north of Oahu, Hawaii",
    # bereiter-2015-co2: a composite of several Antarctic cores (Law Dome, Dome C, WAIS Divide, Siple Dome, Talos
    # Dome, EDML, Vostok), so not one site.
    "ANT_ICECORES": "Antarctic ice cores (composite of several sites)",
    # gardarsdottir-2019: the CEMCAP project's modelled reference plant, the one site its costs describe. It is not a
    # real plant, and not a country or region average.
    "CEMCAP_REF": "CEMCAP reference cement plant (a modelled European plant, 3,000 tonnes of clinker a day)",
}


@dataclass(frozen=True)
class SourceScheme:
    """How one producer names entities: an explicit alias table, then one column of the entity table."""

    base: Literal["iso3", "name"]
    aliases: dict[str, str] = field(default_factory=dict)
    note: str = ""


# Alias tables, declared from the producers' own files (each entry seen in the snapshot named in the note).
SOURCE_SCHEMES: dict[str, SourceScheme] = {
    "gcp-fossil-co2-2025": SourceScheme(
        base="iso3",
        aliases={
            "KSV": "KOS",
            "XIA": "INTL_AIR",
            "XIS": "INTL_SEA",
            # Rows with an empty code column, by their Country value.
            "Kuwaiti Oil Fires": "KWT_OILFIRES",
            "Pacific Islands (Palau)": "PAC_ISLANDS",
            "Ryukyu Islands": "RYUKYU",
        },
        note="mtco2-flat column 'ISO 3166-1 alpha-3' (Country for rows without a code), sha256 20650c19b394…",
    ),
    "climate-trace": SourceScheme(
        base="iso3",
        # Climate TRACE's own codes for areas without an ISO 3166-1 code (its data guide: "Kosovo has been assigned
        # the ISO3 code 'XKX'"; "ZNC name changed to 'Turkish Republic of Northern Cyprus'"). UNK ("unknown"
        # countries, mostly non-broadcasting vessels whose port is not known) has no entity: the transform counts it
        # in the world total only (climate_trace.WORLD_ONLY), so it raises if met anywhere else.
        aliases={"XKX": "KOS", "ZNC": "CYN"},
        note="V5.10.0 country files, column iso3_country (252 codes), e.g. electricity-generation co2e_100yr",
    ),
    "gfw-tree-cover-loss": SourceScheme(
        base="iso3",
        # GADM's own codes for areas without an ISO 3166-1 code. XAD, Z01, Z06 and Z07 have no entity here: the
        # transform leaves them out by name (gfw_tree_cover_loss.LEFT_OUT), so they raise if met anywhere else.
        aliases={"XKO": "KOS", "ZNC": "CYN"},
        note="gadm__tcl__iso_change v20260424 column iso, loss-country-annual sha256 d2356e4f4d7a…",
    ),
    "jones-2025-national-contributions": SourceScheme(
        base="iso3",
        aliases={
            "KSV": "KOS",
            "GLOBAL": "WLD",
            "XKW": "KWT_OILFIRES",
            "XPC": "PAC_ISLANDS",
            "XRY": "RYUKYU",
        },
        # ANNEXI, ANNEXII, BASIC, EIT, LDC, LMDC, NONANNEX and OECD are the producer's own country groups: not
        # entities here, so they raise.
        note="emissions-annual-1830 column ISO3, sha256 c718f698de72…",
    ),
    "wb-cckp": SourceScheme(
        base="iso3",
        # The same two codes as envdash/transforms/country/cckp.py ALIASES (tests/test_geo.py checks they agree): KSV,
        # CCKP's code for Kosovo, is the prefix of Kosovo's regions too (KSV.750081 …); GLOBAL is the key of its
        # /global files.
        aliases={"KSV": "KOS", "GLOBAL": "WLD"},
        note="all_countries and all_countries_subnationals files, keys; subnational-names column code",
    ),
    "kulp-strauss-2019": SourceScheme(
        base="name",
        # Supplementary Data 1's Place names that differ from the crosswalk's. "World" and "SIDS" (small island
        # developing states, the authors' own group) are not countries: the transform reads World on its own and
        # leaves SIDS out by name (kulp_strauss.NOT_PLACES), so they raise if met anywhere else.
        aliases={
            "Bahamas": "BHS",
            "Cape Verde": "CPV",
            "Côte d'Ivoire": "CIV",
            "Hong Kong": "HKG",
            "Macao": "MAC",
            "Micronesia": "FSM",
            "Reunion": "REU",
            "Sao Tome and Principe": "STP",
            "Tanzania": "TZA",
            "United States": "USA",
        },
        note="Supplementary Data 1 column Place (134 places besides World and SIDS), sha256 37a6ab2fc25e…",
    ),
    "wb-gad-v1": SourceScheme(
        base="iso3",
        # World Bank Official Boundaries V1 keys its admin-0 and admin-1 records by ISO_A3, the code CCKP uses for its
        # countries and as the prefix of its region codes. Kosovo is KSV there (its WB_A3 is KOS, ours too). WB_A3 is
        # not used: it is blank for Iraq and Haiti and puts Taiwan, Réunion and other areas under another code.
        aliases={"KSV": "KOS"},
        note="WB_GAD_ADM0.dbf column ISO_A3 (245 codes in 267 records), sha256 23f2adab0b6d…",
    ),
}


class UnknownEntity(KeyError):
    pass


# --- CCKP regions -------------------------------------------------------------------------------------------------

CCKP_COUNTRY = re.compile(r"[A-Z]{3}")
CCKP_REGION = re.compile(r"(?P<prefix>[A-Z]{3})\.\d{1,12}")
"""A CCKP region code: the country's three letters (CCKP's, not always ours), a dot, digits."""

DISPUTED_PREFIX = "DIS"
DISPUTED_UNITS = 32
"""CCKP's DIS.* units (DIS.990011 …): areas the World Bank assigns to no country, named only by their code. They are
left out: there is no country to place them in, and no name to show. In the names workbook of 2026-10-16 and in every
region file registered: 32."""

# Region names that are not names, exactly as CCKP writes them, with how many regions carry each (names workbook of
# 2026-10-16: 50 and 18, 68 regions in 51 countries and territories, for example AIA.9416 to AIA.9425 for Anguilla's
# ten units, BRA.37692 to BRA.37695 for Brazil's four ocean islands, GBR.25639503). A count that differs, or a new
# placeholder wording, stops the build until the names are re-read.
PLACEHOLDER_NAMES: dict[str, int] = {"Administrative unit not available": 50, "Name Unknown": 18}

# Regions whose name in CCKP's workbook is their own code: the single region of Hong Kong, Macao and Taiwan in the
# files that list one (tas, rx1day and cdd; Taiwan also in hi39 and txx). CCKP also names these three countries by
# their code (HKG, MAC, TWN).
NAMED_BY_CODE = frozenset({"HKG.33364910", "MAC.1491879", "TWN.147296925"})

# Units CCKP lists under China for Hong Kong, Macao and Taiwan, with their names in CCKP's workbook. CCKP's region files
# do not agree on where these places sit (read 2026-10-16): hd35 and tr26 list all three only under China; hi39 and txx
# also list Taiwan as TWN.147296925; cdd lists Hong Kong and Macao under China and all three under their own codes;
# rx1day lists all six, with different values (Taiwan, 2040–2059 SSP2-4.5 median: CHN.147295925 151.67 mm,
# TWN.147296925 123.49 mm); tas lists only the three under their own codes. Every file also has HKG, MAC and TWN as
# countries with their own values, and this table has them as entities of their own (Natural Earth's de facto view).
# So these three are left out, and each place keeps one row: its own. Any other unit of China named for one of these
# places, or one of these renamed, stops the build.
CHINA_LISTED: dict[str, str] = {
    "CHN.33364910": "Hong Kong, SAR",
    "CHN.1491879": "Macau, SAR",
    "CHN.147295925": "Taiwan Sheng",
}
CHINA_LISTED_WORDS = re.compile(r"hong ?kong|maca[uo]|taiwan", re.IGNORECASE)


def unnamed_region(country: str, code: str) -> str:
    return f"Unnamed region of {country} ({code})"


def read_cckp_names(xlsx: bytes) -> dict[str, str]:
    """CCKP code -> CCKP name, from the names workbook: one sheet, columns code, name and one value column (not read),
    one row per country and region."""
    wb = openpyxl.load_workbook(io.BytesIO(xlsx), read_only=True)
    try:
        if wb.sheetnames != ["all"]:
            raise ValueError(f"CCKP names workbook: sheets {wb.sheetnames}, expected ['all']")
        rows = list(wb["all"].iter_rows(values_only=True))
    finally:
        wb.close()
    if not rows or len(rows[0]) != 3 or tuple(rows[0][:2]) != ("code", "name"):
        raise ValueError(
            f"CCKP names workbook: header {rows[0] if rows else None}, expected code, name, a value column"
        )
    names: dict[str, str] = {}
    for i, row in enumerate(rows[1:], start=2):
        code, name = row[0], row[1]
        if not isinstance(code, str) or not (CCKP_COUNTRY.fullmatch(code) or CCKP_REGION.fullmatch(code)):
            raise ValueError(f"CCKP names workbook row {i}: code {code!r} is neither three letters nor a region code")
        if not isinstance(name, str) or not name or name != name.strip():
            raise ValueError(f"CCKP names workbook row {i}: {code} has name {name!r}")
        if code in names:
            raise ValueError(f"CCKP names workbook row {i}: {code} appears twice")
        names[code] = name
    return names


def cckp_regions(names: dict[str, str], places: list[Entity]) -> list[Entity]:
    """The region rows for CCKP's names (read_cckp_names), whose countries must be among `places` (by our code)."""
    countries = {c for c in names if CCKP_COUNTRY.fullmatch(c)}
    regions = {c: n for c, n in names.items() if CCKP_REGION.fullmatch(c)}
    prefix = {c: c.split(".", 1)[0] for c in regions}
    disputed = sorted(c for c in regions if prefix[c] == DISPUTED_PREFIX)
    if len(disputed) != DISPUTED_UNITS:
        raise ValueError(f"CCKP has {len(disputed)} {DISPUTED_PREFIX}.* units, not the {DISPUTED_UNITS} declared")
    orphans = sorted(c for c in regions if prefix[c] not in countries | {DISPUTED_PREFIX})
    if orphans:
        raise ValueError(f"CCKP regions whose prefix is not a country of the same file: {orphans[:10]}")
    china = {c: n for c, n in regions.items() if prefix[c] == "CHN" and CHINA_LISTED_WORDS.search(n)}
    if china != CHINA_LISTED:
        raise ValueError(
            f"CCKP lists {china} under China for Hong Kong, Macao or Taiwan, not the declared {CHINA_LISTED}; re-read "
            "the region files before changing envdash/geo.py CHINA_LISTED"
        )
    kept = {c: n for c, n in regions.items() if prefix[c] != DISPUTED_PREFIX and c not in CHINA_LISTED}
    found = Counter(n for n in kept.values() if n in PLACEHOLDER_NAMES)
    if dict(found) != PLACEHOLDER_NAMES:
        raise ValueError(f"CCKP's placeholder region names are {dict(found)}, not the declared {PLACEHOLDER_NAMES}")
    by_code = {c for c, n in kept.items() if n == c}
    if by_code != NAMED_BY_CODE:
        raise ValueError(f"CCKP regions named by their own code are {sorted(by_code)}, not {sorted(NAMED_BY_CODE)}")
    table = EntityTable(tuple(places))
    parents = {e.iso3: e for e in places}
    rows: list[Entity] = []
    for code, name in sorted(kept.items()):
        try:
            parent = parents[resolve(prefix[code], CCKP_SOURCE, entities=table)]
        except UnknownEntity as e:
            raise ValueError(f"CCKP region {code}: its country {prefix[code]!r} has no entity ({e})") from None
        if parent.kind not in ("country", "territory"):
            raise ValueError(f"CCKP region {code}: its parent {parent.iso3} is a {parent.kind}")
        unnamed = name in PLACEHOLDER_NAMES or code in NAMED_BY_CODE
        label = unnamed_region(parent.name, code) if unnamed else name
        rows.append(Entity(code, label, "region", "", "", (parent.iso3,)))
    seen: dict[tuple[str, str], str] = {}
    for e in rows:
        k = (e.member_of[0], e.name)
        if k in seen:
            raise ValueError(f"CCKP regions {seen[k]} and {e.iso3} of {k[0]} are both named {e.name!r}")
        seen[k] = e.iso3
    return rows


@dataclass(frozen=True)
class Entity:
    iso3: str
    name: str
    kind: EntityKind
    ne_adm0_a3: str = ""
    ne_iso_a3_eh: str = ""
    member_of: tuple[str, ...] = ()

    def row(self) -> list[str]:
        return [self.iso3, self.name, self.kind, self.ne_adm0_a3, self.ne_iso_a3_eh, ";".join(self.member_of)]


# --- Natural Earth shapefile attributes (dBase III .dbf) ----------------------------------------------------------


def read_dbf(data: bytes, encoding: str) -> list[dict[str, str]]:
    """Records of a dBase III table as strings (trailing spaces and NUL padding removed); deleted records skipped."""
    count, header_len, record_len = struct.unpack("<IHH", data[4:12])
    fields: list[tuple[str, int]] = []
    pos = 32
    while data[pos] != 0x0D:
        name = data[pos : pos + 11].split(b"\0", 1)[0].decode("ascii")
        fields.append((name, data[pos + 16]))
        pos += 32
    rows: list[dict[str, str]] = []
    for i in range(count):
        rec = data[header_len + i * record_len : header_len + (i + 1) * record_len]
        if rec[:1] == b"*":
            continue
        off, row = 1, {}
        for name, width in fields:
            row[name] = rec[off : off + width].decode(encoding).strip("\x00 ")
            off += width
        rows.append(row)
    return rows


def shapefile_records(zip_bytes: bytes) -> list[dict[str, str]]:
    with zipfile.ZipFile(io.BytesIO(zip_bytes)) as z:
        dbf = [n for n in z.namelist() if n.lower().endswith(".dbf")]
        if len(dbf) != 1:
            raise ValueError(f"expected one .dbf in the zip, found {dbf}")
        cpg = dbf[0][:-4] + ".cpg"
        encoding = z.read(cpg).decode("ascii").strip() if cpg in z.namelist() else "latin-1"
        return read_dbf(z.read(dbf[0]), "utf-8" if encoding.upper().replace("-", "") == "UTF8" else encoding)


def _kind(r: dict[str, str]) -> EntityKind:
    return "country" if r["ADMIN"] == r["SOVEREIGNT"] and r["TYPE"] in COUNTRY_TYPES else "territory"


def entities_from_natural_earth(countries_zip: bytes, tiny_zip: bytes) -> list[Entity]:
    polygons = shapefile_records(countries_zip)
    eh_owners: dict[str, list[str]] = {}
    for r in polygons:
        eh_owners.setdefault(r["ISO_A3_EH"], []).append(r["ADM0_A3"])
    rows: list[Entity] = []
    for r in polygons:
        eh = r["ISO_A3_EH"]
        shared = len(eh_owners[eh]) > 1
        code = r["ADM0_A3"] if eh == "-99" or (shared and r["ADM0_A3"] != eh) else eh
        rows.append(
            Entity(
                iso3=code,
                name=r["ADMIN"],
                kind=_kind(r),
                ne_adm0_a3=r["ADM0_A3"],
                ne_iso_a3_eh="" if eh == "-99" else eh,
                member_of=("EU27",) if code in EU27_MEMBERS and _kind(r) == "country" else (),
            )
        )
    taken = {e.iso3 for e in rows} | {e.ne_iso_a3_eh for e in rows}
    for r in shapefile_records(tiny_zip):
        eh = r["ISO_A3_EH"]
        if eh == "-99" or eh in taken:
            continue
        taken.add(eh)
        rows.append(Entity(iso3=eh, name=r["GEOUNIT"], kind=_kind(r), ne_adm0_a3="", ne_iso_a3_eh=eh))
    return rows


def all_entities(countries_zip: bytes, tiny_zip: bytes, cckp_names_xlsx: bytes) -> list[Entity]:
    rows = entities_from_natural_earth(countries_zip, tiny_zip)
    rows += [Entity(code, name, "territory") for code, (name, _) in EXTRA_TERRITORIES.items()]
    rows += [Entity(code, name, "aggregate") for code, name in AGGREGATES.items()]
    rows += [Entity(code, name, "station") for code, name in STATIONS.items()]
    rows += cckp_regions(read_cckp_names(cckp_names_xlsx), rows)
    codes = [e.iso3 for e in rows]
    dupes = sorted({c for c in codes if codes.count(c) > 1})
    if dupes:
        raise ValueError(f"entity codes defined more than once: {dupes}")
    bad = [c for c in codes if not ENTITY.fullmatch(c)]
    if bad:
        raise ValueError(f"entity codes not matching {ENTITY.pattern}: {bad}")
    members = {e.iso3 for e in rows if "EU27" in e.member_of}
    if members != EU27_MEMBERS:
        raise ValueError(f"EU27 members without a country row: {sorted(EU27_MEMBERS - members)}")
    return sorted(rows, key=lambda e: e.iso3)


def csv_bytes(rows: list[Entity]) -> bytes:
    out = io.StringIO(newline="")
    w = csv.writer(out, lineterminator="\n")
    w.writerow(COLUMNS)
    for e in rows:
        w.writerow(e.row())
    return out.getvalue().encode("utf-8")


# --- build and load -----------------------------------------------------------------------------------------------


def entities_path(paths: Paths) -> Path:
    return paths.repo / "pipeline" / "geo" / "entities.csv"


def _snapshot(paths: Paths, source: str, artifact: str) -> tuple[str, bytes]:
    sha = snapshots.read_current(paths).get(snapshots.key(source, artifact))
    if sha is None:
        raise FileNotFoundError(f"no snapshot of {source}/{artifact}; run envdash fetch -s {source}")
    data = snapshots.cache_path(paths, sha).read_bytes()
    if canonical.sha256_bytes(data) != sha:
        raise ValueError(f"cached bytes of {source}/{artifact} do not hash to {sha}")
    return sha, data


def build(paths: Paths) -> tuple[bytes, dict]:
    """The entities.csv bytes and provenance sidecar built from the current natural-earth snapshots and CCKP's
    names."""
    c_sha, countries = _snapshot(paths, NE_SOURCE, NE_COUNTRIES)
    t_sha, tiny = _snapshot(paths, NE_SOURCE, NE_TINY)
    n_sha, names = _snapshot(paths, CCKP_SOURCE, CCKP_NAMES)
    data = csv_bytes(all_entities(countries, tiny, names))
    side = {
        "file": "pipeline/geo/entities.csv",
        "sha256": canonical.sha256_bytes(data),
        "built_by": "uv run envdash geo build (envdash/geo.py)",
        "inputs": {
            f"{NE_SOURCE}/{NE_COUNTRIES}": c_sha,
            f"{NE_SOURCE}/{NE_TINY}": t_sha,
            f"{CCKP_SOURCE}/{CCKP_NAMES}": n_sha,
        },
    }
    return data, side


def write(paths: Paths) -> bool:
    data, side = build(paths)
    out = entities_path(paths)
    changed = canonical.write_if_changed(out, data)
    changed |= canonical.write_if_changed(out.with_name("entities.provenance.json"), canonical.dump_bytes(side))
    return changed


REGION_INDEXES = frozenset({"iso3"})


@dataclass(frozen=True)
class EntityTable:
    rows: tuple[Entity, ...]
    _indexes: dict[str, Mapping[str, tuple[Entity, ...]]] = field(
        default_factory=dict, init=False, repr=False, compare=False
    )

    def by(self, column: str) -> Mapping[str, tuple[Entity, ...]]:
        """Rows by the value of one column, built once per table (resolve runs once per producer row, and the table
        holds some 3,800 rows). Regions are indexed by their code (iso3) only, never by name or a Natural Earth code,
        so a region named like a country (Georgia, a US state) cannot make a name lookup ambiguous."""
        if column not in self._indexes:
            index: dict[str, list[Entity]] = {}
            for e in self.rows:
                if e.kind == "region" and column not in REGION_INDEXES:
                    continue
                v = getattr(e, column)
                if v:
                    index.setdefault(v, []).append(e)
            self._indexes[column] = MappingProxyType({k: tuple(v) for k, v in index.items()})
        return self._indexes[column]


def read_table(path: Path) -> EntityTable:
    with path.open(encoding="utf-8", newline="") as f:
        reader = csv.reader(f)
        header = next(reader)
        if tuple(header) != COLUMNS:
            raise ValueError(f"{path}: columns {header} != {list(COLUMNS)}")
        rows = [
            Entity(r[0], r[1], r[2], r[3], r[4], tuple(m for m in r[5].split(";") if m))  # type: ignore[arg-type]
            for r in reader
        ]
    return EntityTable(tuple(rows))


@cache
def _default_table() -> EntityTable:
    return read_table(entities_path(Paths.default()))


def table(path: Path | None = None) -> EntityTable:
    return read_table(path) if path else _default_table()


SCHEMES = ("iso3", "ne_adm0_a3", "ne_iso_a3_eh", "name")


def resolve(code_or_name: str, scheme: str, *, entities: EntityTable | None = None) -> str:
    """Our entity code for a producer's code or name under `scheme`; UnknownEntity if absent or ambiguous."""
    t = entities or table()
    if scheme in SOURCE_SCHEMES:
        s = SOURCE_SCHEMES[scheme]
        if code_or_name in s.aliases:
            return resolve(s.aliases[code_or_name], "iso3", entities=t)
        try:
            return resolve(code_or_name, s.base, entities=t)
        except UnknownEntity:
            raise UnknownEntity(
                f"{code_or_name!r} is not an entity under {scheme} (its aliases, then column {s.base}); declare it in "
                "envdash/geo.py SOURCE_SCHEMES if the producer's file has it"
            ) from None
    if scheme not in SCHEMES:
        raise ValueError(f"unknown scheme {scheme!r}: one of {list(SCHEMES)} or a source id in SOURCE_SCHEMES")
    found = t.by(scheme).get(code_or_name, ())
    if not found:
        raise UnknownEntity(f"{code_or_name!r} is not in pipeline/geo/entities.csv column {scheme}")
    if len(found) > 1:
        raise UnknownEntity(
            f"{code_or_name!r} in column {scheme} names {len(found)} entities {[e.iso3 for e in found]}; use iso3"
        )
    return found[0].iso3


def entity(code: str, *, entities: EntityTable | None = None) -> Entity:
    """The row for one of our codes."""
    t = entities or table()
    found = t.by("iso3").get(code)
    if not found:
        raise UnknownEntity(f"{code!r} is not in pipeline/geo/entities.csv")
    return found[0]
