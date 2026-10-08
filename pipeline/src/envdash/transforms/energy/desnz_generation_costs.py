"""UK DESNZ, Electricity Generation Costs 2025, Annex A: the levelised cost of electricity of new plants commissioning
in 2030, 2035, 2040, 2045 and 2050, by technology and load factor, split into its cost components, as DESNZ tabulates
it. A twin of lcoe.eia-aeo-2026.new-plants-us (another country, another producer, never blended with it).

Input: annex-a-additional-estimates-and-key-assumptions-2025.xlsx (registry entry desnz-generation-costs-2025,
artifact annex-a). Sheets read: "Additional Estimates 2030" to "Additional Estimates 2050". Each has a title naming
its commissioning year, the line "All costs in 2024 real GBP prices", a header row ("Levelised cost breakdown",
"Units", then one column per technology and load factor) and eleven rows in pounds per megawatt-hour: pre-development,
construction, fixed O&M, variable O&M, fuel, carbon, CO2 transport and storage, decommissioning and waste, steam
revenue, additional costs, and the total. The Home sheet's general notes say "All costs are in real 2024 GBP" and that
sub-components are rounded and may not sum to the total.

Reading rules:
- Column headers carry footnote numbers glued to the end ("Fixed Offshore Wind5", "Total13") and sometimes a trailing
  space; both are removed. "(FOAK)" marks a first-of-a-kind estimate (dimension "estimate"); "(50% hydrogen / natural
  gas fuel blend)" marks the 2030 hydrogen turbine, a different technology from the later 100% hydrogen ones. What is
  left must be one of the technology names below followed, for dispatchable plants, by its load factor ("5% Load
  Factor", or "5%" for the hydrogen engines in 2035-2045), or the build stops.
- A cell printed '<1' is published as a null value with that reason: DESNZ publishes no number for it.
- Load factors are part of the technology: unabated gas and hydrogen plants at 5%, 30% and 93%, gas with carbon
  capture (CCUS) at 5%, 30% and 88% (DESNZ: the lower baseload factor reflects the assumed availability of the CO2
  transport and storage network).

Checks inside the file, or the build stops: each sheet's title names its year and 2024 real pounds; every unit cell
is '£/MWh'; the eleven row labels are exactly these; each total lies within the sum of its rounded components plus or
minus half a pound per component ('<1' counted as 0 to 1), which catches a misread column.

Nothing is computed: no difference between technologies and no total without carbon costs. Carbon costs are a
published row, so the site can show them beside the total.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass
from decimal import Decimal
from functools import lru_cache
from pathlib import Path

import openpyxl

from envdash.models import Dimension, DimensionValue, Display, Observation, Scope, Unit
from envdash.paths import Paths
from envdash.transform import Input, InputFile, Result, Spec, Transform, Validation

SOURCE = "desnz-generation-costs-2025"
ANNEX = Input(SOURCE, "annex-a")
INDICATOR = "lcoe.desnz-2025.new-plants-uk"
VINTAGE = "Electricity Generation Costs 2025 (DESNZ; Annex A updated 18 March 2026)"
PUBLISHED = "2026-01-14"
ENTITY = "GBR"
YEARS = (2030, 2035, 2040, 2045, 2050)
SHEET = "Additional Estimates {year}"
TITLE = "Additional LCOE estimates - for projects commissioning in {year}"
PRICES = "All costs in 2024 real GBP prices"
HOME_NOTE = "All costs are in real 2024 GBP"
UNIT_CELL = "£/MWh"
LESS_THAN_ONE = "<1"
FOAK = "(FOAK)"
BLEND = "(50% hydrogen / natural gas fuel blend)"

# Row label (footnote digits removed) -> (component id, label). Order is the sheet's.
COMPONENTS: dict[str, tuple[str, str]] = {
    "Pre-development costs": ("pre-development", "Pre-development costs"),
    "Construction costs": ("construction", "Construction costs"),
    "Fixed O&M costs": ("fixed-om", "Fixed operation and maintenance"),
    "Variable O&M costs": ("variable-om", "Variable operation and maintenance"),
    "Fuel costs": ("fuel", "Fuel"),
    "Carbon costs": ("carbon", "Carbon costs (UK ETS price and Carbon Price Support)"),
    "CO2 Transport and Storage": ("co2-transport-storage", "CO2 transport and storage"),
    "Decommissioning and waste": ("decommissioning", "Decommissioning and waste"),
    "Steam Revenue": ("steam-revenue", "Steam revenue"),
    "Additional Costs": ("additional", "Additional costs"),
    "Total": ("total", "Total levelised cost"),
}
HEADER_FIRST = ("Levelised cost breakdown", "Units")

# Base name as printed -> (technology id stem, label, load factors in percent; () for plants without one).
GAS_LF = (5, 30, 93)
CCUS_LF = (5, 30, 88)
BASES: dict[str, tuple[str, str, tuple[int, ...]]] = {
    "Large Scale Solar": ("large-scale-solar", "Large-scale solar", ()),
    "Onshore Wind": ("onshore-wind", "Onshore wind", ()),
    "Fixed Offshore Wind": ("offshore-wind-fixed", "Fixed offshore wind", ()),
    "Floating Offshore Wind": ("offshore-wind-floating", "Floating offshore wind", ()),
    "Gas CCGT": ("gas-ccgt", "Gas combined-cycle (CCGT)", GAS_LF),
    "Gas OCGT 300MW": ("gas-ocgt-300mw", "Gas open-cycle turbine (OCGT), 300 MW", GAS_LF),
    "Gas OCGT 760MW": ("gas-ocgt-760mw", "Gas open-cycle turbine (OCGT), 760 MW", GAS_LF),
    "Gas Reciprocating Engine": ("gas-engine", "Gas reciprocating engine", GAS_LF),
    "CCHT 50% 900MW": (
        "hydrogen-ccht-blend",
        "Combined-cycle hydrogen turbine, 900 MW, 50% hydrogen and natural gas blend",
        GAS_LF,
    ),
    "CCHT 900MW": ("hydrogen-ccht", "Combined-cycle hydrogen turbine (CCHT), 900 MW", GAS_LF),
    "OCHT": ("hydrogen-ocht", "Open-cycle hydrogen turbine (OCHT)", GAS_LF),
    "Hydrogen Reciprocating Engines": ("hydrogen-engine", "Hydrogen reciprocating engine", GAS_LF),
    "Gas CCUS 900MW": ("gas-ccus", "Gas combined-cycle with carbon capture (CCUS), 900 MW", CCUS_LF),
    "Tidal Stream Energy": ("tidal-stream", "Tidal stream", ()),
    "Deep Granite Geothermal": ("deep-geothermal", "Deep granite geothermal", ()),
}
ESTIMATES = {
    "unmarked": "Not marked first of a kind (DESNZ: 'Nth of a Kind' unless marked FOAK)",
    "foak": "First of a kind (FOAK)",
}


def _technologies() -> dict[str, str]:
    out = {}
    for stem, label, lfs in BASES.values():
        if not lfs:
            out[stem] = label
        for lf in lfs:
            out[f"{stem}-{lf}pc-load"] = f"{label}, {lf}% load factor"
    return out


TECHNOLOGIES = _technologies()

_FOOTNOTE = re.compile(r"(?<=[A-Za-z%)])\d+(?:, ?\d+)*$")
_ROW_FOOTNOTE = re.compile(r"(?<=[a-z])\d+$")
_LOAD = re.compile(r"(?P<base>.+?)(?: (?P<lf>\d+)%(?: Load Factor)?)?")


class DesnzFormatError(ValueError):
    pass


@dataclass(frozen=True)
class Column:
    technology: str
    estimate: str
    header: str


def parse_header(header: str) -> Column:
    """A column header as printed -> technology id and estimate kind (see the module docstring)."""
    name = _FOOTNOTE.sub("", header.strip()).strip()
    foak = FOAK in name
    name = name.replace(FOAK, " ")
    blend = BLEND in name
    name = " ".join(name.replace(BLEND, " ").split())
    m = _LOAD.fullmatch(name)
    if m is None or m.group("base") not in BASES:
        raise DesnzFormatError(f"unknown technology column {header!r} (read as {name!r})")
    stem, _, lfs = BASES[m.group("base")]
    if blend != (stem == "hydrogen-ccht-blend"):
        raise DesnzFormatError(f"column {header!r}: the 50% hydrogen blend note does not match {m.group('base')!r}")
    lf = m.group("lf")
    if (lf is None) != (not lfs) or (lf is not None and int(lf) not in lfs):
        raise DesnzFormatError(f"column {header!r}: load factor {lf!r} is not one of {lfs} for {m.group('base')!r}")
    tid = stem if lf is None else f"{stem}-{lf}pc-load"
    return Column(technology=tid, estimate="foak" if foak else "unmarked", header=header)


@dataclass(frozen=True)
class Cell:
    value: Decimal | None
    """None where DESNZ prints '<1'."""


@dataclass(frozen=True)
class Sheet:
    year: int
    columns: list[Column]
    cells: dict[str, list[Cell]]
    """Component id -> one cell per column."""


def read_sheet(rows: list[tuple], year: int) -> Sheet:
    name = SHEET.format(year=year)
    rows = [r for r in rows if any(v is not None for v in r)]  # blank spacer rows carry nothing
    if rows[0][0] != TITLE.format(year=year) or rows[1][0] != PRICES:
        raise DesnzFormatError(f"{name}: title rows {rows[0][0]!r}, {rows[1][0]!r} are not this edition's")
    header = rows[2]
    if tuple(header[:2]) != HEADER_FIRST:
        raise DesnzFormatError(f"{name}: header starts {header[:2]!r}")
    width = len(header)
    while width > 2 and header[width - 1] is None:
        width -= 1
    if any(h is None for h in header[2:width]):
        raise DesnzFormatError(f"{name}: an empty column header inside {header!r}")
    columns = [parse_header(str(h)) for h in header[2:width]]
    keys = [(c.technology, c.estimate) for c in columns]
    if len(set(keys)) != len(keys):
        raise DesnzFormatError(f"{name}: two columns read as the same technology and estimate: {keys}")
    body = rows[3:]
    labels = [_ROW_FOOTNOTE.sub("", str(r[0])) for r in body]
    if labels != list(COMPONENTS):
        raise DesnzFormatError(f"{name}: row labels {labels} != {list(COMPONENTS)}")
    cells: dict[str, list[Cell]] = {}
    for r, label in zip(body, labels, strict=True):
        if r[1] != UNIT_CELL:
            raise DesnzFormatError(f"{name} {label}: unit {r[1]!r} is not {UNIT_CELL!r}")
        if any(v is not None for v in r[width:]):
            raise DesnzFormatError(f"{name} {label}: values beyond the last column header")
        row = []
        for h, v in zip(header[2:width], r[2:width], strict=True):
            if v == LESS_THAN_ONE:
                row.append(Cell(None))
            elif isinstance(v, int | float) and not isinstance(v, bool):
                row.append(Cell(Decimal(repr(v))))
            else:
                raise DesnzFormatError(f"{name} {label} / {h!r}: {v!r} is neither a number nor '<1'")
        cells[COMPONENTS[label][0]] = row
    return Sheet(year=year, columns=columns, cells=cells)


def check_totals(s: Sheet) -> None:
    problems = []
    parts = [cid for cid, _ in COMPONENTS.values() if cid != "total"]
    slack = Decimal("0.5") * len(parts)
    for i, c in enumerate(s.columns):
        lo = sum((s.cells[p][i].value or Decimal(0) for p in parts), Decimal(0))
        hi = lo + sum(Decimal(1) for p in parts if s.cells[p][i].value is None)
        total = s.cells["total"][i].value
        if total is None or not (lo - slack <= total <= hi + slack):
            problems.append(f"{s.year} {c.header!r}: total {total} vs components {lo}..{hi}")
    if problems:
        raise DesnzFormatError("; ".join(problems))


@lru_cache(maxsize=2)
def read_annex(path: Path) -> list[Sheet]:
    # Content-addressed path (pipeline/.snapshots/<sha256>), so caching by path is safe.
    wb = openpyxl.load_workbook(io.BytesIO(path.read_bytes()), read_only=True, data_only=True)
    home = [v for r in wb["Home"].iter_rows(values_only=True) for v in r if isinstance(v, str)]
    if HOME_NOTE not in home:
        raise DesnzFormatError(f"the Home sheet no longer says {HOME_NOTE!r}")
    sheets = [read_sheet([tuple(r) for r in wb[SHEET.format(year=y)].iter_rows(values_only=True)], y) for y in YEARS]
    for s in sheets:
        check_totals(s)
    return sheets


def observations(sheets: list[Sheet]) -> list[Observation]:
    by_series: dict[tuple[str, str, str], list[Observation]] = {}
    for s in sheets:
        for i, c in enumerate(s.columns):
            for label, (cid, _) in COMPONENTS.items():
                cell = s.cells[cid][i]
                by_series.setdefault((c.technology, c.estimate, cid), []).append(
                    Observation(
                        entity=ENTITY,
                        period=str(s.year),
                        value=None if cell.value is None else float(cell.value),
                        missing_reason=None
                        if cell.value is not None
                        else f"DESNZ prints '<1' for {label.lower()}: less than £1 per megawatt-hour, with no "
                        "number published.",
                        status="projection",
                        dims={"technology": c.technology, "estimate": c.estimate, "component": cid},
                    )
                )
    order = {t: i for i, t in enumerate(TECHNOLOGIES)}
    comp = {cid: i for i, (cid, _) in enumerate(COMPONENTS.values())}
    return [
        o
        for key in sorted(by_series, key=lambda k: (order[k[0]], k[1] != "unmarked", comp[k[2]]))
        for o in by_series[key]
    ]


def _run(files: dict[str, InputFile]) -> Result:
    f = files[ANNEX.key]
    sheets = read_annex(f.path)
    n_cols = sum(len(s.columns) for s in sheets)
    return Result(
        observations=observations(sheets),
        vintage=VINTAGE,
        date_published=PUBLISHED,
        steps=[
            "Read the five 'Additional Estimates' sheets (2030 to 2050) of Annex A (sha256 "
            f"{f.snapshot.sha256[:12]}…), {n_cols} technology columns in all, after checking each sheet's title, "
            "its 2024 real pounds line, its eleven row labels and its units.",
            "Read each column header as printed: footnote numbers removed, '(FOAK)' kept as the 'estimate' "
            "dimension, the load factor kept as part of the technology. Cells printed '<1' are null values with that "
            "reason.",
            "Checked that each total lies within the sum of its rounded components plus or minus half a pound per "
            "component. Values are published as DESNZ prints them; nothing is computed across technologies, and no "
            "total without carbon costs is derived (carbon costs are their own published row).",
        ],
    )


def transforms(paths: Paths) -> list[Transform]:
    return [
        Transform(
            spec=Spec(
                id=INDICATOR,
                title="Cost of electricity from new UK power plants, by component (DESNZ)",
                description="What it would cost, per megawatt-hour, to build and run a new power plant in the United "
                "Kingdom starting between 2030 and 2050, by type of plant, as estimated by the Department for Energy "
                "Security and Net Zero: wind, solar, tidal and geothermal beside gas plants with and without carbon "
                "capture and hydrogen plants, split into construction, running, fuel and carbon costs. Gas and "
                "hydrogen plants are shown at three illustrative shares of the year in operation. DESNZ says it is "
                "not appropriate to compare these costs across technologies that play very different roles in the "
                "power system.",
                kind="series",
                unit=Unit(code="GBP2024/MWh", label="pounds (2024) per megawatt-hour", short="£/MWh"),
                display=Display(decimals=0),
                scope=Scope(
                    geography="United Kingdom",
                    basis="Levelised cost of electricity for projects commissioning in the year shown, in real 2024 "
                    "pounds, on DESNZ's central assumptions: fuel prices from DESNZ's fossil fuel price assumptions "
                    "and carbon costs from UK ETS carbon price assumptions plus the Carbon Price Support (a published "
                    "component, included in the total). Technology-specific hurdle rates. Dispatchable plants at "
                    "illustrative load factors of 5% (peaking), 30% (mid-merit) and 93% (baseload), or 88% for gas "
                    "with carbon capture because of the assumed availability of the CO2 transport and storage "
                    "network. Offshore wind excludes grid connection costs; other technologies include them. "
                    "Components are rounded and may not sum to the total. Projection.",
                ),
                geo_coverage="country",
                headline_entity=ENTITY,
                dimensions=(
                    Dimension(
                        id="technology",
                        label="Plant type and load factor",
                        values=[DimensionValue(id=k, label=v) for k, v in TECHNOLOGIES.items()],
                    ),
                    Dimension(
                        id="estimate",
                        label="Estimate",
                        values=[DimensionValue(id=k, label=v) for k, v in ESTIMATES.items()],
                    ),
                    Dimension(
                        id="component",
                        label="Cost component",
                        values=[DimensionValue(id=cid, label=lab) for cid, lab in COMPONENTS.values()],
                    ),
                ),
                headline_dims=(
                    ("technology", "gas-ccgt-93pc-load"),
                    ("estimate", "unmarked"),
                    ("component", "total"),
                ),
            ),
            inputs=(ANNEX,),
            run=_run,
            module_file=Path(__file__),
            validation=Validation(min_rows=1500, value_range=(-1000.0, 2000.0)),
        )
    ]
