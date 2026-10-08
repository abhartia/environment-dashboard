"""EIA Residential Energy Consumption Survey 2020: the average US home's electricity use, as published and as average
power.

Inputs: Table CE2.1 (eia-recs-2020/table-ce2-1, a workbook) and the survey's consumption technical documentation
(eia-recs-2020/ce-technical-documentation, a PDF).

What is read. Sheet 'physical units', row 'All homes' (row 5): column G 'Electricity (kWh)' under 'Average site energy
consumption (per household using the fuel)', with column B (housing units, million) and column C (total electricity,
billion kWh). Every heading, the table title, the release dates and the two footnotes the scope relies on (on-site
solar included; the 50 states and the District of Columbia, primary occupied homes only) must read as expected. Every
home uses electricity, so the average per household using it is the average per home: the build checks that the
total divided by the number of homes equals the printed average to within the rounding of the two printed inputs. The
relative standard error of the average (sheet 'rse', column I) goes in the value's note.

Average power. The technical documentation (page 13) defines the survey's reference period: "For electricity and
natural gas, we defined a complete set of billing records as having a series of monthly bills that covered the 366
total days from January 1 to December 31, 2020." That sentence is required before the conversion: the annual
kilowatt-hours are divided by the hours from 1 January 2020 to 1 January 2021 (8,784) to give kilowatts, in exact
decimal arithmetic (envdash.power).
"""

from __future__ import annotations

import io
from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from pathlib import Path

import openpyxl

from envdash import power, textmatch
from envdash.models import Display, Observation, Scope, Unit
from envdash.paths import Paths
from envdash.transform import Input, InputFile, Result, Spec, Transform, Validation

SOURCE = "eia-recs-2020"
TABLE = Input(SOURCE, "table-ce2-1")
DOCS = Input(SOURCE, "ce-technical-documentation")
ENTITY = "USA"
PERIOD = "2020"
VINTAGE = "RECS 2020, Table CE2.1 (revised March 2024)"
PUBLISHED = "2024-03"
START, END = date(2020, 1, 1), date(2021, 1, 1)

SHEET = "physical units"
EXPECTED: dict[str, str] = {
    "A1": "Data release date: March 2023\nRevised data release date: March 2024",
    "A2": "Table CE2.1  Annual household site fuel consumption in the U.S.—totals and averages, 2020",
    "B3": "Number of housing units (million)",
    "C3": "Total site energy consumptiona",
    "G3": "Average site energy consumptiona\n(per household using the fuel)",
    "B4": "Total U.S.b",
    "C4": "Electricity (billion kWh)",
    "G4": "Electricity (kWh)",
    "A5": "All homes",
}
RSE_SHEET = "rse"
RSE_EXPECTED: dict[str, str] = {
    "H3": "RSEs for average site energy consumptiona",
    "I4": "Electricity",
    "A5": "All homes",
}
FOOTNOTES = (
    "Electricity consumption from on-site solar photovoltaic generation (that is, solar panels) is included.",
    "Total U.S. includes all primary occupied housing units in the 50 states and the District of Columbia. Vacant "
    "housing units, seasonal units, second homes, military houses, and group quarters are excluded.",
)
FOOTER_CELL = "A95"
PERIOD_PAGE = 13
PERIOD_SENTENCE = (
    "For electricity and natural gas, we defined a complete set of billing records as having a series of monthly "
    "bills that covered the 366 total days from January 1 to December 31, 2020."
)


class RecsFormatError(ValueError):
    pass


@dataclass(frozen=True)
class Average:
    kwh: Decimal
    homes_million: Decimal
    total_billion_kwh: Decimal
    rse_percent: Decimal


def _dec(v: object, where: str) -> Decimal:
    if isinstance(v, bool) or not isinstance(v, int | float):
        raise RecsFormatError(f"{where} is {v!r}, not a number")
    return Decimal(str(v))


def _half_unit(v: Decimal) -> Decimal:
    exp = v.normalize().as_tuple().exponent
    assert isinstance(exp, int)
    return Decimal(5) * Decimal(10) ** (min(exp, 0) - 1)


def read(raw: bytes) -> Average:
    wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=False, data_only=True)
    ws, rse = wb[SHEET], wb[RSE_SHEET]
    for sheet, expected in ((ws, EXPECTED), (rse, RSE_EXPECTED)):
        for cell, text in expected.items():
            if sheet[cell].value != text:
                raise RecsFormatError(f"{sheet.title}!{cell} is {sheet[cell].value!r}, expected {text!r}")
    footer = ws[FOOTER_CELL].value or ""
    for f in FOOTNOTES:
        if not textmatch.contains(str(footer), f):
            raise RecsFormatError(f"the footnotes no longer say {f!r}")
    avg = Average(
        kwh=_dec(ws["G5"].value, "G5"),
        homes_million=_dec(ws["B5"].value, "B5"),
        total_billion_kwh=_dec(ws["C5"].value, "C5"),
        rse_percent=_dec(rse["I5"].value, "rse!I5"),
    )
    check_average(avg)
    return avg


def check_average(a: Average) -> None:
    """Total over homes equals the printed average, to within the rounding of the printed total and count."""
    implied = a.total_billion_kwh * 1000 / a.homes_million
    tol = (
        _half_unit(a.total_billion_kwh) * 1000 / a.homes_million
        + a.total_billion_kwh * 1000 * _half_unit(a.homes_million) / a.homes_million**2
        + _half_unit(a.kwh)
    )
    if abs(implied - a.kwh) > tol:
        raise RecsFormatError(
            f"total {a.total_billion_kwh} billion kWh over {a.homes_million} million homes is {implied:.1f} kWh, not "
            f"the printed average {a.kwh} kWh: the average may no longer be over all homes"
        )


def require_period(pdf: bytes) -> None:
    pages = textmatch.pdf_pages_text(pdf)
    if len(pages) < PERIOD_PAGE or not textmatch.contains(pages[PERIOD_PAGE - 1], PERIOD_SENTENCE):
        raise RecsFormatError(f"page {PERIOD_PAGE} of the technical documentation no longer defines 2020 as 366 days")


def _note(a: Average) -> str:
    return f"EIA's relative standard error for this estimate: {a.rse_percent}% (Table CE2.1, sheet 'rse')."


def _read_step(files: dict[str, InputFile], a: Average) -> str:
    t = files[TABLE.key]
    return (
        f"Read Table CE2.1 (sha256 {t.snapshot.sha256[:12]}…), sheet 'physical units', row 'All homes': average site "
        f"electricity consumption {a.kwh} kWh per household, after checking the title, the headings, the release "
        "dates and the footnotes. Every home uses electricity, so this is the average per home: total consumption "
        f"({a.total_billion_kwh} billion kWh) over the number of homes ({a.homes_million} million) agrees with it to "
        "within rounding."
    )


def _run_energy(files: dict[str, InputFile]) -> Result:
    a = read(files[TABLE.key].path.read_bytes())
    return Result(
        observations=[Observation(entity=ENTITY, period=PERIOD, value=float(a.kwh), note=_note(a))],
        vintage=VINTAGE,
        date_published=PUBLISHED,
        steps=[_read_step(files, a), "Published the average as printed."],
    )


def _run_power(files: dict[str, InputFile]) -> Result:
    a = read(files[TABLE.key].path.read_bytes())
    d = files[DOCS.key]
    require_period(d.path.read_bytes())
    hours = power.hours_between(START, END)
    kw = power.average_power(a.kwh, hours)
    return Result(
        observations=[Observation(entity=ENTITY, period=PERIOD, value=float(kw), note=_note(a))],
        vintage=VINTAGE,
        date_published=PUBLISHED,
        steps=[
            _read_step(files, a),
            "Unit conversion to average power: divided the annual kilowatt-hours by the "
            f"{power.hours_words(hours)} from 1 January 2020 to 31 December 2020, giving kilowatts (kWh / h = kW), in "
            "exact decimal arithmetic. The period is EIA's: page 13 of the survey's consumption technical "
            f"documentation (sha256 {d.snapshot.sha256[:12]}…) says the bills cover 'the 366 total days from "
            "January 1 to December 31, 2020', and that sentence was found before converting. The result is the "
            "constant power that would deliver the same electricity over the year, not the home's peak demand.",
        ],
        changes="annual electricity use divided by the 8,784 hours of 2020 to give average power in kilowatts.",
    )


GEOGRAPHY = "United States: all primary occupied homes in the 50 states and the District of Columbia"
BASIS = (
    "Survey estimate (sample of households and their energy suppliers' bills) of site electricity use per home in "
    "calendar year 2020, including electricity from the home's own solar panels. Vacant and seasonal homes, second "
    "homes, military housing and group quarters are excluded."
)


def transforms(paths: Paths) -> list[Transform]:
    here = Path(__file__)
    return [
        Transform(
            spec=Spec(
                id="electricity-use.eia-recs-2020.us-household",
                title="Electricity used by the average US home in a year",
                description="Electricity used at home in 2020 by the average US household, in kilowatt-hours, from "
                "the US Energy Information Administration's Residential Energy Consumption Survey.",
                kind="series",
                unit=Unit(
                    code="kWh-per-household-per-year",
                    label="kilowatt-hours per household per year",
                    short="kWh/household/yr",
                ),
                display=Display(decimals=0),
                scope=Scope(geography=GEOGRAPHY, basis=BASIS),
                geo_coverage="country",
                headline_entity=ENTITY,
            ),
            inputs=(TABLE,),
            run=_run_energy,
            module_file=here,
            validation=Validation(min_rows=1, value_range=(1_000.0, 50_000.0)),
        ),
        Transform(
            spec=Spec(
                id="power-scale.eia-recs-2020.us-household",
                title="Average electric power of the average US home",
                description="The electricity the average US household used at home in 2020, expressed as average "
                "power in kilowatts: the year's kilowatt-hours spread evenly over its 8,784 hours. From the US Energy "
                "Information Administration's Residential Energy Consumption Survey.",
                kind="derived",
                unit=Unit(code="kW", label="kilowatts (average over the year)", short="kW"),
                display=Display(decimals=2),
                scope=Scope(
                    geography=GEOGRAPHY,
                    basis=BASIS + " Average power is the annual energy divided by the hours in 2020 (8,784), not the "
                    "home's peak demand.",
                ),
                geo_coverage="country",
                headline_entity=ENTITY,
            ),
            inputs=(TABLE, DOCS),
            run=_run_power,
            module_file=here,
            validation=Validation(min_rows=1, value_range=(0.1, 10.0)),
            key_files=(Path(power.__file__),),
        ),
    ]
