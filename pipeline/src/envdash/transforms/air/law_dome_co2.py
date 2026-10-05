"""Law Dome carbon dioxide over the last 2000 years: CSIRO's 20-year smoothing spline, one value per calendar year.

Input: Law_Dome_GHG_2000years.xlsx (CSIRO Data Access Portal, collection version 3), sheet "Splines fits". Row 1 is
the sheet title, row 2 says "The spline fits follow Enting (1987) and attenuate variations with periods of less than
20 years by 50% for CO2 and CH4, ...", row 4 is the column header, and from row 5 each gas has its own three columns
(year AD, spline value, growth rate). The CO2 columns are A ("Year AD") and B ("CO2 spline (20 yr, ppm)"), years 154 to
1996 with one row per year. We publish column B as printed, for each year of column A, as the ISO year (0154 to 1996).
The growth-rate column and the other gases are not published here.

Why the spline. The measurement sheets (CO2byAge, CO2byCore) give one value per ice or firn sample at a fractional
gas age, sometimes several in a year and none in others; a calendar-year series of those would need us to bin or
interpolate, which we do not do. The spline is the producer's own annual series.

Vintage. The workbook carries no version. The CSIRO Data Access Portal serves version 3 of the collection as
collection 63432 (published 4 September 2024, a metadata-only change; the ReadMe sheet still says "LAST UPDATE:
11/2018"), so the version comes from the collection number in the download URL through COLLECTIONS below. An unlisted
collection, or a ReadMe that no longer says the expected last update, stops the build until a person checks it.

No uncertainty is published: the spline sheet gives none. The producer's description of the spline (row 2) is the
condition for reading the sheet; if it changes, the transform stops.
"""

from __future__ import annotations

import io
import re
from datetime import date
from pathlib import Path

import openpyxl

from envdash.models import Display, Observation, Scope, Unit
from envdash.paths import Paths
from envdash.transform import Input, InputFile, Result, Spec, Transform, Validation

SOURCE = "law-dome-2k"
WORKBOOK = Input(SOURCE, "law-dome-ghg-2000years")
SHEET = "Splines fits"
TITLE = "Spline fits to the Law Dome firn and ice core records"
METHOD = (
    "The spline fits follow Enting (1987) and attenuate variations with periods of less than 20 years by 50% for CO2 "
    "and CH4"
)
YEAR_HEADER, CO2_HEADER = "Year AD", "CO2 spline (20 yr, ppm)"
FIRST_DATA_ROW = 5

# CSIRO DAP collection number -> (version label, publication date of that version, the ReadMe's "LAST UPDATE").
# https://data.csiro.au/dap/ws/v2/collections/63432 (read 2026-10-04): version 3, published 2024-09-04.
COLLECTIONS: dict[str, tuple[str, date, str]] = {
    "63432": ("v3", date(2024, 9, 4), "11/2018"),
}
COLLECTION_IN_URL = re.compile(r"/collections/(\d+)/data/")


class LawDomeFormatError(ValueError):
    pass


def collection_of(url: str | None) -> tuple[str, date, str]:
    m = COLLECTION_IN_URL.search(url or "")
    if not m:
        raise LawDomeFormatError(f"cannot read a CSIRO DAP collection number from {url!r}")
    if m.group(1) not in COLLECTIONS:
        raise LawDomeFormatError(
            f"collection {m.group(1)} is not in COLLECTIONS; add its version from the CSIRO Data Access Portal"
        )
    return COLLECTIONS[m.group(1)]


def read_spline(raw: bytes, last_update: str) -> list[Observation]:
    wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    readme = [str(r[0]) for r in wb["ReadMe"].iter_rows(max_col=1, values_only=True) if r[0] is not None]
    if f"LAST UPDATE: {last_update}" not in readme:
        raise LawDomeFormatError(f"the ReadMe sheet no longer says 'LAST UPDATE: {last_update}'")
    if SHEET not in wb.sheetnames:
        raise LawDomeFormatError(f"no sheet {SHEET!r} in {wb.sheetnames}")
    rows = list(wb[SHEET].iter_rows(min_col=1, max_col=2, values_only=True))
    if rows[0][0] != TITLE or not str(rows[1][0] or "").startswith(METHOD):
        raise LawDomeFormatError(f"{SHEET}: title or spline description changed; re-read the sheet")
    if rows[3] != (YEAR_HEADER, CO2_HEADER):
        raise LawDomeFormatError(f"{SHEET}: columns A-B are {rows[3]}, expected {(YEAR_HEADER, CO2_HEADER)}")
    obs: list[Observation] = []
    previous: int | None = None
    ended = False
    for n, (year, co2) in enumerate(rows[FIRST_DATA_ROW - 1 :], start=FIRST_DATA_ROW):
        if year is None and co2 is None:
            ended = True
            continue
        if ended or not isinstance(year, int) or not isinstance(co2, int | float):
            raise LawDomeFormatError(f"{SHEET} row {n}: unexpected cells {(year, co2)!r} in the CO2 columns")
        if previous is not None and year != previous + 1:
            raise LawDomeFormatError(f"{SHEET} row {n}: year {year} does not follow {previous}")
        if not 1 <= year <= 9999:
            raise LawDomeFormatError(f"{SHEET} row {n}: year {year} is outside ISO years 0001-9999")
        obs.append(Observation(entity="LAWDOME", period=f"{year:04d}", value=float(co2)))
        previous = year
    return obs


def _run(files: dict[str, InputFile]) -> Result:
    f = files[WORKBOOK.key]
    version, published, last_update = collection_of(str(f.snapshot.url) if f.snapshot.url else None)
    obs = read_spline(f.path.read_bytes(), last_update)
    return Result(
        observations=obs,
        vintage=version,
        date_published=published.isoformat(),
        steps=[
            f"Read the sheet '{SHEET}' of Law_Dome_GHG_2000years.xlsx, CSIRO collection version {version} (data last "
            f"updated {last_update}).",
            f"Published the column '{CO2_HEADER}' as printed for each year of its 'Year AD' column, {obs[0].period} "
            f"to {obs[-1].period}, as ISO calendar years. The spline is CSIRO's fit to the ice and firn measurements, "
            "which halves variations with periods shorter than 20 years; the growth-rate column and the individual "
            "samples are not published here.",
        ],
    )


def transforms(paths: Paths) -> list[Transform]:
    return [
        Transform(
            spec=Spec(
                id="co2.law-dome.2k",
                title="Carbon dioxide over the last 2,000 years (Law Dome ice cores)",
                description="Carbon dioxide in air trapped in ice and firn at Law Dome, East Antarctica, as CSIRO's "
                "smoothing-spline fit for each year from 154 to 1996 CE.",
                kind="series",
                unit=Unit(code="ppm", label="parts per million", short="ppm"),
                display=Display(decimals=1),
                scope=Scope(
                    geography="Law Dome, East Antarctica (66.7° S, 112.8° E)",
                    basis="Smoothing spline (Enting 1987) through ice-core (DSS, DSS0506, DE08, DE08-2) and firn-air "
                    "(DE08-2, DSSW20K) measurements by gas age; variations with periods under 20 years are halved.",
                ),
                geo_coverage="global-only",
                headline_entity="LAWDOME",
            ),
            inputs=(WORKBOOK,),
            run=_run,
            module_file=Path(__file__),
            validation=Validation(min_rows=1800, value_range=(250.0, 400.0)),
        )
    ]
