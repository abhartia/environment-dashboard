"""IDMC Global Internal Displacement Database (GIDD): internal displacements caused by weather-related disasters, for
the world, by year and hazard type, calculated from IDMC's disaster event file.

Input: artifact gidd-disaster-events of source idmc-gidd, the GIDD disaster export (xlsx). Until IDMC issues our API
key, a person downloads it from the GIDD page (Cause: Disaster, then Excel) and records it with envdash snapshot add.
It has two sheets. README holds the title, "DATE EXTRACTED: <date>", "LAST UPDATE: <date>", the file's licence line
and IDMC's definitions. 1_Disaster_Displacement_data has one row per disaster event in a country and year, with the
columns in COLUMNS. No fixture: the class is noncommercial, so the tests that read the file are marked snapshot.

What is counted. Internal displacements are movements, not people: IDMC counts each time a person is forced to leave
home within their own country, so one person displaced twice in a year is counted twice, and pre-emptive evacuations
are included. Each row's figure is IDMC's estimate for that event in that country and year. Year is the year IDMC
reports the displacement in; an event can have started the year before.

What is summed. Only rows whose Hazard Category is exactly "Weather related", IDMC's own category (CRED EM-DAT
classification). That leaves out IDMC's "Geophysical" rows (earthquakes, volcanic activity and dry mass movement) and
its "Mixed disasters" rows (displacement IDMC does not assign to one hazard), with no judgement of ours. The raw column
(Disaster Internal Displacements (Raw)) of the remaining rows is added up as whole numbers per Year and Hazard Type,
and over the nine types (hazard "all"). The rounded column is not used: IDMC rounds it at national level.

China before 2017. IDMC files China's displacement for 2008 and for 2010-2016 as one countrywide row a year under
"Mixed disasters" (26,822,000 in 2008), not split by hazard; from 2017 (and in 2009) China has one row per event. So
that figure is not in the weather-related sums of those years, which are not comparable with later ones. Nothing is
estimated in its place. The description says so (the years are STATED_CHINA_MIXED, checked against the file), and
each year's total carries a note with the "Mixed disasters" figures left out and the largest of them.

Missing figures. A row with no raw figure is left out, never counted as zero. A published value whose year and type
have such rows carries a note with their number, the steps give the count per year, and the description states the
total, checked against the file. A hazard type with no row at all in a year has no value for that year.

Checks rather than assumptions. The README's title and licence line, the column headers, the hazard categories, the
types in each and the sub-types of each weather-related type must be the ones listed here, so a change in the file's
vocabulary is read by a person before it is counted (the dimension labels rely on the sub-types). Raw and rounded
figures are whole numbers, never negative, present or absent together. Years run from FIRST_YEAR without gaps and end
before the year the file was extracted. The counts the description states (STATED_*) are checked against the file,
so a new release stops the build until a person re-reads it and updates them.

No publisher check. IDMC's Global Report on Internal Displacement prints weather-related and per-hazard totals that
are built differently (GRID 2026 gives 27.2 million weather-related displacements in 2025, where these rows add up to
26,933,321), and IDMC publishes no total built from this file the same way, so none is declared and GRID figures are
never mixed in. The API documentation's caution against summing figures that share dates and locations is about
IDMC's figure-level data (it names the "Recommended Figure" role); each row here is one event's total for a country
and year, as the README defines it.
"""

from __future__ import annotations

import io
import re
from collections import Counter, defaultdict
from collections.abc import Iterable, Iterator
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

import openpyxl

from envdash.models import Dimension, DimensionValue, Display, Observation, Scope, Unit
from envdash.paths import Paths
from envdash.transform import Input, InputFile, Result, Spec, Transform, Validation

SOURCE = "idmc-gidd"
EVENTS = Input(SOURCE, "gidd-disaster-events")
INDICATOR = "displacement.idmc-gidd.weather-by-hazard"
FIRST_YEAR = 2008

README_SHEET = "README"
DATA_SHEET = "1_Disaster_Displacement_data"
COLUMNS = (
    "ISO3",
    "Country / Territory",
    "Year",
    "Event Name",
    "Date of Event (start)",
    "Disaster Internal Displacements",
    "Disaster Internal Displacements (Raw)",
    "Hazard Category",
    "Hazard Type",
    "Hazard Sub Type",
    "Event Codes (Code:Type)",
    "Event ID",
    "Displacement occurred",
)
TITLE = "TITLE: Global Internal Displacement Database (GIDD) - Disasters"
LICENCE_LINE = "USE LICENSE: This content is licensed under CC BY-NC."
LAST_UPDATE = re.compile(r"LAST UPDATE: ([A-Z][a-z]+ \d{1,2}, \d{4})")
DATE_EXTRACTED = re.compile(r"DATE EXTRACTED: ([A-Z][a-z]+ \d{1,2}, \d{4})")
ISO3 = re.compile(r"[A-Z][A-Z0-9]{2}")

WEATHER = "Weather related"
MIXED = "Mixed disasters"
# IDMC's Hazard Type, our dimension value id and label, and the sub-types IDMC files under it, in dimension order.
HAZARDS: tuple[tuple[str, str, str, frozenset[str]], ...] = (
    ("Flood", "flood", "Flood", frozenset({"Flood", "Dam release flood"})),
    (
        "Storm",
        "storm",
        "Storm",
        frozenset(
            {
                "Storm",
                "Tornado",
                "Typhoon/Hurricane/Cyclone",
                "Hailstorm",
                "Storm surge",
                "Winter storm/Blizzard",
                "Sand/dust storm",
            }
        ),
    ),
    ("Wildfire", "wildfire", "Wildfire", frozenset({"Wildfire"})),
    (
        "Mass Movement",
        "mass-movement",
        "Landslide or avalanche (wet mass movement)",
        frozenset({"Landslide/Wet mass movement", "Avalanche"}),
    ),
    ("Drought", "drought", "Drought", frozenset({"Drought"})),
    ("Erosion", "erosion", "Erosion", frozenset({"Erosion"})),
    (
        "Extreme Temperature",
        "extreme-temperature",
        "Extreme temperature (cold wave or heat wave)",
        frozenset({"Cold wave", "Heat wave"}),
    ),
    ("Wave action", "wave-action", "Wave action", frozenset({"Rogue Wave"})),
    ("Sea level Rise", "sea-level-rise", "Sea level rise", frozenset({"Sea level rise"})),
)
# IDMC's other categories and the types in each: read and checked, never summed.
OTHER_CATEGORIES: dict[str, frozenset[str]] = {
    "Geophysical": frozenset({"Earthquake", "Mass Movement", "Volcanic activity"}),
    MIXED: frozenset({"Mixed disasters"}),
}
ALL = ("all", "All weather-related hazards")

# Counts the description states, read from the release whose README says "LAST UPDATE: August 31, 2026". The build
# checks them against the file, so a new release stops the build until a person re-reads it and updates them.
STATED_WITHOUT_FIGURE = 701
"""Weather-related rows with no reported figure."""
STATED_LAST_YEAR = 2025
STATED_ROWS = {FIRST_YEAR: 1044, STATED_LAST_YEAR: 7064}
"""Weather-related rows (with or without a figure) in the first and last year."""
CHINA = "CHN"
STATED_CHINA_MIXED = (2008, 2010, 2011, 2012, 2013, 2014, 2015, 2016)
"""The years in which China's displacement is a "Mixed disasters" row with a figure."""


class IdmcFormatError(ValueError):
    pass


# --- reading the export -------------------------------------------------------------------------------------------


@dataclass(frozen=True)
class Readme:
    last_update: date
    extracted: date


@dataclass(frozen=True)
class Event:
    iso3: str
    country: str
    year: int
    category: str
    hazard_type: str
    raw: int | None


def _date(text: str) -> date:
    return datetime.strptime(text, "%B %d, %Y").date()


def _one(lines: list[str], rx: re.Pattern[str]) -> date:
    found = [m.group(1) for line in lines if (m := rx.fullmatch(line))]
    if len(found) != 1:
        raise IdmcFormatError(f"README has {len(found)} lines matching {rx.pattern!r}; expected one")
    return _date(found[0])


def read_readme(lines: list[str]) -> Readme:
    """lines: the README sheet's first-column texts, stripped."""
    if TITLE not in lines:
        raise IdmcFormatError(f"README has no line {TITLE!r}: not the GIDD disaster export")
    if not any(line.startswith(LICENCE_LINE) for line in lines):
        raise IdmcFormatError(
            f"README has no line starting {LICENCE_LINE!r}: the file's own licence statement changed. Re-read it and "
            "the registry entry before building."
        )
    return Readme(last_update=_one(lines, LAST_UPDATE), extracted=_one(lines, DATE_EXTRACTED))


def parse_events(header: tuple, rows: Iterable[tuple], readme: Readme) -> list[Event]:
    if tuple(header) != COLUMNS:
        raise IdmcFormatError(f"{DATA_SHEET!r} columns {header} != expected {COLUMNS}")
    col = {c: i for i, c in enumerate(COLUMNS)}
    weather_types = {t: subs for t, _, _, subs in HAZARDS}
    out: list[Event] = []
    for n, r in enumerate(rows, start=2):
        if len(r) != len(COLUMNS):
            raise IdmcFormatError(f"row {n} has {len(r)} cells, not {len(COLUMNS)}")
        iso3, country, year = r[col["ISO3"]], r[col["Country / Territory"]], r[col["Year"]]
        if not isinstance(iso3, str) or not ISO3.fullmatch(iso3):
            raise IdmcFormatError(f"row {n}: ISO3 {iso3!r} is not a three-character code")
        if not isinstance(country, str) or not country.strip():
            raise IdmcFormatError(f"row {n}: Country / Territory {country!r} is not a name")
        if isinstance(year, bool) or not isinstance(year, int):
            raise IdmcFormatError(f"row {n}: Year {year!r} is not a whole year")
        if not FIRST_YEAR <= year < readme.extracted.year:
            raise IdmcFormatError(
                f"row {n}: Year {year} is outside {FIRST_YEAR}–{readme.extracted.year - 1} (the file was extracted on "
                f"{readme.extracted}; a year not yet over would be published as final)"
            )
        cat, typ, sub = r[col["Hazard Category"]], r[col["Hazard Type"]], r[col["Hazard Sub Type"]]
        if cat == WEATHER:
            if typ not in weather_types:
                raise IdmcFormatError(f"row {n}: weather-related Hazard Type {typ!r} is not one this transform reads")
            if sub not in weather_types[typ]:
                raise IdmcFormatError(f"row {n}: {typ} Hazard Sub Type {sub!r} is not one this transform reads")
        elif cat in OTHER_CATEGORIES:
            if typ not in OTHER_CATEGORIES[cat]:
                raise IdmcFormatError(f"row {n}: {cat} Hazard Type {typ!r} is not one this transform reads")
        else:
            raise IdmcFormatError(f"row {n}: Hazard Category {cat!r} is not one this transform reads")
        raw, rounded = r[col["Disaster Internal Displacements (Raw)"]], r[col["Disaster Internal Displacements"]]
        for name, v in (("raw", raw), ("rounded", rounded)):
            if v is not None and (isinstance(v, bool) or not isinstance(v, int) or v < 0):
                raise IdmcFormatError(f"row {n}: {name} displacements {v!r} is not a non-negative whole number")
        if (raw is None) != (rounded is None):
            raise IdmcFormatError(f"row {n}: raw {raw!r} and rounded {rounded!r} figures are not present together")
        out.append(Event(iso3=iso3, country=country.strip(), year=year, category=cat, hazard_type=typ, raw=raw))
    if not out:
        raise IdmcFormatError(f"{DATA_SHEET!r} has no event rows")
    return out


def read_sheets(raw: bytes) -> tuple[list[str], tuple, Iterator[tuple]]:
    """The README's first-column texts (stripped), the data sheet's header and an iterator over its other rows."""
    wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    for sheet in (README_SHEET, DATA_SHEET):
        if sheet not in wb.sheetnames:
            raise IdmcFormatError(f"no sheet {sheet!r} in {wb.sheetnames}")
    lines = [r[0].strip() for r in wb[README_SHEET].iter_rows(values_only=True) if r and isinstance(r[0], str)]
    rows = wb[DATA_SHEET].iter_rows(values_only=True)
    header = next(rows, None)
    if header is None:
        raise IdmcFormatError(f"{DATA_SHEET!r} is empty")
    return lines, header, rows


def read_export(raw: bytes) -> tuple[Readme, list[Event]]:
    lines, header, rows = read_sheets(raw)
    readme = read_readme(lines)
    return readme, parse_events(header, rows, readme)


# --- adding up ----------------------------------------------------------------------------------------------------


@dataclass
class Tally:
    """Weather-related rows by (year, IDMC Hazard Type): the sum of the raw figures, and how many rows had a figure."""

    sums: dict[tuple[int, str], int] = field(default_factory=lambda: defaultdict(int))
    with_figure: Counter[tuple[int, str]] = field(default_factory=Counter)
    without_figure: Counter[tuple[int, str]] = field(default_factory=Counter)
    left_out: Counter[str] = field(default_factory=Counter)
    """Rows of the other categories, by category."""
    mixed: dict[int, list[Event]] = field(default_factory=lambda: defaultdict(list))
    """Rows of the category "Mixed disasters" with a figure, by year: displacement not assigned to one hazard."""

    def years(self) -> list[int]:
        return sorted({y for y, _ in (*self.with_figure, *self.without_figure)})

    def rows(self, year: int) -> tuple[int, int]:
        """(rows with a figure, rows without one) in a year, over every weather-related type."""
        w = sum(n for (y, _), n in self.with_figure.items() if y == year)
        wo = sum(n for (y, _), n in self.without_figure.items() if y == year)
        return w, wo


def tally(events: list[Event]) -> Tally:
    t = Tally()
    for e in events:
        if e.category != WEATHER:
            t.left_out[e.category] += 1
            if e.category == MIXED and e.raw is not None:
                t.mixed[e.year].append(e)
        elif e.raw is None:
            t.without_figure[(e.year, e.hazard_type)] += 1
        else:
            t.sums[(e.year, e.hazard_type)] += e.raw
            t.with_figure[(e.year, e.hazard_type)] += 1
    years = t.years()
    if not years or years != list(range(FIRST_YEAR, years[-1] + 1)):
        raise IdmcFormatError(f"weather-related rows cover the years {years}, not {FIRST_YEAR} onwards without gaps")
    for y in years:
        if t.rows(y)[0] == 0:
            raise IdmcFormatError(f"no weather-related row in {y} has a reported figure")
    return t


def check_stated(t: Tally) -> None:
    """The counts the description states must be the file's."""
    without = sum(t.without_figure.values())
    found = {y: sum(t.rows(y)) for y in STATED_ROWS}
    problems = []
    if without != STATED_WITHOUT_FIGURE:
        problems.append(f"{without} weather-related rows have no figure, not {STATED_WITHOUT_FIGURE}")
    if t.years()[-1] != STATED_LAST_YEAR:
        problems.append(f"the last year is {t.years()[-1]}, not {STATED_LAST_YEAR}")
    problems += [f"{y} has {found[y]} weather-related rows, not {n}" for y, n in STATED_ROWS.items() if found[y] != n]
    # One row a year, in exactly the stated years (a year listed twice means two rows that year).
    china = sorted(y for y, rows in t.mixed.items() for e in rows if e.iso3 == CHINA)
    if china != list(STATED_CHINA_MIXED):
        problems.append(f"China's '{MIXED}' rows with a figure are in {china}, not one in each of {STATED_CHINA_MIXED}")
    if problems:
        raise IdmcFormatError(
            "the description's counts no longer match the file (" + "; ".join(problems) + "). Re-read the new release "
            "and update STATED_* and the description in this module."
        )


def _missing_note(missing: int, total: int, what: str, year: int) -> str | None:
    if not missing:
        return None
    verb = "has" if missing == 1 else "have"
    return (
        f"{missing:,} of the {total:,} {what} event rows IDMC lists for {year} {verb} no reported figure and "
        f"{'is' if missing == 1 else 'are'} not in this sum."
    )


def _mixed_note(rows: list[Event], year: int) -> str | None:
    """The "Mixed disasters" figures of a year, which no weather-related sum includes, and the largest of them."""
    if not rows:
        return None
    total = sum(e.raw or 0 for e in rows)
    top = max(rows, key=lambda e: e.raw or 0)
    return (
        f"Also not in this sum: {total:,} displacements in {len(rows):,} {'row' if len(rows) == 1 else 'rows'} IDMC "
        f"files for {year} under '{MIXED}', not assigned to one hazard (the largest: {top.country}, {top.raw:,})."
    )


def _total_note(t: Tally, year: int) -> str | None:
    w, wo = t.rows(year)
    parts = [_missing_note(wo, w + wo, "weather-related", year), _mixed_note(t.mixed.get(year, []), year)]
    return " ".join(p for p in parts if p) or None


def observations(t: Tally) -> list[Observation]:
    obs: list[Observation] = []
    for y in t.years():
        value = sum(v for (yy, _), v in t.sums.items() if yy == y)
        obs.append(
            Observation(
                entity="WLD",
                period=f"{y:04d}",
                value=float(value),
                note=_total_note(t, y),
                dims={"hazard": ALL[0]},
            )
        )
    for typ, hid, _label, _subs in HAZARDS:
        for y in t.years():
            w, wo = t.with_figure[(y, typ)], t.without_figure[(y, typ)]
            if w == 0 and wo == 0:
                continue
            if w == 0:
                obs.append(
                    Observation(
                        entity="WLD",
                        period=f"{y:04d}",
                        value=None,
                        missing_reason=f"None of the {wo:,} {typ.lower()} event rows IDMC lists for {y} has a "
                        "reported figure.",
                        dims={"hazard": hid},
                    )
                )
                continue
            obs.append(
                Observation(
                    entity="WLD",
                    period=f"{y:04d}",
                    value=float(t.sums[(y, typ)]),
                    note=_missing_note(wo, w + wo, typ.lower(), y),
                    dims={"hazard": hid},
                )
            )
    return obs


def _run(files: dict[str, InputFile]) -> Result:
    f = files[EVENTS.key]
    readme, events = read_export(f.path.read_bytes())
    t = tally(events)
    check_stated(t)
    years = t.years()
    weather = sum(t.with_figure.values()) + sum(t.without_figure.values())
    per_year = "; ".join(f"{y} {t.rows(y)[0]:,}/{t.rows(y)[1]:,}" for y in years)
    mixed = "; ".join(f"{y} {sum(e.raw or 0 for e in rows):,}" for y, rows in sorted(t.mixed.items()))
    steps = [
        f"Read the sheet {DATA_SHEET!r} of IDMC's GIDD disaster export (README: last update "
        f"{readme.last_update.isoformat()}, extracted {readme.extracted.isoformat()}; recorded "
        f"{f.snapshot.date_accessed.isoformat()}, sha256 {f.snapshot.sha256[:12]}…): {len(events):,} rows for "
        f"{years[0]}–{years[-1]}, one per disaster event in a country and year.",
        f"Kept the {weather:,} rows whose Hazard Category is 'Weather related' (IDMC's own category, CRED EM-DAT "
        f"classification). Left out {t.left_out['Geophysical']:,} 'Geophysical' rows (earthquakes, volcanic activity "
        f"and dry mass movement) and {t.left_out[MIXED]:,} '{MIXED}' rows (displacement IDMC does not assign to one "
        f"hazard). For China, IDMC gives one '{MIXED}' row a year in {_years_text(STATED_CHINA_MIXED)}, so China's "
        "displacement in those years is not in the sums; nothing is estimated in its place. 'Mixed disasters' "
        f"figures left out, by year: {mixed}; each year's total carries a note with them.",
        "Added up the raw figure (Disaster Internal Displacements (Raw)) of these rows by Year, the year IDMC reports "
        "the displacement in, and by IDMC's Hazard Type, and over all nine types, as whole numbers. The rounded "
        "column is not used.",
        f"Left out the {sum(t.without_figure.values()):,} weather-related rows with no reported figure; they are "
        "never counted as zero, and each value whose rows include some carries a note with their number. Rows with a "
        f"figure and without one, by year: {per_year}.",
        "A hazard type with no row in a year has no value for that year. No check against the totals printed in "
        "IDMC's Global Report on Internal Displacement: they are built differently and are not used.",
    ]
    return Result(
        observations=observations(t),
        vintage=f"GIDD disaster export, last update {readme.last_update.isoformat()}",
        date_published=readme.last_update.isoformat(),
        steps=steps,
        changes="Calculated from IDMC event data: added up the raw displacement figures of weather-related disaster "
        "events by year and hazard type, and over all types.",
    )


def _years_text(years: tuple[int, ...]) -> str:
    """(2008, 2010, 2011, 2012) -> "2008 and 2010–2012"."""
    runs: list[list[int]] = []
    for y in sorted(years):
        if runs and y == runs[-1][-1] + 1:
            runs[-1].append(y)
        else:
            runs.append([y])
    parts = [f"{r[0]}" if len(r) == 1 else f"{r[0]}–{r[-1]}" for r in runs]
    return parts[0] if len(parts) == 1 else ", ".join(parts[:-1]) + " and " + parts[-1]


DESCRIPTION = (
    "How many times people were forced to leave their homes, staying within their own country, because of floods, "
    "storms, wildfires, landslides, droughts and other weather-related disasters, worldwide, each year since "
    f"{FIRST_YEAR}, by hazard. Calculated from IDMC's event data by adding up IDMC's estimate for each event with a "
    "reported figure. These are movements, not people: one person displaced twice is counted twice, and pre-emptive "
    "evacuations are included. IDMC records far more events than it did as monitoring and reporting widened (its "
    f"file has {STATED_ROWS[FIRST_YEAR]:,} weather-related event rows for {FIRST_YEAR} and "
    f"{STATED_ROWS[STATED_LAST_YEAR]:,} for {STATED_LAST_YEAR}), so a rise is not evidence that disasters are "
    "becoming more frequent. The figures are weather-related; they are not attributed to climate change. For "
    f"{_years_text(STATED_CHINA_MIXED)}, IDMC gives China's displacement as one figure a year that is not split by "
    "hazard, so it is not in those years' sums and they are not comparable with later ones; each year's note gives "
    f"the figures left out. {STATED_WITHOUT_FIGURE} weather-related event rows in this release have no "
    "reported figure and are left out, not counted as zero. Totals printed in IDMC's Global Report on Internal "
    "Displacement are built differently and are not used."
)
BASIS = (
    "IDMC's estimates of internal displacements (forced movements within a country, pre-emptive evacuations "
    "included) for disaster events whose Hazard Category is 'Weather related' (CRED EM-DAT classification), added up "
    "by the year IDMC reports them in and by IDMC's Hazard Type. Geophysical events (earthquakes, volcanic activity, "
    "dry mass movement) and IDMC's 'Mixed disasters' (displacement not assigned to one hazard, including China's "
    f"yearly figure for {_years_text(STATED_CHINA_MIXED)}) are left out. Mass movement here is landslides, wet mass "
    "movement and avalanches; extreme temperature is cold waves and heat waves. Only events with a reported figure "
    "are counted. Displacement by drought, which builds up slowly, is harder to monitor than displacement by sudden "
    "events. Weather-related, not attributed to climate change."
)


def transforms(paths: Paths) -> list[Transform]:
    return [
        Transform(
            spec=Spec(
                id=INDICATOR,
                title="Times people were forced from home by weather-related disasters, by hazard (estimated)",
                description=DESCRIPTION,
                kind="derived",
                unit=Unit(
                    code="displacements",
                    label="internal displacements (movements, not people)",
                    short="displacements",
                ),
                display=Display(decimals=0),
                scope=Scope(
                    geography="World: the countries and territories in IDMC's event file (ISO 3166 codes, plus AB9 "
                    "for the Abyei Area)",
                    basis=BASIS,
                ),
                geo_coverage="global-only",
                headline_entity="WLD",
                dimensions=(
                    Dimension(
                        id="hazard",
                        label="Hazard",
                        values=[DimensionValue(id=ALL[0], label=ALL[1])]
                        + [DimensionValue(id=hid, label=label) for _, hid, label, _ in HAZARDS],
                    ),
                ),
                headline_dims=(("hazard", ALL[0]),),
            ),
            inputs=(EVENTS,),
            run=_run,
            module_file=Path(__file__),
            validation=Validation(min_rows=100, value_range=(0.0, 1.0e9)),
        )
    ]
