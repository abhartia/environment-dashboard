"""Lancet Countdown 2025 transforms (envdash/transforms/impacts/lancet_countdown.py).

The source is CC BY-NC-SA 4.0 (class noncommercial). Committed fixtures are cut only from open-class sources
(AGENTS.md rule 9; tests/fixtures/make_fixture.py refuses other classes), so the value tests below read the full
real workbooks from the snapshot cache and are marked `snapshot` (run with `uv run pytest -m snapshot`). The
key-finding sentences are compared with the same workbook's own Global sheet, i.e. the producer's statement for the
same vintage. The tests without raw data check the declared constants: publisher checks, producer labels, what is
left out, and the wording limits.
"""

from __future__ import annotations

import io
from decimal import ROUND_HALF_UP, Decimal

import openpyxl
import pytest

from envdash import geo, snapshots
from envdash.paths import Paths
from envdash.registry import load_registry
from envdash.transforms.impacts import lancet_countdown as lc

# openpyxl warns that it drops drawing shapes when it re-saves a workbook (the refusal tests below); values are kept.
pytestmark = pytest.mark.filterwarnings("ignore:DrawingML support is incomplete:UserWarning")

SOURCE = "lancet-countdown-2025"
IDS = {
    "heat.lancet-2025.deaths-global",
    "heat.lancet-2025.labour-hours-global",
    "heat.lancet-2025.heatwave-days-global",
    "heat.lancet-2025.heatwave-days-countries",
    "heat.lancet-2025.labour-hours-per-worker",
    "heat.lancet-2025.heatwave-days-vulnerable",
    "smoke.lancet-2025.wildfire-deaths",
    "drought.lancet-2025.land-share",
    "rain.lancet-2025.extreme-events",
    "air-pollution.lancet-2025.deaths-by-sector-fuel",
    "air-pollution.lancet-2025.fossil-deaths",
}


def _current_bytes(source_id: str, artifact_id: str) -> bytes:
    # Copied from tests/test_snapshot_golden.py.
    p = Paths.default()
    return snapshots.cache_path(p, snapshots.read_current(p)[f"{source_id}/{artifact_id}"]).read_bytes()


def _replaced(raw: bytes, sheet: str, row_with: str, old, new) -> bytes:
    """The real workbook with one cell changed, to show that a changed file is refused: in the one row of `sheet`
    containing the text `row_with`, the cell holding `old` is set to `new`."""
    wb = openpyxl.load_workbook(io.BytesIO(raw))
    rows = [r for r in wb[sheet].iter_rows() if any(c.value == row_with for c in r)]
    assert len(rows) == 1, f"{len(rows)} rows of {sheet!r} contain {row_with!r}"
    cells = [c for c in rows[0] if c.value == old]
    assert len(cells) == 1
    cells[0].value = new
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _with_cell(raw: bytes, sheet: str, row_starts: tuple, column: int, new) -> bytes:
    """The real workbook with one cell changed: in the one row of `sheet` whose first cells are `row_starts`, the cell
    in (0-based) `column` is set to `new`."""
    wb = openpyxl.load_workbook(io.BytesIO(raw))
    rows = [r for r in wb[sheet].iter_rows() if tuple(c.value for c in r[: len(row_starts)]) == row_starts]
    assert len(rows) == 1, f"{len(rows)} rows of {sheet!r} start with {row_starts!r}"
    rows[0][column].value = new
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _round(x: Decimal, places: str) -> Decimal:
    return x.quantize(Decimal(places), rounding=ROUND_HALF_UP)


def _mean(values: list[Decimal]) -> Decimal:
    return sum(values, Decimal()) / len(values)


# --- no raw data needed -------------------------------------------------------------------------------------------


def test_transforms_declare_registered_inputs():
    paths = Paths.default()
    ts = {t.spec.id: t for t in lc.transforms(paths)}
    assert set(ts) == IDS
    src = load_registry(paths).sources[SOURCE]
    artifacts = {a.id for a in src.artifacts}
    for t in ts.values():
        assert [i.source_id for i in t.inputs] == [SOURCE]
        assert {i.artifact_id for i in t.inputs} <= artifacts
    # The registry notice carries the Copernicus (ERA5) acknowledgement every one of these indicators needs.
    assert src.obligations.notice and "Copernicus Climate Change Service information {year}" in src.obligations.notice
    check = ts["heat.lancet-2025.labour-hours-global"].checks[0]
    assert check.vintage == lc.VINTAGE and check.stated == "640" and check.tolerance == 0.5
    assert check.url == str(next(a.url for a in src.artifacts if a.id == lc.WORK_HOURS.artifact_id))


def test_publisher_checks_of_the_new_workbooks():
    paths = Paths.default()
    ts = {t.spec.id: t for t in lc.transforms(paths)}
    url = {a.id: str(a.url) for a in load_registry(paths).sources[SOURCE].artifacts}
    got = {
        (i, c.period, c.stated, c.tolerance, dict(c.dims).get("months") or dict(c.dims).get("fuel"))
        for i, t in ts.items()
        for c in t.checks
        if i != "heat.lancet-2025.labour-hours-global"
    }
    # Each stated value is written with exactly the digits the producer printed ("154,000", "2.68 million",
    # "60.7%", "five"), so each tolerance is half its last printed digit.
    assert got == {
        ("smoke.lancet-2025.wildfire-deaths", "2024", "154E3", 500.0, None),
        ("drought.lancet-2025.land-share", "2024", "60.7", 0.05, "1m"),
        ("rain.lancet-2025.extreme-events", "2024", "5", 0.5, None),
        ("air-pollution.lancet-2025.fossil-deaths", "2010", "2.68E6", 5000.0, "total"),
        ("air-pollution.lancet-2025.fossil-deaths", "2022", "2.52E6", 5000.0, "total"),
    }
    for i, inp in (
        ("smoke.lancet-2025.wildfire-deaths", lc.SMOKE),
        ("drought.lancet-2025.land-share", lc.DROUGHT),
        ("rain.lancet-2025.extreme-events", lc.RAIN),
        ("air-pollution.lancet-2025.fossil-deaths", lc.AIR),
    ):
        assert ts[i].inputs == (inp,)
        assert all(c.vintage == lc.VINTAGE and c.url == url[inp.artifact_id] for c in ts[i].checks)
    assert "2.68 million in 2010" in lc.AIR_KEY_FINDING and "2.52 million in 2022" in lc.AIR_KEY_FINDING
    assert "record-high 154,000 in 2024" in lc.SMOKE_KEY_FINDING
    assert "record-breaking 60.7% in 2024" in lc.DROUGHT_KEY_FINDING
    assert lc.RAIN_QUOTE.startswith("in 2024 ") and lc.RAIN_QUOTE.endswith("record high of five")


def test_declared_constants_of_the_new_workbooks():
    ts = {t.spec.id: t for t in lc.transforms(Paths.default())}
    # The fossil-fuel deaths are the one calculated indicator; everything else is copied.
    assert {i for i, t in ts.items() if t.spec.kind == "derived"} == {"air-pollution.lancet-2025.fossil-deaths"}
    # Dimension labels are the producer's own (the Global sheet's Sector values and fuel headings).
    assert [label for _, label in lc.AIR_SECTORS] == [
        "Power plants",
        "Industry",
        "Transport",
        "Households",
        "Waste",
        "Agriculture",
        "Other",
    ]
    assert tuple(label for _, label, _ in lc.AIR_FUELS) == lc.AIR_HEADER[1][3:7]
    assert [lc.AIR_HEADER[0][col] for *_, col in lc.AIR_FUELS] == ["Deaths", None, None, None]
    assert set(lc.FOSSIL_FUELS) <= {f for f, *_ in lc.AIR_FUELS}
    assert [d for d, _ in lc.FOSSIL_DIMS] == ["total", *lc.FOSSIL_FUELS]
    # 1.2.2: the 12-month share is never published; 1.2.3: no row is expected for 1991-1994.
    assert [k for k, *_ in lc.DURATIONS] == ["1m", "3m", "6m"] and lc.LEFT_OUT_DURATION == "12m"
    assert lc.RAIN_PERIODS == ((1961, 1990), (1995, 2024))
    assert [g for g, *_ in lc.GROUPS] == ["over-65", "infants"]
    assert all("75" not in col for _, _, col, _ in lc.GROUPS)
    headlines = {i: dict(t.spec.headline_dims) for i, t in ts.items()}
    assert headlines["air-pollution.lancet-2025.deaths-by-sector-fuel"] == {"sector": "power-plants", "fuel": "coal"}
    assert headlines["air-pollution.lancet-2025.fossil-deaths"] == {"fuel": "total"}
    # The wording claims only what the producer measured.
    d = {i: t.spec.description for i, t in ts.items()}
    assert "not attributed" in d["smoke.lancet-2025.wildfire-deaths"]
    assert "not attributed" in d["heat.lancet-2025.heatwave-days-vulnerable"]
    assert "not drought caused by climate change" in d["drought.lancet-2025.land-share"]
    assert "warmer air itself pushes it towards drought" in d["drought.lancet-2025.land-share"]
    assert "not floods" in d["rain.lancet-2025.extreme-events"]
    for i in ("air-pollution.lancet-2025.deaths-by-sector-fuel", "air-pollution.lancet-2025.fossil-deaths"):
        assert "a harm of the burning itself, not of the warming" in d[i]
        assert "modelled" in d[i]
    # The notice names ERA5 only for the indicators computed from it, and credits SPEI for 1.2.2.
    notice = load_registry(Paths.default()).sources[SOURCE].obligations.notice or ""
    assert "Indicators 1.1.1, 1.1.3, 1.1.5 and 1.2.3 are computed from ERA5 or ERA5-Land" in notice
    assert "Indicator 1.2.2 is computed from the Standardised Precipitation-Evapotranspiration Index" in notice


# --- full real workbooks (snapshot cache) -------------------------------------------------------------------------


@pytest.mark.snapshot
def test_deaths_and_the_key_finding():
    raw = _current_bytes(SOURCE, lc.MORTALITY.artifact_id)
    obs, g = lc.parse_mortality(raw)
    assert g.report_year == 2025 and g.published.isoformat() == "2025-10-29"
    by = {o.period: o.value for o in obs}
    assert len(obs) == 32 and obs[0].period == "1990" and obs[-1].period == "2021"
    assert by["2021"] == 558304 and by["1990"] == 260952
    assert g.key_finding == (
        "In 2012-2021, global heat-related mortality reached an estimated average 546,000 deaths annually, up 63.2% "
        "from 335,000 in 1990-1999"
    )
    early = sum(Decimal(int(by[str(y)])) for y in range(1990, 2000)) / 10
    late = sum(Decimal(int(by[str(y)])) for y in range(2012, 2022)) / 10
    assert early == Decimal("334690.3")
    assert _round(early, "1E3") == Decimal("335000")  # matches "335,000 in 1990-1999"
    # The same workbook's annual values do not reproduce the other two figures of its key finding: the 2012-2021
    # mean is 545,478.4 (545,000 to the nearest thousand, not 546,000) and the rise is 62.98%, not 63.2%. The WHO,
    # Lancet Countdown region and HDI sheets sum to the same means. Recorded here, not adjusted.
    assert late == Decimal("545478.4")
    assert _round(late, "1E3") == Decimal("545000")
    assert _round((late / early - 1) * 100, "0.01") == Decimal("62.98")


@pytest.mark.snapshot
def test_work_hours_and_the_key_finding():
    raw = _current_bytes(SOURCE, lc.WORK_HOURS.artifact_id)
    obs, g = lc.parse_work_hours(raw)
    assert len(obs) == 5 * 35
    total = {o.period: Decimal(repr(o.value)) for o in obs if o.dims == {"sector": "total"}}
    assert g.key_finding == (
        "A record-high 640 billion potential work hours were lost in 2024, a 98% increase compared to the 1990–99 "
        "annual average."
    )
    assert _round(total["2024"], "1") == Decimal("640")
    assert total["2024"] == max(total.values())  # "record-high"
    base = sum(total[str(y)] for y in range(1990, 2000)) / 10
    assert _round((total["2024"] / base - 1) * 100, "1") == Decimal("98")
    # TotalSunAgCon 639855005.6499999 thousand hours in the Global sheet.
    assert total["2024"] == Decimal("639.85500565")
    sectors = {o.dims["sector"]: o.value for o in obs if o.period == "2024"}
    assert sum(v for k, v in sectors.items() if k != "total") == pytest.approx(sectors["total"], rel=1e-12)


@pytest.mark.snapshot
def test_heatwave_days_and_the_key_finding():
    raw = _current_bytes(SOURCE, lc.HEATWAVE.artifact_id)
    obs, g = lc.parse_heatwave(raw)
    assert len(obs) == 15
    by = {(o.dims["scenario"], o.period): Decimal(repr(o.value)) for o in obs}
    assert by[("observed", "2024")] == Decimal("30.7905888162652")
    assert by[("attributable", "2020")] == Decimal("14.2293997022799")
    assert g.key_finding.startswith(
        "Globally, 84% of the heatwave days that people were exposed to on average annually in 2020-2024, would have "
        "not been expected to occur without climate change."
    )
    years = [str(y) for y in range(2020, 2025)]
    share = sum(by[("attributable", y)] for y in years) / sum(by[("observed", y)] for y in years)
    assert _round(share * 100, "1") == Decimal("84")


def test_country_sheets_are_declared():
    ts = {t.spec.id: t for t in lc.transforms(Paths.default())}
    hw, work = ts["heat.lancet-2025.heatwave-days-countries"], ts["heat.lancet-2025.labour-hours-per-worker"]
    assert hw.inputs == (lc.HEATWAVE,) and work.inputs == (lc.WORK_HOURS,)
    assert hw.spec.geo_coverage == work.spec.geo_coverage == "mixed"
    assert hw.spec.headline_entity == work.spec.headline_entity == "WLD"
    # The headline is the part the producer attributes to climate change; negatives are allowed through validation.
    assert dict(hw.spec.headline_dims) == {"scenario": "attributable"} and hw.validation.value_range[0] < 0
    assert "negative" in hw.spec.description and "not attributed to climate change" in work.spec.description
    # The value columns follow the producer's five key columns (1.1.1 spells two of them in lower case).
    assert lc.HEATWAVE_COUNTRY_COLUMNS[5:] == lc.HEATWAVE_COLUMNS[1:]
    assert (
        lc.WORK_COUNTRY_COLUMNS[: len(lc.COUNTRY_KEYS)] == lc.COUNTRY_KEYS and lc.PER_WORKER in lc.WORK_COUNTRY_COLUMNS
    )
    assert set(lc.MISSING) == {None, "#DIV/0!"}


@pytest.mark.snapshot
def test_heatwave_days_by_country():
    raw = _current_bytes(SOURCE, lc.HEATWAVE.artifact_id)
    obs, _ = lc.parse_heatwave_countries(raw)
    world, _ = lc.parse_heatwave(raw)
    assert [o for o in obs if o.entity == "WLD"] == world
    assert len({o.entity for o in obs}) == 221 and len(obs) == 221 * 3 * 5
    by = {(o.entity, o.dims["scenario"], o.period): Decimal(repr(o.value)) for o in obs}
    # The 2025_Report_Data_Country row United Kingdom, GBR, 2024.
    assert by[("GBR", "attributable", "2024")] == Decimal("5.81840417198212")
    assert by[("GBR", "observed", "2024")] == Decimal("6.62396931964666")
    # Fewer heatwave days than the modelled world without the warming: published as negative, not dropped.
    assert by[("PAK", "attributable", "2021")] == Decimal("-0.652866793187558")
    assert sorted(k for k, v in by.items() if k[1] == "attributable" and v < 0) == [
        ("PAK", "attributable", "2021"),
        ("PAK", "attributable", "2023"),
    ]


@pytest.mark.snapshot
def test_a_changed_country_row_is_refused():
    raw = _current_bytes(SOURCE, lc.HEATWAVE.artifact_id)
    sheet = lc.HEATWAVE_COUNTRY_SHEET
    # One more observed day breaks Observed = Counterfactual + Attributable_to_CC.
    with pytest.raises(lc.LancetFormatError, match="GBR 2024 Observed"):
        lc.parse_heatwave_countries(
            _with_cell(raw, sheet, ("United Kingdom", "GBR", "Europe", "Very high", 2024), 5, 7.62396931964666)
        )
    # A code outside the crosswalk is never guessed.
    with pytest.raises(geo.UnknownEntity, match="XGB"):
        lc.parse_heatwave_countries(
            _with_cell(raw, sheet, ("United Kingdom", "GBR", "Europe", "Very high", 2024), 1, "XGB")
        )
    # A blank value stops the build rather than leaving a hole in a sheet that is complete.
    with pytest.raises(lc.LancetFormatError, match="empty value"):
        lc.parse_heatwave_countries(
            _with_cell(raw, sheet, ("United Kingdom", "GBR", "Europe", "Very high", 2024), 6, None)
        )


@pytest.mark.snapshot
def test_work_hours_per_worker_by_country():
    raw = _current_bytes(SOURCE, lc.WORK_HOURS.artifact_id)
    obs, _, left_out = lc.parse_work_per_worker(raw)
    by = {(o.entity, o.period): o for o in obs}
    # The Global sheet's TotalSunWHLpp for 2024, and the United Kingdom's.
    assert Decimal(repr(by[("WLD", "2024")].value)) == Decimal("178.05728508386184")
    assert by[("GBR", "2024")].value == pytest.approx(0.15226597195484218, rel=1e-15)
    assert len({o.entity for o in obs}) == 220
    # No one employed in the four sectors: a division by zero in the sheet, published as missing with that reason.
    assert {k for k, v in left_out.items()} == {"#DIV/0!", "None"} and len(left_out["#DIV/0!"]) == 204
    andorra = [o for o in obs if o.entity == "AND"]
    assert andorra and all(o.value is None and "division by zero" in (o.missing_reason or "") for o in andorra)
    assert by[("ZWE", "2022")].value is None and "blank" in (by[("ZWE", "2022")].missing_reason or "")
    assert by[("ZWE", "2021")].value is not None


@pytest.mark.snapshot
def test_a_per_worker_value_that_is_not_the_total_over_the_employed_is_refused():
    raw = _current_bytes(SOURCE, lc.WORK_HOURS.artifact_id)
    zw = ("Zimbabwe", "ZWE", "Africa", "Medium", 2024)
    with pytest.raises(lc.LancetFormatError, match="ZWE 2024 hours per worker"):
        lc.parse_work_per_worker(
            _with_cell(raw, lc.WORK_COUNTRY_SHEET, zw, lc.WORK_COUNTRY_COLUMNS.index(lc.PER_WORKER), 70.0)
        )
    with pytest.raises(lc.LancetFormatError, match="divides by zero with 5781000"):
        lc.parse_work_per_worker(
            _with_cell(raw, lc.WORK_COUNTRY_SHEET, zw, lc.WORK_COUNTRY_COLUMNS.index(lc.PER_WORKER), "#DIV/0!")
        )


@pytest.mark.snapshot
def test_a_different_report_year_is_refused():
    raw = _current_bytes(SOURCE, lc.MORTALITY.artifact_id)
    report = "The 2025 report of the Lancet Countdown on health and climate change"
    edited = _replaced(raw, "DATA GUIDANCE", report, report, report.replace("2025", "2026"))
    with pytest.raises(lc.LancetFormatError, match="2026 report"):
        lc.parse_mortality(edited)


@pytest.mark.snapshot
def test_a_changed_unit_is_refused():
    raw = _current_bytes(SOURCE, lc.WORK_HOURS.artifact_id)
    edited = _replaced(raw, "DATA GUIDANCE", "TotalsunAgCon", "Hours in 1000s", "Hours")
    with pytest.raises(lc.LancetFormatError, match="TotalsunAgCon"):
        lc.parse_work_hours(edited)


@pytest.mark.snapshot
def test_inconsistent_heatwave_row_is_refused():
    raw = _current_bytes(SOURCE, lc.HEATWAVE.artifact_id)
    # One more observed day in 2020 breaks Observed = Counterfactual + Attributable_to_CC.
    edited = _replaced(raw, lc.HEATWAVE_SHEET, 2020, 16.1167785054611, 17.1167785054611)
    with pytest.raises(lc.LancetFormatError, match="2020 Observed"):
        lc.parse_heatwave(edited)


@pytest.mark.snapshot
def test_vulnerable_heatwave_days_and_the_key_finding():
    raw = _current_bytes(SOURCE, lc.VULNERABLE.artifact_id)
    obs, g = lc.parse_vulnerable(raw)
    assert g.report_year == 2025 and len(obs) == 2 * 45
    by = {(o.dims["group"], o.period): Decimal(repr(o.value)) for o in obs}
    assert {p for _, p in by} == {str(y) for y in range(1980, 2025)}
    assert by[("over-65", "2024")] == Decimal("20.76785436154853")
    assert by[("infants", "2024")] == Decimal("20.45295544481495")
    assert by[("over-65", "1980")] == Decimal("2.818277721614311")
    assert g.key_finding == (
        "In 2024, people older than 65 years and infants younger than 1-year experienced record-high heatwave "
        "exposures – up by 304% and 389%, respectively, from the 1986-2005 baseline."
    )
    # The key finding's percentages are the per-person averages (the published column), against 1986-2005.
    rise = {}
    for group in ("over-65", "infants"):
        series = {p: v for (gr, p), v in by.items() if gr == group}
        assert series["2024"] == max(series.values())  # "record-high"
        base = _mean([series[str(y)] for y in range(1986, 2006)])
        rise[group] = _round((series["2024"] / base - 1) * 100, "1")
    assert rise == {"over-65": Decimal("304"), "infants": Decimal("389")}


@pytest.mark.snapshot
def test_vulnerable_over_75_columns_are_left_out_for_their_stated_reason():
    raw = _current_bytes(SOURCE, lc.VULNERABLE.artifact_id)
    wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    assert "exposure_average_75" not in lc.read_guidance(wb).units
    rows = lc.read_table(wb, lc.VULNERABLE_SHEET, lc.VULNERABLE_COLUMNS, "test")
    i75, i65 = lc.VULNERABLE_COLUMNS.index("exposure_average_75"), lc.VULNERABLE_COLUMNS.index("exposure_average_65")
    assert [r[0] for r in rows if r[i75] == 0] == list(range(1980, 2000))
    assert all(r[i65] > 0 for r in rows)


@pytest.mark.snapshot
def test_wildfire_smoke_deaths_and_the_key_finding():
    raw = _current_bytes(SOURCE, lc.SMOKE.artifact_id)
    obs, g = lc.parse_smoke(raw)
    assert g.report_year == 2025 and g.published.isoformat() == "2025-10-29"
    assert [o.period for o in obs] == [str(y) for y in range(2003, 2025)]
    by = {o.period: Decimal(repr(o.value)) for o in obs}
    assert by["2003"] == Decimal("120418.0420719701") and by["2024"] == Decimal("153718.63398755688")
    assert g.key_finding == lc.SMOKE_KEY_FINDING
    assert by["2024"] == max(by.values())  # "record-high"
    assert _round(by["2024"], "1E3") == Decimal("154000")
    base = _mean([by[str(y)] for y in range(2003, 2013)])
    assert _round((by["2024"] / base - 1) * 100, "1") == Decimal("36")


@pytest.mark.snapshot
def test_drought_land_share_and_the_key_finding():
    raw = _current_bytes(SOURCE, lc.DROUGHT.artifact_id)
    obs, g = lc.parse_drought(raw)
    assert g.report_year == 2025 and len(obs) == 3 * 74
    by = {(o.dims["months"], o.period): Decimal(repr(o.value)) for o in obs}
    assert {p for _, p in by} == {str(y) for y in range(1951, 2025)}
    assert by[("1m", "2024")] == Decimal("60.6941829279543")
    assert by[("3m", "2024")] == Decimal("41.2974184052838")
    assert by[("6m", "2024")] == Decimal("22.8414971224506")
    assert by[("1m", "1951")] == Decimal("19.238488925986")
    assert g.key_finding == lc.DROUGHT_KEY_FINDING
    one = {p: v for (k, p), v in by.items() if k == "1m"}
    assert one["2024"] == max(one.values())  # "record-breaking"
    base = _mean([one[str(y)] for y in range(1951, 1961)])
    assert _round((one["2024"] / base - 1) * 100, "1") == Decimal("299")
    # The 12-month share is left out because it reads NA in 15 of its 74 years.
    wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    twelve = [r for r in wb[lc.DROUGHT_SHEET].iter_rows(values_only=True) if r[1] == lc.LEFT_OUT_DURATION]
    assert len(twelve) == 74 and sum(1 for r in twelve if r[2] == "NA") == 15


@pytest.mark.snapshot
def test_extreme_rain_events_and_the_key_finding():
    raw = _current_bytes(SOURCE, lc.RAIN.artifact_id)
    obs, g = lc.parse_rain(raw)
    assert g.report_year == 2025 and len(obs) == 60
    by = {o.period: Decimal(repr(o.value)) for o in obs}
    assert not {"1991", "1992", "1993", "1994"} & set(by)  # the gap stays a gap
    assert by["1961"] == Decimal("3.33440949301272") and by["2024"] == Decimal("4.96840408088483")
    assert lc.RAIN_QUOTE in g.key_finding
    assert by["2024"] == max(by.values())  # "record high"
    assert _round(by["2024"], "1") == Decimal("5")  # "of five"


@pytest.mark.snapshot
def test_air_pollution_deaths_and_the_key_finding():
    raw = _current_bytes(SOURCE, lc.AIR.artifact_id)
    cells, g = lc.parse_deaths_by_sector_fuel(raw)
    assert g.report_year == 2025 and len(cells) == 7 * 4 * 16
    by = {(o.dims["sector"], o.dims["fuel"], o.period): o.value for o in cells}
    assert {p for *_, p in by} == {str(y) for y in range(2007, 2023)}
    assert by[("power-plants", "coal", "2022")] == 610189
    assert by[("transport", "liquids-gas", "2022")] == 1197085
    assert by[("households", "biomass", "2007")] == 865450
    assert by[("transport", "coal", "2022")] == 0  # a value, per the guidance sheet

    fossil, g2 = lc.parse_fossil_deaths(raw)
    assert g2 == g and len(fossil) == 3 * 16
    f = {(o.dims["fuel"], o.period): o.value for o in fossil}
    for year in range(2007, 2023):
        y = str(year)
        for fuel in lc.FOSSIL_FUELS:
            assert f[(fuel, y)] == sum(by[(s, fuel, y)] for s, _ in lc.AIR_SECTORS)
        assert f[("total", y)] == f[("coal", y)] + f[("liquids-gas", y)]
    assert g.key_finding == lc.AIR_KEY_FINDING
    assert (f[("total", "2010")], f[("total", "2022")]) == (2677081, 2522076)
    assert _round(Decimal(int(f[("total", "2010")])), "1E4") == Decimal("2.68E6")
    assert _round(Decimal(int(f[("total", "2022")])), "1E4") == Decimal("2.52E6")
    change = (Decimal(int(f[("total", "2022")])) / Decimal(int(f[("total", "2010")])) - 1) * 100
    assert _round(change, "0.1") == Decimal("-5.8")  # "decreased by 5.8 %"


@pytest.mark.snapshot
def test_changed_air_pollution_sector_is_refused():
    raw = _current_bytes(SOURCE, lc.AIR.artifact_id)
    edited = _with_cell(raw, lc.AIR_SHEET, (2022, "Power plants"), 1, "Power stations")
    with pytest.raises(lc.LancetFormatError, match="Power stations"):
        lc.parse_deaths_by_sector_fuel(edited)


@pytest.mark.snapshot
def test_missing_drought_share_is_refused_not_filled():
    raw = _current_bytes(SOURCE, lc.DROUGHT.artifact_id)
    edited = _with_cell(raw, lc.DROUGHT_SHEET, (2024, "3m"), 2, "NA")
    with pytest.raises(lc.LancetFormatError, match="2024 3m is 'NA'"):
        lc.parse_drought(edited)


@pytest.mark.snapshot
def test_drought_shares_that_rise_with_duration_are_refused():
    raw = _current_bytes(SOURCE, lc.DROUGHT.artifact_id)
    # More land with six months of drought than with three cannot be.
    edited = _with_cell(raw, lc.DROUGHT_SHEET, (2024, "6m"), 2, 50.0)
    with pytest.raises(lc.LancetFormatError, match="2024 shares do not fall"):
        lc.parse_drought(edited)


@pytest.mark.snapshot
def test_a_rain_year_inside_the_gap_is_refused():
    raw = _current_bytes(SOURCE, lc.RAIN.artifact_id)
    edited = _with_cell(raw, lc.RAIN_SHEET, (1995,), 0, 1994)
    with pytest.raises(lc.LancetFormatError, match="not exactly"):
        lc.parse_rain(edited)
