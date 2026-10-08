"""US EIA, Levelized Costs of New Generation Resources in the Annual Energy Outlook 2026: the levelised cost of
electricity (and of storage) of new plants entering service in 2031, by technology, before and including US tax
credits, as EIA tabulates it.

Inputs (registry entry eia-aeo-2026):
- lcoe-figures: LCOE-LCOS-LACE_figures.xlsx, the data behind the report's figures. Sheets read:
  - "lcoe-components": for each plant type, the capacity factor, the levelised capital, fixed O&M, variable (with
    fuel) and transmission costs, "Total system LCOE or LCOS" (before tax credits), the levelised tax credit, the
    levelised capture credit (45Q) and "Total LCOE or LCOS including tax credit". These are simple averages of the 25
    supply regions (report p. 2: "a simple average (unweighted) of the regional values across the 25 U.S. supply
    regions").
  - "lcoe-regional": the regional minimum and maximum without tax credits, and the minimum, simple average,
    capacity-weighted average and maximum with tax credits. "NB" (not built) marks plant types with no capacity
    added in 2031, which have no capacity-weighted average.
- lcoe-report: the report PDF. Page 8 labels each plant type's average LCOE including tax credits, in dollars and
  cents; the build checks the workbook against those labels (publisher checks).

Published: one indicator, entity USA, period 2031 (the online year), status projection, in 2025 dollars per
megawatt-hour, with dimensions technology, cost_basis and average:
- before-tax-credits / simple-average: "Total system LCOE or LCOS", with the regional minimum and maximum without tax
  credits as its range ("range": across the 25 regions);
- with-tax-credits / simple-average: "Total LCOE or LCOS including tax credit", with the regional minimum and maximum
  with tax credits as its range;
- with-tax-credits / capacity-weighted: lcoe-regional's capacity-weighted average; null where EIA marks it NB.
EIA publishes no capacity-weighted average before tax credits, so there is none here. Values are published as stored
(the workbook keeps full precision; the report rounds to cents). Each value's note gives EIA's capacity factor for
that plant type, because plants are modelled at different capacity factors (combined-cycle 40%, with carbon capture
87%) and EIA warns (report p. 6) that "Direct comparisons of LCOE or LCOS across technologies are misleading as a
method to assess the economic competitiveness". Nothing is computed across rows.

Checks inside the file, or the build stops: the titles name AEO2026, online year 2031 and 2025 dollars (EIA reuses
the file name for each edition); the header rows, section rows and plant types are exactly where this edition has
them; each "Total system" equals the sum of its four components, and each total including tax credit equals the total
plus the credits it lists (within 1e-6); lcoe-regional's simple average with tax credits equals lcoe-components' total
including tax credit; every average lies within its regional minimum and maximum; both sheets carry EIA's data source
line.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass
from decimal import Decimal
from functools import lru_cache
from pathlib import Path

import openpyxl

from envdash import textmatch
from envdash.models import Dimension, DimensionValue, Display, Observation, Scope, Unit
from envdash.paths import Paths
from envdash.transform import Input, InputFile, PublisherCheck, Result, Spec, Transform, Validation

SOURCE = "eia-aeo-2026"
FIGURES = Input(SOURCE, "lcoe-figures")
REPORT = Input(SOURCE, "lcoe-report")
REPORT_URL = "https://www.eia.gov/outlooks/aeo/electricity_generation/pdf/LCOE_report.pdf"
INDICATOR = "lcoe.eia-aeo-2026.new-plants-us"
VINTAGE = "AEO2026 (April 2026)"
PUBLISHED = "2026-04"
ENTITY = "USA"
PERIOD = "2031"
DATA_SOURCE = "Data source:  U.S. Energy Information Administration, Annual Energy Outlook 2026"

COMPONENTS_SHEET = "lcoe-components"
COMPONENTS_TITLE = (
    "Estimated average levelized cost of electricity (LCOE) and levelized cost of storage (LCOS) for new resources",
    "entering service in 2031, AEO2026 Counterfactual Baseline case (2025 dollars per megawatthour)",
)
COMPONENTS_HEADER = (
    "Plant type",
    "Capacity factor (percent)",
    "Levelized capital cost",
    "Levelized fixed O&M",
    "Levelized variable cost",
    "Levelized transmission cost",
    "Total system LCOE or LCOS",
    "Levelized tax credit",
    "Levelized Capture Credit",
    "Total LCOE or LCOS including tax credit",
)
REGIONAL_SHEET = "lcoe-regional"
REGIONAL_TITLE = (
    "Regional variation in levelized cost of electricity (LCOE) and levelized cost of storage (LCOS) for new resources",
    "entering service in 2031, AEO2026 Counterfactual Baseline case (2025 dollars per megawatthour)",
)
REGIONAL_GROUPS = (None, "Without tax credits", None, "With tax credits", None, None, None)
REGIONAL_HEADER = ("Plant type", "Minimum", "Maximum", "Mininum", "Simple average", "Capacity weighted", "Maximum")
SECTIONS = ("Dispatchable technologies", "Resource-contrained technologies", "Capacity resource technologies")
NOT_BUILT = "NB"
NOT_ELIGIBLE = "NA"

# Workbook plant type -> (dimension value id, label, label on report page 8). Order is the workbook's.
TECHNOLOGIES: dict[str, tuple[str, str, str]] = {
    "Advanced nuclear": ("advanced-nuclear", "Advanced nuclear", "advanced nuclear"),
    "Biomass": ("biomass", "Biomass", "biomass"),
    "Combined-cycle": ("gas-combined-cycle", "Natural gas combined-cycle", "combined-cycle"),
    "Combined-cycle with CCS": (
        "gas-combined-cycle-ccs",
        "Natural gas combined-cycle with carbon capture and sequestration (CCS)",
        "combined-cycle with CCS",
    ),
    "Geothermal": ("geothermal", "Geothermal", "geothermal"),
    "Wind, offshore": ("offshore-wind", "Offshore wind", "wind, offshore"),
    "Hydroelectric": ("hydroelectric", "Hydroelectric", "hydroelectric"),
    "Solar, hybrid": (
        "solar-pv-battery",
        "Solar photovoltaic with a four-hour battery (hybrid)",
        "PV-battery hybrid",
    ),
    "Solar, standalone": ("solar-pv", "Solar photovoltaic, standalone", "solar PV"),
    "Wind, onshore": ("onshore-wind", "Onshore wind", "wind, onshore"),
    "Combustion turbine": ("combustion-turbine", "Natural gas combustion turbine", "combustion turbine"),
    "Battery storage": ("battery-storage", "Battery storage (levelised cost of storage)", "battery storage"),
}
COST_BASIS = {
    "before-tax-credits": "Before tax credits (total system cost)",
    "with-tax-credits": "Including tax credits (levelised tax credit and 45Q captured-carbon credit)",
}
AVERAGE = {
    "simple-average": "Simple average of the 25 US supply regions",
    "capacity-weighted": "Average weighted by new capacity in 2031",
}
TOLERANCE = 1e-6
REPORT_PAGE = 8


class EiaLcoeFormatError(ValueError):
    pass


# --- reading ------------------------------------------------------------------------------------------------------


@dataclass(frozen=True)
class Plant:
    capacity_factor: Decimal
    capital: Decimal
    fixed_om: Decimal
    variable: Decimal
    transmission: Decimal
    total: Decimal
    tax_credit: Decimal | None
    capture_credit: Decimal | None
    total_with_credits: Decimal


@dataclass(frozen=True)
class Regional:
    min_without: Decimal
    max_without: Decimal
    min_with: Decimal
    average_with: Decimal
    weighted_with: Decimal | None
    max_with: Decimal


@dataclass(frozen=True)
class Workbook:
    plants: dict[str, Plant]
    regional: dict[str, Regional]


def _rows(ws) -> list[tuple]:
    return [tuple(r) for r in ws.iter_rows(values_only=True)]


def _num(v: object, where: str) -> Decimal:
    if isinstance(v, bool) or not isinstance(v, int | float):
        raise EiaLcoeFormatError(f"{where}: {v!r} is not a number")
    return Decimal(repr(v))


def _num_or(v: object, marker: str, where: str) -> Decimal | None:
    return None if v == marker else _num(v, where)


def _body(rows: list[tuple], title: tuple[str, str], header_rows: list[tuple], width: int, sheet: str) -> list[tuple]:
    """The plant rows of a sheet after its title and header rows, section rows checked; stops at the source line."""
    if (rows[0][0], rows[1][0]) != title:
        raise EiaLcoeFormatError(f"{sheet}: title {(rows[0][0], rows[1][0])!r} is not the AEO2026 title {title!r}")
    start = 2
    for h in header_rows:
        if tuple(rows[start][: len(h)]) != h:
            raise EiaLcoeFormatError(f"{sheet}: row {start + 1} {rows[start]!r} is not the header {h!r}")
        start += 1
    out: list[tuple] = []
    sections: list[str] = []
    for r in rows[start:]:
        label = r[0]
        if label == DATA_SOURCE:
            break
        if label in SECTIONS and all(v is None for v in r[1:]):
            sections.append(label)
            continue
        if label not in TECHNOLOGIES:
            raise EiaLcoeFormatError(f"{sheet}: unexpected row {r!r}")
        out.append(tuple(r[:width]))
    else:
        raise EiaLcoeFormatError(f"{sheet}: no line {DATA_SOURCE!r}")
    if tuple(sections) != SECTIONS:
        raise EiaLcoeFormatError(f"{sheet}: sections {sections} != {list(SECTIONS)}")
    if [r[0] for r in out] != list(TECHNOLOGIES):
        raise EiaLcoeFormatError(f"{sheet}: plant types {[r[0] for r in out]} != {list(TECHNOLOGIES)}")
    return out


def read_components(rows: list[tuple]) -> dict[str, Plant]:
    out = {}
    for r in _body(rows, COMPONENTS_TITLE, [COMPONENTS_HEADER], len(COMPONENTS_HEADER), COMPONENTS_SHEET):
        w = f"{COMPONENTS_SHEET} {r[0]}"
        out[r[0]] = Plant(
            capacity_factor=_num(r[1], w),
            capital=_num(r[2], w),
            fixed_om=_num(r[3], w),
            variable=_num(r[4], w),
            transmission=_num(r[5], w),
            total=_num(r[6], w),
            tax_credit=_num_or(r[7], NOT_ELIGIBLE, w),
            capture_credit=_num_or(r[8], NOT_ELIGIBLE, w),
            total_with_credits=_num(r[9], w),
        )
    return out


def read_regional(rows: list[tuple]) -> dict[str, Regional]:
    out = {}
    for r in _body(rows, REGIONAL_TITLE, [REGIONAL_GROUPS, REGIONAL_HEADER], len(REGIONAL_HEADER), REGIONAL_SHEET):
        w = f"{REGIONAL_SHEET} {r[0]}"
        out[r[0]] = Regional(
            min_without=_num(r[1], w),
            max_without=_num(r[2], w),
            min_with=_num(r[3], w),
            average_with=_num(r[4], w),
            weighted_with=_num_or(r[5], NOT_BUILT, w),
            max_with=_num(r[6], w),
        )
    return out


@lru_cache(maxsize=2)
def read_workbook(path: Path) -> Workbook:
    # Content-addressed path (pipeline/.snapshots/<sha256>), so caching by path is safe.
    wb = openpyxl.load_workbook(io.BytesIO(path.read_bytes()), read_only=True, data_only=True)
    return Workbook(
        plants=read_components(_rows(wb[COMPONENTS_SHEET])),
        regional=read_regional(_rows(wb[REGIONAL_SHEET])),
    )


def check_workbook(w: Workbook) -> None:
    problems: list[str] = []
    for name, p in w.plants.items():
        parts = p.capital + p.fixed_om + p.variable + p.transmission
        if abs(parts - p.total) > Decimal(str(TOLERANCE)):
            problems.append(f"{name}: components sum to {parts}, total system is {p.total}")
        credits = sum((c for c in (p.tax_credit, p.capture_credit) if c is not None), Decimal(0))
        if abs(p.total + credits - p.total_with_credits) > Decimal(str(TOLERANCE)):
            problems.append(f"{name}: total {p.total} + credits {credits} != {p.total_with_credits}")
        g = w.regional[name]
        if abs(g.average_with - p.total_with_credits) > Decimal(str(TOLERANCE)):
            problems.append(f"{name}: regional simple average {g.average_with} != {p.total_with_credits}")
        if not g.min_without <= p.total <= g.max_without:
            problems.append(f"{name}: {p.total} outside {g.min_without}..{g.max_without}")
        if not g.min_with <= p.total_with_credits <= g.max_with:
            problems.append(f"{name}: {p.total_with_credits} outside {g.min_with}..{g.max_with}")
        if g.weighted_with is not None and not g.min_with <= g.weighted_with <= g.max_with:
            problems.append(f"{name}: capacity-weighted {g.weighted_with} outside {g.min_with}..{g.max_with}")
    if problems:
        raise EiaLcoeFormatError("; ".join(problems))


# --- observations and checks --------------------------------------------------------------------------------------


def _cf_note(p: Plant) -> str:
    pct = p.capacity_factor * 100
    shown = f"{pct.normalize():f}" if pct == pct.to_integral_value() else f"{pct}"
    return f"EIA models this plant type at a capacity factor of {shown}%."


def observations(w: Workbook) -> list[Observation]:
    obs = []
    for name, (tid, _, _) in TECHNOLOGIES.items():
        p, g = w.plants[name], w.regional[name]
        note = _cf_note(p)

        def o(
            basis: str,
            average: str,
            value: Decimal | None,
            lo: Decimal | None,
            hi: Decimal | None,
            tid: str = tid,
            note: str = note,
            **kw,
        ) -> Observation:
            return Observation(
                entity=ENTITY,
                period=PERIOD,
                value=None if value is None else float(value),
                lower=None if lo is None else float(lo),
                upper=None if hi is None else float(hi),
                interval=None if lo is None else "range",
                status="projection",
                note=note,
                dims={"technology": tid, "cost_basis": basis, "average": average},
                **kw,
            )

        obs.append(o("before-tax-credits", "simple-average", p.total, g.min_without, g.max_without))
        obs.append(o("with-tax-credits", "simple-average", p.total_with_credits, g.min_with, g.max_with))
        if g.weighted_with is None:
            obs.append(
                o(
                    "with-tax-credits",
                    "capacity-weighted",
                    None,
                    None,
                    None,
                    missing_reason="EIA marks it NB (not built): no new capacity of this type is projected to come "
                    "online in 2031, so there is no capacity-weighted average.",
                )
            )
        else:
            obs.append(o("with-tax-credits", "capacity-weighted", g.weighted_with, None, None))
    return obs


_DOLLARS = re.compile(r"(-?)\$(\d+\.\d\d)")


def report_labels(page_text: str) -> list[tuple[str, str]]:
    """(label, printed value) pairs on report page 8: the twelve values in order, then the twelve labels in the same
    order. Raises if the page no longer lists exactly this edition's plant types in order."""
    values = [m.group(1) + m.group(2) for m in _DOLLARS.finditer(page_text)]
    labels = [lab for _, _, lab in TECHNOLOGIES.values()]
    lines = [ln.strip() for ln in page_text.splitlines()]
    try:
        first = lines.index(labels[0])
    except ValueError:
        raise EiaLcoeFormatError(f"report page {REPORT_PAGE}: no label {labels[0]!r}") from None
    if lines[first : first + len(labels)] != labels:
        raise EiaLcoeFormatError(f"report page {REPORT_PAGE}: labels {lines[first : first + len(labels)]} != {labels}")
    # The axis ticks (-$40 ... $240) follow the twelve labelled values and carry no cents.
    if len(values) != len(labels):
        raise EiaLcoeFormatError(f"report page {REPORT_PAGE}: {len(values)} values in cents, expected {len(labels)}")
    return list(zip(labels, values, strict=True))


REPORT_TITLE = (
    "Average levelized cost of electricity and levelized cost of storage for new resources entering service in 2031 "
    "by technology, AEO2026 Counterfactual Baseline case"
)
REPORT_NOTE = "The stated LCOE values include the levelized tax credit component for eligible technologies."


def require_report(pdf: bytes) -> list[tuple[str, str]]:
    pages = textmatch.pdf_pages_text(pdf)
    page = pages[REPORT_PAGE - 1]
    for q in (REPORT_TITLE, REPORT_NOTE):
        if not textmatch.contains(page, q):
            raise EiaLcoeFormatError(f"report page {REPORT_PAGE} no longer says {q!r}")
    if not textmatch.contains(
        pages[5], "Direct comparisons of LCOE or LCOS across technologies are misleading as a method to assess"
    ):
        raise EiaLcoeFormatError("report page 6 no longer carries EIA's warning against direct comparisons")
    return report_labels(page)


def _checks() -> tuple[PublisherCheck, ...]:
    # The labels on report page 8, as printed on 2026-10-08 (the build re-reads them and fails if they change).
    stated = ("87.81", "84.54", "77.46", "58.47", "40.38", "118.79", "64.77", "94.20", "58.33", "56.75", "172.57")
    stated += ("152.61",)
    return tuple(
        PublisherCheck(
            source_id=SOURCE,
            vintage=VINTAGE,
            entity=ENTITY,
            period=PERIOD,
            stated=s,
            quote=f"${s}",
            url=REPORT_URL,
            dims=(("average", "simple-average"), ("cost_basis", "with-tax-credits"), ("technology", tid)),
        )
        for s, (tid, _, _) in zip(stated, TECHNOLOGIES.values(), strict=True)
    )


CHECKS = _checks()


def _run(files: dict[str, InputFile]) -> Result:
    f = files[FIGURES.key]
    w = read_workbook(f.path)
    check_workbook(w)
    labels = require_report(files[REPORT.key].path.read_bytes())
    printed = [v for _, v in labels]
    if printed != [c.stated for c in CHECKS]:
        raise EiaLcoeFormatError(f"report page {REPORT_PAGE} now prints {printed}; update CHECKS after review")
    return Result(
        observations=observations(w),
        vintage=VINTAGE,
        date_published=PUBLISHED,
        steps=[
            f"Read sheets '{COMPONENTS_SHEET}' and '{REGIONAL_SHEET}' of LCOE-LCOS-LACE_figures.xlsx (sha256 "
            f"{f.snapshot.sha256[:12]}…), after checking that their titles name AEO2026, online year 2031 and 2025 "
            "dollars, and that the plant types and columns are where this edition has them.",
            "Checked inside the file: each total system cost equals the sum of its capital, fixed O&M, variable and "
            "transmission components; each total including tax credits equals the total plus the credits listed; "
            "the regional sheet's simple average with tax credits equals the components sheet's total including "
            "tax credits; every average lies within its regional minimum and maximum.",
            "Published, for each plant type: the total system cost before tax credits and the total including tax "
            "credits (simple averages of the 25 supply regions, each with its regional minimum and maximum as the "
            "range), and the capacity-weighted average including tax credits (null where EIA marks it NB, not "
            "built). Values as stored in the workbook; nothing is computed across plant types.",
            f"Checked the averages including tax credits against the labels on page {REPORT_PAGE} of the report "
            "(publisher checks).",
        ],
    )


def transforms(paths: Paths) -> list[Transform]:
    return [
        Transform(
            spec=Spec(
                id=INDICATOR,
                title="Cost of electricity from new US power plants coming online in 2031 (EIA)",
                description="What it would cost, per megawatt-hour, to build and run a new power plant in the United "
                "States that starts operating in 2031, by type of plant, as projected by the US Energy Information "
                "Administration: solar, wind, nuclear, geothermal and hydro beside natural gas plants with and "
                "without carbon capture, and batteries. Shown before and after US tax credits. Plants of different "
                "types run for different shares of the year and play different roles on the grid, so EIA warns "
                "that these costs alone do not say which plant is the better choice.",
                kind="series",
                unit=Unit(
                    code="USD2025/MWh",
                    label="US dollars (2025) per megawatt-hour",
                    short="USD/MWh",
                ),
                display=Display(decimals=2),
                scope=Scope(
                    geography="United States (25 electricity supply regions)",
                    basis="Levelised cost of electricity (levelised cost of storage for batteries) for plants "
                    "entering service in 2031 in the AEO2026 Counterfactual Baseline case: capital, fixed and "
                    "variable operation and maintenance including fuel, and transmission, over a 30-year cost "
                    "recovery period at an after-tax weighted average cost of capital of 7.27%. Laws and regulations "
                    "effective as of December 2025, including the One Big Beautiful Bill Act's changes to tax "
                    "credits and the EPA's 2024 Section 111 rule (new gas combined-cycle plants without carbon "
                    "capture limited to a 40% capacity factor). No carbon price. Projection; the cost to build and "
                    "run a plant, not the value of its output to the grid.",
                ),
                geo_coverage="country",
                headline_entity=ENTITY,
                dimensions=(
                    Dimension(
                        id="technology",
                        label="Plant type",
                        values=[DimensionValue(id=t, label=lab) for t, lab, _ in TECHNOLOGIES.values()],
                    ),
                    Dimension(
                        id="cost_basis",
                        label="Tax credits",
                        values=[DimensionValue(id=k, label=v) for k, v in COST_BASIS.items()],
                    ),
                    Dimension(
                        id="average",
                        label="Average across regions",
                        values=[DimensionValue(id=k, label=v) for k, v in AVERAGE.items()],
                    ),
                ),
                headline_dims=(
                    ("technology", "gas-combined-cycle"),
                    ("cost_basis", "before-tax-credits"),
                    ("average", "simple-average"),
                ),
            ),
            inputs=(FIGURES, REPORT),
            run=_run,
            module_file=Path(__file__),
            validation=Validation(min_rows=36, value_range=(0.0, 1000.0)),
            checks=CHECKS,
        )
    ]
