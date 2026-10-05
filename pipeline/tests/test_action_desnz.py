"""UK Government conversion factors 2026 (DESNZ): flight emission factors (envdash/transforms/action/desnz_flights.py).

No fixture is committed (the flat file is a 515 kB workbook that cannot be cut into a smaller valid one); the tests
read the full snapshot from the local cache and are marked `snapshot`.
"""

from __future__ import annotations

import io

import openpyxl
import pytest

from envdash import snapshots
from envdash.paths import Paths
from envdash.transform import InputFile, validate
from envdash.transforms.action import desnz_flights as dz

pytestmark = pytest.mark.snapshot


def _current() -> InputFile:
    p = Paths.default()
    sha = snapshots.read_current(p)[dz.FLAT.key]
    return InputFile(snapshots.cache_path(p, sha), snapshots.read_manifest(p, sha))


def _edited(edit) -> bytes:
    """The real workbook with one change, to show what is refused."""
    wb = openpyxl.load_workbook(io.BytesIO(_current().path.read_bytes()))
    edit(wb[dz.SHEET])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _row_of(ws, factor_id: str) -> int:
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, 1).value == factor_id:
            return r
    raise LookupError(factor_id)


def test_28_factors_with_their_ids():
    (t,) = dz.transforms(Paths.default())
    result = t.run({dz.FLAT.key: _current()})
    validate(t, result.observations)
    obs = {(o.dims["haul"], o.dims["seat_class"], o.dims["radiative_forcing"]): o for o in result.observations}
    assert len(obs) == 28
    long_economy = obs[("long-haul", "economy", "with-rf")]
    assert long_economy.value == 0.11704 and "21_316_3170_11_1" in (long_economy.note or "")
    assert obs[("long-haul", "economy", "without-rf")].value == 0.06926
    assert obs[("domestic", "average", "with-rf")].value == 0.22928
    assert obs[("international-non-uk", "first", "with-rf")].value == 0.43663
    # With radiative forcing is always the larger factor.
    for (h, c, rf), o in obs.items():
        if rf == "with-rf":
            assert o.value > obs[(h, c, "without-rf")].value


def test_a_missing_factor_is_refused():
    def drop(ws):
        ws.cell(_row_of(ws, "21_316_3176_11_1"), 3).value = "Business travel- sea"

    with pytest.raises(dz.DesnzFormatError, match="missing"):
        dz.read_factors(_edited(drop))


def test_an_unknown_class_is_refused():
    def rename(ws):
        ws.cell(_row_of(ws, "21_316_3170_11_1"), 6).value = "Economy plus"

    with pytest.raises(dz.DesnzFormatError, match="unknown haul"):
        dz.read_factors(_edited(rename))


def test_a_blank_factor_is_refused():
    def blank(ws):
        ws.cell(_row_of(ws, "21_316_3170_11_1"), 10).value = None

    with pytest.raises(dz.DesnzFormatError, match="not a positive number"):
        dz.read_factors(_edited(blank))
