"""Lelieveld et al. 2023 (BMJ): deaths from fossil-fuel air pollution, world, 2019, from the authors' data deposit.

One indicator, air-pollution.lelieveld-2023.fossil-deaths, from two files of source lelieveld-2023 (Edmond,
doi:10.17617/3.T7XAXH, CC0 1.0): mortality_region_2019_ref.xlsx, the reference model run, and
mortality_region_2019_noFF.xlsx, the same run with all fossil-fuel emissions removed. The paper defines the deaths
attributable to fossil fuels as the difference between the two, and so does this transform: the WORLD row of the column
'NCD+LRI+LRI5-+CPO3' (excess deaths from all causes due to outdoor fine particles and ozone) in the reference file,
less the same cell in the no-fossil file, in exact decimal arithmetic on the stored values. It is a calculation inside
one source, so the indicator is kind `derived` and credited "Calculated by Environment Dashboard from ...".

Files. Each has one sheet whose header names every column: region, population (millions) and, for each cause of death
and the totals, the value and its 95% confidence limits; the cause columns carry the run's suffix ("_ref", "_noFF"),
the last total does not. The header must be exactly as expected, every region must be named once, WORLD among them,
and the two runs must list the same regions with the same population, so a changed file or a file of another run stops
the build until a person re-reads it.

Checks against the paper (BMJ 2023;383:e077784, CC BY-NC 4.0; not fetched, its sentences read on 2026-10-12 and quoted
here). The abstract prints the fossil-fuel deaths as "5.13 million (3.63 to 6.32)": a publisher cross-check on the
difference, tolerance half the last printed digit (±5,000). It prints all outdoor air-pollution deaths as "8.34
million (95% confidence interval 5.63 to 11.19)": the reference run's WORLD total and its confidence limits must round
to those, which confirms that the column read is the one the paper reports.

Ranges. Confidence limits cannot be subtracted (the limits of a difference are not the difference of the limits), so
no range is calculated. The observation's note quotes the paper's own interval.
"""

from __future__ import annotations

import io
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

import openpyxl

from envdash.models import Display, Observation, Scope, Unit
from envdash.paths import Paths
from envdash.transform import Input, InputFile, PublisherCheck, Result, Spec, Transform, Validation

SOURCE = "lelieveld-2023"
REF = Input(SOURCE, "mortality-region-2019-ref")
NO_FOSSIL = Input(SOURCE, "mortality-region-2019-no-fossil")
INDICATOR = "air-pollution.lelieveld-2023.fossil-deaths"
VINTAGE = "Edmond dataset version 2 (2023-11-23)"
PUBLISHED = "2023-11-23"
"""Release of version 2 of the deposit (Dataverse releaseTime 2023-11-23T13:04:44Z)."""
PERIOD = "2019"
"""The year the model represents: 2019 emissions, meteorology and population."""

SHEET = "Sheet 1"
WORLD = "WORLD"
CI = ("CI lower limit", "CI upper limit")
CAUSES = ("NCD+LRI", "LRI5-", "LRI25+", "IHD", "ST", "DM", "LC", "COPD", "CPO3", "OTHERS", "NCD+LRI+LRI5-")
"""The cause columns in file order, before the run's suffix: non-communicable diseases and lower respiratory
infections, LRI in children under 5, LRI in adults over 25, ischaemic heart disease, stroke, type 2 diabetes, lung
cancer, COPD, COPD from ozone, other causes, and the fine-particle total."""
TOTAL = "NCD+LRI+LRI5-+CPO3"
"""All causes, fine particles and ozone: the column the paper's totals come from."""
RUNS = {REF.key: "ref", NO_FOSSIL.key: "noFF"}
"""The suffix each file's cause columns carry."""


def header(run: str) -> tuple[str, ...]:
    causes = tuple(x for c in CAUSES for x in (f"{c}_{run}", *CI))
    return ("region", "pop (1E6)", *CI, *causes, TOTAL, *CI)


TOTAL_COLUMN = len(header("ref")) - 3

URL = "https://doi.org/10.1136/bmj-2023-077784"
ABSTRACT_FOSSIL = (
    "An estimated 5.13 million (3.63 to 6.32) excess deaths per year globally are attributable to ambient air "
    "pollution from fossil fuel use and therefore could potentially be avoided by phasing out fossil fuels."
)
ABSTRACT_ALL = (
    "Globally, all cause excess deaths due to fine particulate and ozone air pollution are estimated at 8.34 million "
    "(95% confidence interval 5.63 to 11.19) deaths per year."
)
PRINTED_ALL = (Decimal("8.34"), Decimal("5.63"), Decimal("11.19"))
"""ABSTRACT_ALL's total and 95% confidence limits, in millions as printed."""
NOTE = (
    f'Lelieveld et al. 2023 (BMJ 383: e077784), abstract: "{ABSTRACT_FOSSIL}" The range in brackets is the paper\'s '
    "95% confidence interval."
)


class LelieveldFormatError(ValueError):
    pass


def dec(x: object, where: str) -> Decimal:
    """A stored number as an exact decimal (the shortest repr of the double)."""
    if isinstance(x, bool) or not isinstance(x, int | float):
        raise LelieveldFormatError(f"{where} is {x!r}, not a number")
    return Decimal(repr(x)) if isinstance(x, float) else Decimal(x)


def read_regions(raw: bytes, run: str, name: str) -> dict[str, tuple]:
    """Each region's row (name stripped of the surrounding whitespace the files carry), checked against the header."""
    wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    if wb.sheetnames != [SHEET]:
        raise LelieveldFormatError(f"{name}: sheets {wb.sheetnames}, expected [{SHEET!r}]")
    rows = list(wb[SHEET].iter_rows(values_only=True))
    if not rows or tuple(rows[0]) != header(run):
        raise LelieveldFormatError(f"{name}: header {rows[0] if rows else None} != expected {header(run)}")
    out: dict[str, tuple] = {}
    for r in rows[1:]:
        if not isinstance(r[0], str) or not r[0].strip():
            raise LelieveldFormatError(f"{name}: a row has region {r[0]!r}")
        region = r[0].strip()
        if region in out:
            raise LelieveldFormatError(f"{name}: region {region!r} appears twice")
        for i, c in enumerate(r[1:], start=1):
            dec(c, f"{name}: {region} {rows[0][i]}")
        out[region] = tuple(r)
    if WORLD not in out:
        raise LelieveldFormatError(f"{name}: no {WORLD} row")
    return out


def _millions(x: Decimal) -> Decimal:
    return (x / Decimal(1_000_000)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def check_reference_total(world: tuple) -> None:
    """The reference run's WORLD total and its 95% limits round to the paper's printed 8.34 million (5.63 to 11.19)."""
    got = tuple(_millions(dec(world[TOTAL_COLUMN + i], f"reference {WORLD} {TOTAL}")) for i in range(3))
    if got != PRINTED_ALL:
        raise LelieveldFormatError(
            f"the reference {WORLD} {TOTAL} and its limits are {got} million; the paper prints {PRINTED_ALL} "
            f"({ABSTRACT_ALL!r})"
        )


def fossil_deaths(ref_raw: bytes, no_fossil_raw: bytes) -> tuple[Decimal, Decimal, Decimal]:
    """(reference deaths, deaths without fossil fuels, their difference): world, all causes, particles and ozone."""
    ref = read_regions(ref_raw, RUNS[REF.key], "reference")
    nof = read_regions(no_fossil_raw, RUNS[NO_FOSSIL.key], "no fossil fuels")
    if list(ref) != list(nof):
        raise LelieveldFormatError(f"the two runs list different regions: {list(ref)} and {list(nof)}")
    for region in ref:
        if ref[region][1:4] != nof[region][1:4]:
            raise LelieveldFormatError(
                f"{region}: the two runs give different population, {ref[region][1:4]} and {nof[region][1:4]}"
            )
    check_reference_total(ref[WORLD])
    a = dec(ref[WORLD][TOTAL_COLUMN], f"reference {WORLD} {TOTAL}")
    b = dec(nof[WORLD][TOTAL_COLUMN], f"no fossil fuels {WORLD} {TOTAL}")
    if b > a:
        raise LelieveldFormatError(f"{WORLD}: more deaths without fossil fuels ({b}) than in the reference run ({a})")
    return a, b, a - b


def observations(ref_raw: bytes, no_fossil_raw: bytes) -> list[Observation]:
    _, _, diff = fossil_deaths(ref_raw, no_fossil_raw)
    return [Observation(entity="WLD", period=PERIOD, value=float(diff), note=NOTE)]


def _run(files: dict[str, InputFile]) -> Result:
    obs = observations(files[REF.key].path.read_bytes(), files[NO_FOSSIL.key].path.read_bytes())
    printed = ", ".join(str(x) for x in PRINTED_ALL[1:])
    return Result(
        observations=obs,
        vintage=VINTAGE,
        date_published=PUBLISHED,
        steps=[
            "Read the WORLD row of mortality_region_2019_ref.xlsx (the reference model run) and of "
            "mortality_region_2019_noFF.xlsx (the run with all fossil-fuel emissions removed) from the Edmond "
            f"deposit of Lelieveld et al. 2023: column {TOTAL!r}, excess deaths in {PERIOD} from all causes due to "
            "outdoor fine particles and ozone. Checked that both files list the same regions with the same "
            "population, and that the reference total and its 95% confidence limits round to the paper's printed "
            f"{PRINTED_ALL[0]} million ({printed} million).",
            "Subtracted the deaths without fossil fuels from those of the reference run, in exact decimal arithmetic "
            "on the stored values: the deaths the paper attributes to fossil-fuel emissions. Confidence limits cannot "
            "be subtracted, so no range is calculated; the observation quotes the paper's own 95% confidence "
            "interval.",
        ],
        changes="the deaths in the run without fossil fuels subtracted from those in the reference run (world, all "
        "causes, fine particles and ozone).",
    )


BASIS = (
    "Modelled excess deaths from all causes due to long-term exposure to outdoor fine particles (PM2.5) and ozone in "
    "2019: exposure from the EMAC atmospheric chemistry model (2019 emissions and meteorology) combined with NASA "
    "satellite-based fine particles and population, and deaths from Burnett and colleagues' FUSION relative-risk "
    "model applied to Global Burden of Disease 2019 baseline mortality. Deaths from fossil fuels are the reference run "
    "less a run with every fossil-fuel emission removed at once. The authors give three reasons their estimate is "
    "higher than most earlier ones: the new relative-risk model, counting deaths from all causes as well as from "
    "specific diseases, and the large fall in exposure when fossil fuels are removed entirely."
)


def transforms(paths: Paths) -> list[Transform]:
    return [
        Transform(
            spec=Spec(
                id=INDICATOR,
                title="Deaths from fossil-fuel air pollution (fine particles and ozone), world (Lelieveld et al. 2023)",
                description="Excess deaths worldwide in 2019 from breathing outdoor fine particles (PM2.5) and ozone "
                "from burning coal, oil and gas, as modelled by Lelieveld and colleagues (BMJ, 2023): a harm of the "
                "burning itself, not of the warming. Calculated from the authors' deposited model runs as the deaths "
                "in the reference run less those in a run with all fossil-fuel emissions removed, as the paper "
                "defines them. The paper's 95% confidence interval is quoted with the value; a range cannot be "
                "calculated by subtracting two ranges, so none is calculated here. A twin of the Lancet Countdown's "
                "estimate (indicator 3.2.1), from a different model; the two are never combined.",
                kind="derived",
                unit=Unit(code="deaths", label="deaths", short="deaths"),
                display=Display(decimals=0),
                scope=Scope(geography="World", basis=BASIS),
                geo_coverage="global-only",
                headline_entity="WLD",
            ),
            inputs=(REF, NO_FOSSIL),
            run=_run,
            module_file=Path(__file__),
            validation=Validation(min_rows=1, value_range=(1.0, 20_000_000.0)),
            # Printed "5.13 million": written in deaths with its three significant digits, so the tolerance is half the
            # last printed digit (±5,000).
            checks=(
                PublisherCheck(
                    source_id=SOURCE,
                    vintage=VINTAGE,
                    entity="WLD",
                    period=PERIOD,
                    stated="5.13E6",
                    quote=ABSTRACT_FOSSIL,
                    url=URL,
                ),
            ),
        )
    ]
