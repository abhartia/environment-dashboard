"""The future map's shapes: World Bank Official Boundaries (GAD V1) simplified into TopoJSON, with CCKP's region codes
matched to its admin-1 shapes.

Inputs, the current snapshots of: wb-gad-v1's ADM0, ADM1 and NDLSA shapefiles (shp, shx and dbf of each; one prj and
one cpg, whose bytes are the same in all three layers) and wb-cckp's subnational-names (CCKP's code and name for every
country and region). Outputs, under data/v1/geo/ (public, listed in SHA256SUMS):

  world.topo.json          object countries: one feature per entity (id: our entity code, through
                           geo.SOURCE_SCHEMES["wb-gad-v1"] on ISO_A3; every part of it in one MultiPolygon, nothing
                           merged), property rp; object disputed: the World Bank's non-determined legal status areas
                           (id "NDLSA:<record>"), properties name (the World Bank's own), rp, and e (the entity of the
                           area's own ISO_A3, when it has one).
  adm1/<entity>.topo.json  for each country with at least one region matched: object regions (id: the CCKP region
                           code, or "V1:<record>" for a shape no CCKP region was matched to; properties m, how it was
                           matched, or "none"; rp; and name, the World Bank's, for an unmatched shape), and object
                           disputed: the disputed areas within 0.01° of the country.
  index.json               models.GeoIndex: origins, processing, licence and credit, each file's sha256, size and
                           counts, and per country the matches by method and what was left unmatched.
  cckp-adm1-crosswalk.csv  every CCKP region and every admin-1 shape, with the pair made and how.

`envdash geo shapes` builds them; `envdash build` runs this when its build key changes (the input snapshots and their
manifests, this module, the declared pairs, pipeline/geo/entities.csv, the data contract, the two source YAMLs and
uv.lock), and the outputs must then be byte-identical to the last build with the same key, or the build stops.

Reading. A .shp is read record by record at the offsets its .shx gives; every record must be a polygon (type 5) or
null. Rings follow the shapefile rule: clockwise rings are outer rings and counter-clockwise rings are holes, each hole
belonging to the smallest outer ring that contains it; a hole inside no outer ring (wrongly wound) is drawn as an outer
ring, and the number of such rings is a processing step. Coordinates stay as written, WGS 84 longitude and latitude:
the .prj must be exactly WGS84_PRJ and the .cpg exactly UTF-8, or the build stops. A polygon record with no rings has no
shape and is counted.

Matching CCKP regions to admin-1 shapes, country by country (CCKP's region prefix is the shape's ISO_A3; CCKP's 32
DIS.* units, areas it assigns to no country, are left out by rule). A code match counts only when the names agree: the
CCKP name equals, after normalisation (norm), one of the shape's own names (NAME_FIELDS); CCKP's placeholders
(NOT_NAMES) agree only with the same placeholder in the World Bank's file, whose names CCKP's are. The name rule never
uses a placeholder, or a name that is the region's own code (Hong Kong's, Macao's and Taiwan's single regions).
  1. gaul: CCKP's code equals {ISO_A3}.{GAUL_0}{GAUL_1} of shapes (both GAUL codes set).
  2. adm1cd: CCKP's code equals {ISO_A3}.{GAUL_0}{int(ADM1CD_t[3:])} of shapes.
  3. declared: pairs in pipeline/geo/cckp-adm1-declared.csv. accept joins the pair; reject refuses a code match.
  4. name: the CCKP name is unique among the country's CCKP regions and is a name of exactly one shape of the
     country, not yet used. Only where the country has as many CCKP regions as shapes: where the counts differ, a
     region was split or merged between the two sets (Indonesia's Papua, split in 2022, is one region in CCKP and
     four shapes here), and a shared name can name a different shape.
A region matched by a code rule takes every shape with that code whose names agree (Wisconsin is two records). The
build stops when the two code rules point at different shapes, a shape would be used twice, a code match whose names
differ is not declared, a declared row no longer matches the files or is no longer needed, or a World Bank code has
no entity. Regions left unmatched are listed (index.json, crosswalk) and never drawn on a shape.

Simplifying. Each file is one topology (topojson): coordinates snapped to a grid (prequantize), shared borders found
once, arcs simplified by Douglas-Peucker without changing which shapes touch (toposimplify, rings kept at least a
triangle), and quantised for the file (topoquantize). Before that, polygons and holes smaller than the file's
minimum area are left out, since they would be under a pixel. The world file uses WORLD; a country file scales the
same steps to the country's extent (its longest side, measured across the 180° line where that is shorter). Every
feature keeps rp, a representative point of its full, unsimplified shape, so a place too small to draw is a dot and is
never lost. The build stops when a feature is missing from its file, a ring does not close, or a file is over its
size limit.
"""

from __future__ import annotations

import csv
import io
import json
import math
import re
import struct
import unicodedata
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import numpy as np
import shapefile
import shapely
import topojson
from shapely.geometry import MultiPolygon, Polygon

from envdash import canonical, geo, snapshots
from envdash import models as models_mod
from envdash.build import PUBLIC_FILES_BASE, lock_sha
from envdash.models import (
    GeoCountry,
    GeoFile,
    GeoIndex,
    Origin,
    ProcessingStep,
    Snapshot,
    Source,
    strictest,
)
from envdash.paths import Paths
from envdash.registry import Registry, render

SOURCE = "wb-gad-v1"
CCKP = "wb-cckp"
NAMES_ARTIFACT = "subnational-names"
LAYERS = ("adm0", "adm1", "ndlsa")
PARTS = ("shp", "shx", "dbf")
INPUT_KEYS: tuple[str, ...] = (
    *(f"{SOURCE}/{layer}-{part}" for layer in LAYERS for part in PARTS),
    f"{SOURCE}/prj",
    f"{SOURCE}/cpg",
    f"{CCKP}/{NAMES_ARTIFACT}",
)
VINTAGE = "GAD V1 (2025-05-07)"

WGS84_PRJ = (
    'GEOGCS["GCS_WGS_1984",DATUM["D_WGS_1984",SPHEROID["WGS_1984",6378137.0,298.257223563]],'
    'PRIMEM["Greenwich",0.0],UNIT["Degree",0.0174532925199433]]'
)
CODE_PAGE = "UTF-8"

NAME_FIELDS = ("NAM_1", "ADM1NM", "NAM_1_SRCE", "P_NAME_1")
"""The record's own names in the World Bank's file: its name, its administrative name, its name in the boundary's
source and in the P-code list. NAM_1_GAUL and NAM_1_STAT are not used: they can name a larger unit the record lies in
(Bogotá's record has Cundinamarca in both, and Cundinamarca's GAUL code)."""
NOT_NAMES = frozenset({"Administrative unit not available", "Name Unknown"})
"""CCKP's placeholders for a region without a name (50 and 18 regions on 2026-10-16)."""
DISPUTED_UNITS = 32
"""CCKP's DIS.* units (areas it assigns to no country), left out by rule; another count stops the build."""

DECLARED = "pipeline/geo/cckp-adm1-declared.csv"
DECLARED_COLUMNS = ("cckp_code", "v1_record", "v1_adm1cd_t", "v1_name", "decision", "reason")
CROSSWALK_COLUMNS = (
    "country",
    "cckp_code",
    "cckp_name",
    "v1_id",
    "v1_adm1cd_t",
    "v1_name",
    "method",
    "names_equal",
)


@dataclass(frozen=True)
class Simplify:
    """One file's simplification, in degrees: snap to `prequantize` steps over the data's box, simplify arcs by
    `epsilon`, leave out polygons and holes under `min_area` square degrees, quantise the output to `quantize`
    steps."""

    prequantize: int
    epsilon: float
    min_area: float
    quantize: int


WORLD = Simplify(prequantize=100_000, epsilon=0.25, min_area=0.05, quantize=10_000)
COUNTRY_EPSILON = 1 / 1500
"""A country file's epsilon, as a share of its extent (its longest side in degrees)."""
COUNTRY_MIN_SIDE = 1 / 300
"""A country file's minimum area is the square of this share of its extent."""
COUNTRY_QUANTIZE = 4000
"""Quantisation steps across the country's extent (more across its box when it straddles the 180° line)."""
TOUCH_DEGREES = 0.01
"""A disputed area goes in a country's file when it lies within this distance of one of the country's shapes."""
WORLD_LIMIT = 250_000
COUNTRY_LIMIT = 200_000
RP_DECIMALS = 4
OBJECT = "__object"
"""A feature's object while all of a file's features go through one topology together."""


class GeoShapesError(Exception):
    """The shapes cannot be built from these files without a person deciding something."""


# --- reading ---------------------------------------------------------------------------------------------------------


@dataclass(frozen=True)
class Record:
    fid: int
    """The record's position in the shapefile, from 0 (ArcGIS's FID)."""
    attrs: dict[str, Any]
    polygons: tuple[Polygon, ...]
    orphan_holes: int = 0
    flat_rings: int = 0

    @property
    def iso(self) -> str:
        return str(self.attrs["ISO_A3"])


def _signed_area(ring: np.ndarray) -> float:
    x, y = ring[:, 0], ring[:, 1]
    return 0.5 * float(np.dot(x[:-1], y[1:]) - np.dot(x[1:], y[:-1]))


def read_rings(shp: bytes, shx: bytes, name: str) -> list[list[np.ndarray]]:
    """The rings of each record of a polygon shapefile, in record order (an empty list for a null record)."""
    if len(shp) < 100 or len(shx) < 100:
        raise GeoShapesError(f"{name}: shorter than a shapefile header")
    for label, data in (("shp", shp), ("shx", shx)):
        code, length = struct.unpack(">i20xi", data[:28])
        version, stype = struct.unpack("<ii", data[28:36])
        if code != 9994 or length * 2 != len(data) or version != 1000 or stype != 5:
            raise GeoShapesError(
                f"{name}.{label}: header (code {code}, length {length * 2}, version {version}, type {stype}) is not a "
                f"polygon shapefile of {len(data):,} bytes"
            )
    if (len(shx) - 100) % 8:
        raise GeoShapesError(f"{name}.shx: {len(shx)} bytes is not a header plus 8-byte entries")
    out: list[list[np.ndarray]] = []
    for i in range((len(shx) - 100) // 8):
        offset, content = struct.unpack(">ii", shx[100 + 8 * i : 108 + 8 * i])
        offset *= 2
        number, content2 = struct.unpack(">ii", shp[offset : offset + 8])
        if number != i + 1 or content2 != content:
            raise GeoShapesError(f"{name}.shp: record {i + 1} at byte {offset} does not match the .shx")
        (stype,) = struct.unpack("<i", shp[offset + 8 : offset + 12])
        if stype == 0:
            out.append([])
            continue
        if stype != 5:
            raise GeoShapesError(f"{name}.shp: record {i + 1} has shape type {stype}, not a polygon")
        nparts, npoints = struct.unpack("<ii", shp[offset + 44 : offset + 52])
        parts = np.frombuffer(shp, dtype="<i4", count=nparts, offset=offset + 52)
        xy = np.frombuffer(shp, dtype="<f8", count=2 * npoints, offset=offset + 52 + 4 * nparts).reshape(-1, 2)
        bounds = [*parts.tolist(), npoints]
        rings = [xy[bounds[k] : bounds[k + 1]] for k in range(nparts)]
        for k, ring in enumerate(rings):
            if len(ring) < 4 or not np.array_equal(ring[0], ring[-1]):
                raise GeoShapesError(f"{name}.shp: record {i + 1} ring {k} is not a closed ring of 4 or more points")
        out.append(rings)
    return out


def assemble(rings: list[np.ndarray]) -> tuple[tuple[Polygon, ...], int, int]:
    """Polygons from a record's rings (clockwise outer rings, counter-clockwise holes); how many holes lay in no outer
    ring and were made outer rings; how many rings had no area (their points on one line), which enclose nothing and
    are left out."""
    shells: list[np.ndarray] = []
    holes: list[np.ndarray] = []
    flat = 0
    for ring in rings:
        area = _signed_area(ring)
        if area == 0:
            flat += 1
            continue
        (shells if area < 0 else holes).append(ring)
    shell_polys = [Polygon(s) for s in shells]
    tree = shapely.STRtree(shell_polys) if shell_polys else None
    owned: list[list[np.ndarray]] = [[] for _ in shells]
    orphans: list[np.ndarray] = []
    for hole in holes:
        inside = Polygon(hole).representative_point()
        hits = [] if tree is None else [int(k) for k in tree.query(inside, predicate="within")]
        if not hits:
            orphans.append(hole)
            continue
        owned[min(hits, key=lambda k: (shell_polys[k].area, k))].append(hole)
    polys = [Polygon(s, owned[k]) for k, s in enumerate(shells)] + [Polygon(h) for h in orphans]
    return tuple(polys), len(orphans), flat


def read_layer(files: dict[str, bytes], layer: str) -> list[Record]:
    name = f"WB_GAD_{layer.upper()}"
    rings = read_rings(files[f"{SOURCE}/{layer}-shp"], files[f"{SOURCE}/{layer}-shx"], name)
    dbf = shapefile.Reader(dbf=io.BytesIO(files[f"{SOURCE}/{layer}-dbf"]), encoding="utf-8")
    attrs = [r.as_dict() for r in dbf.records()]
    if len(attrs) != len(rings):
        raise GeoShapesError(f"{name}: {len(attrs)} records in the .dbf, {len(rings)} in the .shp")
    out = []
    for fid, (a, rs) in enumerate(zip(attrs, rings, strict=True)):
        polys, orphans, flat = assemble(rs)
        out.append(Record(fid, a, polys, orphans, flat))
    return out


def check_projection(prj: bytes, cpg: bytes) -> None:
    text = prj.decode("ascii").strip()
    if text != WGS84_PRJ:
        raise GeoShapesError(f"the .prj is not WGS 84 longitude and latitude as expected: {text[:200]!r}")
    if cpg.decode("ascii").strip() != CODE_PAGE:
        raise GeoShapesError(f"the .cpg is {cpg[:40]!r}, not {CODE_PAGE}")


def read_cckp_names(data: bytes) -> dict[str, str]:
    """CCKP's code -> name, from the code and name columns of its xlsx (its value column is not read)."""
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True)
    rows = list(wb.worksheets[0].iter_rows(values_only=True))
    if not rows or tuple(rows[0][:2]) != ("code", "name"):
        first = rows[0][:2] if rows else None
        raise GeoShapesError(f"subnational-names: the first columns are {first}, not code, name")
    names: dict[str, str] = {}
    for r in rows[1:]:
        code, name = r[0], r[1]
        if not isinstance(code, str) or not isinstance(name, str) or not name.strip():
            raise GeoShapesError(f"subnational-names: row {r[:2]!r} has no code or no name")
        if code in names:
            raise GeoShapesError(f"subnational-names: {code} is listed twice")
        names[code] = name
    return names


# --- matching --------------------------------------------------------------------------------------------------------


def norm(name: str) -> str:
    """The form names are compared in: Unicode compatibility decomposition with accents dropped, case folded,
    hyphens, dashes and apostrophes read as spaces, runs of spaces as one. Nothing else (no fuzzy matching)."""
    s = unicodedata.normalize("NFKD", name)
    s = "".join(ch for ch in s if not unicodedata.combining(ch)).casefold()
    s = re.sub(r"[-‐‑‒–—―'’`]", " ", s)
    return " ".join(s.split())


def is_name(code: str, name: str) -> bool:
    return name not in NOT_NAMES and name != code


def shape_names(r: Record) -> dict[str, str]:
    """The shape's names in compared form -> the first field holding each."""
    out: dict[str, str] = {}
    for f in NAME_FIELDS:
        v = r.attrs.get(f) or ""
        if isinstance(v, str) and v.strip():
            out.setdefault(norm(v), f)
    return out


def gaul_key(r: Record) -> str | None:
    g0, g1 = r.attrs.get("GAUL_0") or 0, r.attrs.get("GAUL_1") or 0
    return f"{r.iso}.{g0}{g1}" if g0 > 0 and g1 > 0 else None


ADM1CD_T = re.compile(r"[A-Z]{3}\d+")


def adm1cd_key(r: Record) -> str | None:
    g0, t = r.attrs.get("GAUL_0") or 0, r.attrs.get("ADM1CD_t") or ""
    return f"{r.iso}.{g0}{int(t[3:])}" if g0 > 0 and ADM1CD_T.fullmatch(t) else None


@dataclass(frozen=True)
class Declared:
    cckp_code: str
    record: int
    adm1cd_t: str
    name: str
    decision: str
    reason: str


def read_declared(path: Path) -> list[Declared]:
    with path.open(encoding="utf-8", newline="") as f:
        reader = csv.reader(f)
        header = next(reader)
        if tuple(header) != DECLARED_COLUMNS:
            raise GeoShapesError(f"{DECLARED}: columns {header} != {list(DECLARED_COLUMNS)}")
        out = []
        for row in reader:
            d = Declared(row[0], int(row[1]), row[2], row[3], row[4], row[5])
            if d.decision not in {"accept", "reject"} or not d.reason.strip():
                raise GeoShapesError(f"{DECLARED}: {d.cckp_code} record {d.record} needs accept or reject and a reason")
            out.append(d)
    pairs = Counter((d.cckp_code, d.record) for d in out)
    twice = sorted(p for p, n in pairs.items() if n > 1)
    if twice:
        raise GeoShapesError(f"{DECLARED}: pairs declared twice: {twice}")
    return out


@dataclass(frozen=True)
class Pair:
    """One row of the crosswalk: a CCKP region, a shape, or both."""

    country: str
    cckp_code: str | None
    cckp_name: str | None
    record: Record | None
    method: str
    names_equal: bool | None


@dataclass
class Join:
    pairs: list[Pair]
    countries: list[GeoCountry]
    regions_of: dict[str, list[tuple[str, list[Record], str]]]
    """entity -> [(feature id, its records, method or "none")] for its map file."""
    disputed_left_out: int


def join(names: dict[str, str], adm1: list[Record], declared: list[Declared], entity_of: dict[str, str]) -> Join:
    """Match CCKP's regions to the admin-1 shapes (see the module docstring); raises GeoShapesError on any stop."""
    regions = {c: n for c, n in names.items() if "." in c}
    disputed = sorted(c for c in regions if c.startswith("DIS."))
    if len(disputed) != DISPUTED_UNITS:
        raise GeoShapesError(f"CCKP lists {len(disputed)} DIS.* units, not {DISPUTED_UNITS}: {disputed[:5]}…")
    for c in disputed:
        del regions[c]
    by_iso: dict[str, list[Record]] = defaultdict(list)
    for r in adm1:
        if not r.iso:
            raise GeoShapesError(f"WB_GAD_ADM1 record {r.fid} has no ISO_A3")
        by_iso[r.iso].append(r)
    cckp_by_iso: dict[str, list[str]] = defaultdict(list)
    for c in sorted(regions):
        cckp_by_iso[c.split(".")[0]].append(c)
    for iso in sorted(set(by_iso) | set(cckp_by_iso)):
        if iso not in entity_of:
            raise GeoShapesError(f"{iso} (a CCKP region prefix or a WB_GAD_ADM1 ISO_A3) has no entity")
    records = {r.fid: r for r in adm1}
    gaul: dict[str, list[Record]] = defaultdict(list)
    adm1cd: dict[str, list[Record]] = defaultdict(list)
    for r in adm1:
        if (k := gaul_key(r)) is not None:
            gaul[k].append(r)
        if (k := adm1cd_key(r)) is not None:
            adm1cd[k].append(r)
    names_of = {r.fid: shape_names(r) for r in adm1}

    accept: dict[str, list[Declared]] = defaultdict(list)
    reject: dict[tuple[str, int], Declared] = {}
    for d in declared:
        r = records.get(d.record)
        if r is None or d.cckp_code not in regions:
            raise GeoShapesError(f"{DECLARED}: {d.cckp_code} or record {d.record} is no longer in the files")
        if (r.attrs.get("ADM1CD_t") or "", (r.attrs.get("NAM_1") or "").strip()) != (d.adm1cd_t, d.name):
            raise GeoShapesError(
                f"{DECLARED}: record {d.record} is now {r.attrs.get('ADM1CD_t')!r} {r.attrs.get('NAM_1')!r}, not "
                f"{d.adm1cd_t!r} {d.name!r}"
            )
        if d.cckp_code.split(".")[0] != r.iso:
            raise GeoShapesError(f"{DECLARED}: {d.cckp_code} and record {d.record} ({r.iso}) are in two countries")
        if d.decision == "accept":
            accept[d.cckp_code].append(d)
        else:
            reject[(d.cckp_code, d.record)] = d
    used_rejects: set[tuple[str, int]] = set()

    joined: dict[str, tuple[str, list[Record]]] = {}
    owner: dict[int, str] = {}

    def take(code: str, method: str, rs: list[Record]) -> None:
        for r in rs:
            if r.fid in owner:
                raise GeoShapesError(
                    f"WB_GAD_ADM1 record {r.fid} ({r.attrs.get('NAM_1')!r}) would be used twice: by {owner[r.fid]} and "
                    f"by {code} ({method})"
                )
            owner[r.fid] = code
        joined[code] = (method, rs)

    # 1-2. Code rules, then 3. declared accepts.
    for code in sorted(regions):
        name = regions[code]
        agreeing_by_rule: list[tuple[str, list[Record]]] = []
        for rule, index in (("gaul", gaul), ("adm1cd", adm1cd)):
            agree: list[Record] = []
            for r in index.get(code, []):
                if (code, r.fid) in reject:
                    used_rejects.add((code, r.fid))
                    continue
                if any(d.record == r.fid for d in accept.get(code, [])):
                    continue  # declared accept: joined in step 3
                if norm(name) in names_of[r.fid]:
                    agree.append(r)
                else:
                    raise GeoShapesError(
                        f"{code} {name!r} matches WB_GAD_ADM1 record {r.fid} ({r.attrs.get('ADM1CD_t')!r} "
                        f"{r.attrs.get('NAM_1')!r}) by the {rule} rule but the names differ: declare the pair in "
                        f"{DECLARED} (accept or reject, with the reason)"
                    )
            if agree:
                agreeing_by_rule.append((rule, agree))
        if len(agreeing_by_rule) == 2 and {r.fid for r in agreeing_by_rule[0][1]} != {
            r.fid for r in agreeing_by_rule[1][1]
        }:
            raise GeoShapesError(
                f"{code}: the gaul rule gives records {[r.fid for r in agreeing_by_rule[0][1]]} and the adm1cd rule "
                f"{[r.fid for r in agreeing_by_rule[1][1]]}"
            )
        if agreeing_by_rule:
            if code in accept:
                raise GeoShapesError(f"{DECLARED}: {code} is already matched by the {agreeing_by_rule[0][0]} rule")
            take(code, *agreeing_by_rule[0])
    for code in sorted(accept):
        take(code, "declared", [records[d.record] for d in accept[code]])
    stale = sorted(set(reject) - used_rejects)
    if stale:
        raise GeoShapesError(f"{DECLARED}: rejected pairs no code rule makes any more: {stale}")

    # 4. Names, where the country has as many CCKP regions as shapes.
    withheld: set[str] = set()
    for iso in sorted(cckp_by_iso):
        codes, shapes = cckp_by_iso[iso], by_iso.get(iso, [])
        if len(codes) != len(shapes):
            withheld.add(iso)
            continue
        counts = Counter(norm(regions[c]) for c in codes if is_name(c, regions[c]))
        for code in codes:
            if code in joined or not is_name(code, regions[code]):
                continue
            n = norm(regions[code])
            hits = [r for r in shapes if n in names_of[r.fid]]
            if counts[n] == 1 and len(hits) == 1 and hits[0].fid not in owner:
                take(code, "name", hits)

    # The crosswalk, the map features and the per-country summary.
    pairs: list[Pair] = []
    regions_of: dict[str, list[tuple[str, list[Record], str]]] = defaultdict(list)
    countries: list[GeoCountry] = []
    for iso in sorted(set(by_iso) | set(cckp_by_iso)):
        entity = entity_of[iso]
        codes, shapes = cckp_by_iso.get(iso, []), by_iso.get(iso, [])
        methods: Counter[str] = Counter()
        unjoined: list[str] = []
        shapeless: list[str] = []
        for code in codes:
            name = regions[code]
            if code in joined:
                method, rs = joined[code]
                methods[method] += 1
                if any(r.polygons for r in rs):
                    regions_of[entity].append((code, rs, method))
                else:
                    shapeless.append(code)
                for r in rs:
                    equal = norm(name) in names_of[r.fid]
                    pairs.append(Pair(entity, code, name, r, method, equal))
            else:
                unjoined.append(code)
                pairs.append(Pair(entity, code, name, None, "none", None))
        unmatched = [r for r in shapes if r.fid not in owner]
        for r in unmatched:
            if r.polygons:  # a record without rings has nothing to draw (counted in the processing steps)
                regions_of[entity].append((f"V1:{r.fid}", [r], "none"))
            pairs.append(Pair(entity, None, None, r, "none", None))
        countries.append(
            GeoCountry(
                entity=entity,
                cckp_regions=len(codes),
                shapes=len(shapes),
                joined={m: methods[m] for m in ("gaul", "adm1cd", "declared", "name") if methods[m]},
                name_rule="withheld" if iso in withheld else "applied",
                unjoined=sorted(unjoined),
                shapeless=sorted(shapeless),
                unmatched=sorted(f"V1:{r.fid}" for r in unmatched if r.polygons),
                file=None,
            )
        )
    pairs.sort(key=lambda p: (p.country, p.cckp_code or "~", p.record.fid if p.record else -1))
    return Join(pairs, countries, dict(regions_of), len(disputed))


# --- simplifying -----------------------------------------------------------------------------------------------------


def representative_point(polys: list[Polygon]) -> list[float] | None:
    """A point inside the full, unsimplified shape (inside its valid form when the rings cross themselves)."""
    if not polys:
        return None
    g = shapely.make_valid(MultiPolygon(polys))
    p = g.representative_point()
    return [round(p.x, RP_DECIMALS), round(p.y, RP_DECIMALS)]


def extent(polys: list[Polygon]) -> tuple[float, float, float]:
    """(longest side in degrees, measured the short way round the 180° line; width of the box; height of the box)."""
    b = np.array([p.bounds for p in polys])
    lat = float(b[:, 3].max() - b[:, 1].min())
    spans = sorted((float(x0), float(x1)) for x0, _, x1, _ in b)
    covered: list[list[float]] = []
    for x0, x1 in spans:
        if covered and x0 <= covered[-1][1]:
            covered[-1][1] = max(covered[-1][1], x1)
        else:
            covered.append([x0, x1])
    gaps = [covered[k + 1][0] - covered[k][1] for k in range(len(covered) - 1)]
    gaps.append(covered[0][0] + 360 - covered[-1][1])
    lon = 360 - max(gaps)
    box_width = covered[-1][1] - covered[0][0]
    return max(lon, lat), box_width, lat


def country_simplify(polys: list[Polygon]) -> Simplify:
    side, width, height = extent(polys)
    steps = max(COUNTRY_QUANTIZE, math.ceil(COUNTRY_QUANTIZE * max(width, height) / side))
    return Simplify(
        prequantize=25 * steps,
        epsilon=side * COUNTRY_EPSILON,
        min_area=(side * COUNTRY_MIN_SIDE) ** 2,
        quantize=steps,
    )


def keep(polys: list[Polygon], min_area: float) -> list[Polygon]:
    """The polygons and holes at least min_area square degrees in size."""
    out = []
    for p in polys:
        if p.area >= min_area:
            out.append(Polygon(p.exterior, [h for h in p.interiors if Polygon(h).area >= min_area]))
    return out


@dataclass
class Feature:
    id: str
    polygons: list[Polygon]
    properties: dict[str, Any] = field(default_factory=dict)


def topology(objects: dict[str, list[Feature]], s: Simplify) -> dict:
    """A TopoJSON topology of the features, one object per key (features sorted by id), simplified by s, the objects
    sharing their borders. A feature with nothing left at this scale is kept as a null geometry with its rp."""
    drawn: list[dict] = []
    out: dict[str, list[dict]] = {name: [] for name in objects}
    for name, feats in objects.items():
        for f in feats:
            rp = representative_point(f.polygons)
            if rp is None:
                raise GeoShapesError(f"{name} {f.id}: no shape at all in the file")
            props = {**f.properties, "rp": rp}
            kept = keep(f.polygons, s.min_area)
            if kept:
                geom = MultiPolygon(kept).__geo_interface__
                drawn.append({"type": "Feature", "id": f.id, "properties": {**props, OBJECT: name}, "geometry": geom})
            else:
                out[name].append({"type": None, "id": f.id, "properties": props})
    # One collection for every object (the library's several-object mode needs data frames), split again after.
    topo = topojson.Topology(
        {"type": "FeatureCollection", "features": drawn},
        object_name="all",
        prequantize=s.prequantize,
        toposimplify=s.epsilon,
        topoquantize=s.quantize,
        prevent_oversimplify=True,
    ).to_dict()
    for g in topo["objects"]["all"]["geometries"]:
        out[g["properties"].pop(OBJECT)].append(g)
    topo["objects"] = {
        name: {"type": "GeometryCollection", "geometries": sorted(geoms, key=lambda g: g["id"])}
        for name, geoms in out.items()
    }
    return topo


def check_topology(topo: dict, expected: dict[str, list[str]], name: str) -> dict[str, tuple[int, int]]:
    """Every expected feature present (by id), every arc index valid, every ring closed. Returns, per object,
    (features, features with at least one ring)."""
    arcs = topo["arcs"]
    decoded = []
    for a in arcs:
        if len(a) < 2:
            raise GeoShapesError(f"{name}: an arc of {len(a)} point(s)")
        pts = np.cumsum(np.array(a, dtype=np.int64), axis=0)
        decoded.append((tuple(pts[0]), tuple(pts[-1])))

    def ends(i: int) -> tuple[tuple, tuple]:
        k = i if i >= 0 else ~i
        if k >= len(arcs):
            raise GeoShapesError(f"{name}: arc index {i} out of range")
        a, b = decoded[k]
        return (a, b) if i >= 0 else (b, a)

    out = {}
    for obj, ids in expected.items():
        geoms = topo["objects"][obj]["geometries"]
        got = [g["id"] for g in geoms]
        if got != sorted(ids):
            missing = sorted(set(ids) - set(got))
            extra = sorted(set(got) - set(ids))
            raise GeoShapesError(f"{name} {obj}: features lost {missing[:10]}, unexpected {extra[:10]}")
        drawn = 0
        for g in geoms:
            if g["type"] is None:
                continue
            polys = g["arcs"] if g["type"] == "MultiPolygon" else [g["arcs"]]
            if g["type"] not in {"Polygon", "MultiPolygon"} or not polys:
                raise GeoShapesError(f"{name} {obj} {g['id']}: geometry {g['type']} without rings")
            for poly in polys:
                for ring in poly:
                    start = ends(ring[0])[0]
                    cur = start
                    for i in ring:
                        a, b = ends(i)
                        if a != cur:
                            raise GeoShapesError(f"{name} {obj} {g['id']}: arcs of a ring do not join")
                        cur = b
                    if cur != start:
                        raise GeoShapesError(f"{name} {obj} {g['id']}: a ring does not close")
            drawn += 1
        out[obj] = (len(geoms), drawn)
    return out


def topo_bytes(topo: dict, meta: dict[str, str]) -> bytes:
    """Compact, key-sorted JSON (deterministic for the same topology), with a credit block."""
    doc = {**topo, "meta": meta}
    return (json.dumps(doc, ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False) + "\n").encode()


# --- building --------------------------------------------------------------------------------------------------------


def geo_dir(paths: Paths) -> Path:
    return paths.data / "v1" / "geo"


def declared_path(paths: Paths) -> Path:
    return paths.repo / DECLARED


def record_path(paths: Paths) -> Path:
    return paths.builds / "geo-shapes.json"


@dataclass(frozen=True)
class Inputs:
    files: dict[str, bytes]
    snaps: dict[str, Snapshot]


def load_inputs(paths: Paths) -> Inputs:
    pointers = snapshots.read_current(paths)
    files, snaps = {}, {}
    for key in INPUT_KEYS:
        sha = pointers.get(key)
        if sha is None:
            raise GeoShapesError(f"no snapshot of {key}; run envdash fetch -s {key.split('/')[0]}")
        snap = snapshots.read_manifest(paths, sha)
        if snap is None:
            raise GeoShapesError(f"current snapshot {sha} of {key} has no manifest")
        p = snapshots.cache_path(paths, sha)
        if not p.exists():
            raise GeoShapesError(f"raw bytes {sha} of {key} are not in the local cache; run envdash snapshots pull")
        data = p.read_bytes()
        if canonical.sha256_bytes(data) != sha:
            raise GeoShapesError(f"cached bytes of {key} do not hash to {sha}")
        files[key], snaps[key] = data, snap
    return Inputs(files, snaps)


def build_key(paths: Paths, registry: Registry) -> tuple[str, dict]:
    pointers = snapshots.read_current(paths)
    parts = {
        "inputs": {
            k: {
                "sha256": pointers.get(k),
                "manifest_sha256": canonical.sha256_file(snapshots.manifest_path(paths, pointers[k]))
                if k in pointers and snapshots.manifest_path(paths, pointers[k]).exists()
                else None,
            }
            for k in INPUT_KEYS
        },
        "module_sha256": canonical.sha256_file(Path(__file__)),
        "declared_sha256": canonical.sha256_file(declared_path(paths)),
        "entities_sha256": canonical.sha256_file(geo.entities_path(paths)),
        "contract_sha256": canonical.sha256_file(Path(models_mod.__file__)),
        "lock_sha256": lock_sha(paths),
        "sources": {sid: canonical.sha256_file(registry.source_files[sid]) for sid in (CCKP, SOURCE)},
    }
    return canonical.sha256_bytes(canonical.dump_bytes(parts)), parts


def _origin(src: Source, snap: Snapshot, vintage: str) -> Origin:
    return Origin(
        source_id=src.id,
        artifact_id=snap.artifact_id,
        producer=src.publisher,
        title=src.title,
        version_producer=vintage,
        citation_full=render(src.citation.text, version=vintage, date_accessed=snap.date_accessed, year=None),
        url_main=src.landing_url,
        url_download=snap.url,
        date_published=src.citation.issued if src.id == SOURCE else None,
        date_accessed=snap.date_accessed,
        licence=src.licence,
        sha256=snap.sha256,
        bytes=snap.bytes,
        etag=snap.etag,
        last_modified=snap.last_modified,
        wayback_url=snap.wayback.url if snap.wayback and snap.wayback.status == "captured" else None,
        r2_url=f"{PUBLIC_FILES_BASE}{snap.r2_key}" if snap.r2_bucket == "envdash-public" and snap.r2_key else None,
        doi=src.citation.DOI,
        acquisition=snap.acquisition,
    )


CHANGES = (
    "boundaries simplified for a small screen (shared borders kept shared), islands and lakes too small to draw left "
    "out, and CCKP's region codes matched to the World Bank's shapes"
)


@dataclass
class Built:
    files: dict[str, bytes]
    """Path under data/ -> bytes."""
    index: GeoIndex
    join: Join


def entity_of_codes(paths: Paths, codes: set[str]) -> dict[str, str]:
    """Our entity for each World Bank ISO_A3 (CCKP's region prefixes are the same codes); stops on an unknown one."""
    table = geo.read_table(geo.entities_path(paths))
    out = {}
    for iso in sorted(codes):
        try:
            out[iso] = geo.resolve(iso, SOURCE, entities=table)
        except geo.UnknownEntity as e:
            raise GeoShapesError(f"World Bank code {iso!r} has no entity: {e}") from None
    return out


@dataclass(frozen=True)
class Matching:
    """What the join reads: CCKP's names, the admin-1 records, the declared pairs and the entity of each code."""

    names: dict[str, str]
    adm1: list[Record]
    declared: list[Declared]
    entity_of: dict[str, str]

    def join(self) -> Join:
        return join(self.names, self.adm1, self.declared, self.entity_of)


def matching(paths: Paths, inputs: Inputs, layers: dict[str, list[Record]]) -> Matching:
    names = read_cckp_names(inputs.files[f"{CCKP}/{NAMES_ARTIFACT}"])
    codes = {r.iso for recs in layers.values() for r in recs if r.iso}
    codes |= {c.split(".")[0] for c in names if "." in c and not c.startswith("DIS.")}
    return Matching(names, layers["adm1"], read_declared(declared_path(paths)), entity_of_codes(paths, codes))


def read_layers(inputs: Inputs, layers: tuple[str, ...] = LAYERS) -> dict[str, list[Record]]:
    check_projection(inputs.files[f"{SOURCE}/prj"], inputs.files[f"{SOURCE}/cpg"])
    return {layer: read_layer(inputs.files, layer) for layer in layers}


def build(paths: Paths, registry: Registry, inputs: Inputs | None = None, *, only: set[str] | None = None) -> Built:
    """Every output file's bytes (nothing written). With `only`, just the world file (if "world" is in it) and those
    entities' country files, for tests."""
    for sid in (SOURCE, CCKP):
        if sid not in registry.sources:
            raise GeoShapesError(registry.source_errors.get(sid, f"{sid} is not registered"))
    inputs = inputs or load_inputs(paths)
    layers = read_layers(inputs)
    m = matching(paths, inputs, layers)
    entity_of = m.entity_of

    def entity(iso: str) -> str:
        return entity_of[iso]

    j = m.join()

    sources = {sid: registry.sources[sid] for sid in (CCKP, SOURCE)}
    cls = strictest([s.licence_class for s in sources.values()])
    licences = {s.licence for s in sources.values() if s.licence_class == cls}
    if len(licences) != 1:
        raise GeoShapesError(f"{CCKP} and {SOURCE} carry different licences at class {cls}; declare the combined terms")
    licence = licences.pop()
    gad, cckp = sources[SOURCE].obligations, sources[CCKP].obligations
    accessed = max(s.date_accessed for k, s in inputs.snaps.items() if k.startswith(SOURCE))
    gad_line = render(gad.attribution_modified or gad.attribution, version=VINTAGE, date_accessed=accessed, year=None)
    attribution = f"{gad_line} Changes: {CHANGES}. Region codes and names: {cckp.attribution}"
    notices = [n for n in (gad.notice, cckp.notice) if n]
    notice = " ".join(dict.fromkeys(notices)) or None
    meta = {
        "credit": attribution,
        "licence": f"{licence.name} <{licence.url}>",
        "about": "https://environmentdashboard.org/data/v1/geo/index.json",
    }
    if notice:
        meta["notice"] = notice

    out: dict[str, bytes] = {}
    files: list[GeoFile] = []

    def add(path: str, data: bytes, counts: dict[str, tuple[int, int]], main: str, limit: int) -> GeoFile:
        if len(data) > limit:
            raise GeoShapesError(f"{path} is {len(data):,} bytes, over its limit of {limit:,}")
        gf = GeoFile(
            path=path,
            sha256=canonical.sha256_bytes(data),
            bytes=len(data),
            features=counts[main][0],
            drawn=counts[main][1],
            disputed=counts["disputed"][0],
        )
        out[path] = data
        files.append(gf)
        return gf

    ndlsa = []
    for r in layers["ndlsa"]:
        props: dict[str, Any] = {"name": str(r.attrs["NAM_0"]).strip()}
        if r.iso:
            props["e"] = entity(r.iso)
        ndlsa.append(Feature(f"NDLSA:{r.fid}", list(r.polygons), props))

    if only is None or "world" in only:
        by_entity: dict[str, list[Polygon]] = defaultdict(list)
        for r in layers["adm0"]:
            by_entity[entity(r.iso)].extend(r.polygons)
        countries = [Feature(e, polys) for e, polys in sorted(by_entity.items())]
        topo = topology({"countries": countries, "disputed": ndlsa}, WORLD)
        counts = check_topology(
            topo, {"countries": [c.id for c in countries], "disputed": [d.id for d in ndlsa]}, "world.topo.json"
        )
        add("v1/geo/world.topo.json", topo_bytes(topo, meta), counts, "countries", WORLD_LIMIT)

    ndlsa_tree = shapely.STRtree([MultiPolygon(d.polygons) for d in ndlsa])
    summaries = []
    for c in j.countries:
        feats = j.regions_of.get(c.entity, [])
        joined_any = any(m != "none" for _, _, m in feats)
        if not joined_any or (only is not None and c.entity not in only):
            summaries.append(c)
            continue
        regions = []
        for fid, rs, method in feats:
            props: dict[str, Any] = {"m": method}
            if method == "none":
                props["name"] = str(rs[0].attrs.get("NAM_1") or "").strip()
            regions.append(Feature(fid, [p for r in rs for p in r.polygons], props))
        regions.sort(key=lambda x: x.id)
        allpolys = [p for x in regions for p in x.polygons]
        whole = MultiPolygon(allpolys)
        near = sorted(int(k) for k in ndlsa_tree.query(whole, predicate="dwithin", distance=TOUCH_DEGREES))
        disputed = [ndlsa[k] for k in near]
        path = f"v1/geo/adm1/{c.entity}.topo.json"
        topo = topology({"regions": regions, "disputed": disputed}, country_simplify(allpolys))
        counts = check_topology(topo, {"regions": [x.id for x in regions], "disputed": [d.id for d in disputed]}, path)
        add(path, topo_bytes(topo, meta), counts, "regions", COUNTRY_LIMIT)
        summaries.append(c.model_copy(update={"file": path}))

    crosswalk = crosswalk_bytes(j, attribution, licence.name, str(licence.url))
    out["v1/geo/cckp-adm1-crosswalk.csv"] = crosswalk
    files.append(
        GeoFile(
            path="v1/geo/cckp-adm1-crosswalk.csv",
            sha256=canonical.sha256_bytes(crosswalk),
            bytes=len(crosswalk),
            features=len(j.pairs),
            drawn=0,
            disputed=0,
        )
    )
    origins = [
        _origin(sources[k.split("/")[0]], s, VINTAGE if k.startswith(SOURCE) else f"fetched {s.date_accessed}")
        for k, s in sorted(inputs.snaps.items())
    ]
    index = GeoIndex(
        description="Map shapes for the future view: the World Bank's country, state and province and disputed-area "
        "boundaries (Official Boundaries, GAD V1 of 2025-05-07), simplified, with CCKP's region codes matched to the "
        "states and provinces they describe.",
        licence_class=cls,
        licence=licence,
        attribution=attribution,
        notice=notice,
        origins=origins,
        processing=steps(paths, inputs, j, layers),
        files=sorted(files, key=lambda x: x.path),
        countries=sorted(summaries, key=lambda x: x.entity),
    )
    out["v1/geo/index.json"] = canonical.dump_bytes(index)
    return Built(out, index, j)


def crosswalk_bytes(j: Join, attribution: str, licence: str, licence_url: str) -> bytes:
    buf = io.StringIO(newline="")
    for line in (
        "CCKP region codes matched to World Bank Official Boundaries (GAD V1, 2025-05-07) admin-1 shapes",
        f"Credit: {attribution}",
        f"Licence: {licence} <{licence_url}>",
        "method: gaul or adm1cd (CCKP's code rebuilt from the shape's GAUL or ADM1CD_t code, names agreeing), declared "
        "(pipeline/geo/cckp-adm1-declared.csv), name (the same name, unique on both sides), none (not matched)",
        "v1_id: V1:<record>, the shape's record number (from 0) in WB_GAD_ADM1; a region may take several records",
        "Provenance and processing: https://environmentdashboard.org/data/v1/geo/index.json",
    ):
        buf.write(f"# {line}\n")
    w = csv.writer(buf, lineterminator="\n")
    w.writerow(CROSSWALK_COLUMNS)
    for p in j.pairs:
        r = p.record
        w.writerow(
            [
                p.country,
                p.cckp_code or "",
                p.cckp_name or "",
                f"V1:{r.fid}" if r else "",
                (r.attrs.get("ADM1CD_t") or "") if r else "",
                str(r.attrs.get("NAM_1") or "").strip() if r else "",
                p.method,
                "" if p.names_equal is None else str(p.names_equal).lower(),
            ]
        )
    return buf.getvalue().encode("utf-8")


def steps(paths: Paths, inputs: Inputs, j: Join, layers: dict[str, list[Record]]) -> list[ProcessingStep]:
    lock = lock_sha(paths)
    module = Path(__file__)
    shas = [inputs.snaps[k].sha256 for k in INPUT_KEYS]
    method_counts: Counter[str] = Counter()
    for c in j.countries:
        method_counts.update(c.joined)
    unjoined = sum(len(c.unjoined) for c in j.countries)
    unmatched = sum(len(c.unmatched) for c in j.countries)
    shapeless = sorted(code for c in j.countries for code in c.shapeless)
    withheld = sorted(c.entity for c in j.countries if c.name_rule == "withheld" and c.cckp_regions and c.shapes)

    def per_layer(count) -> str:
        return ", ".join(f"{layer.upper()} {count(recs)}" for layer, recs in layers.items())

    empty = "; ".join(
        f"{layer.upper()} records {', '.join(str(r.fid) for r in recs if not r.polygons)}"
        for layer, recs in layers.items()
        if any(not r.polygons for r in recs)
    )
    texts = [
        "Read the ADM0, ADM1 and NDLSA layers of the World Bank's GAD V1 shapefile as written (WGS 84 longitude and "
        "latitude, checked against the .prj; UTF-8 attributes, checked against the .cpg). Clockwise rings are outer "
        "rings and counter-clockwise rings holes, each hole in the smallest outer ring containing it. Holes inside no "
        f"outer ring, drawn as outer rings: {per_layer(lambda rs: sum(r.orphan_holes for r in rs))}. Rings with no "
        f"area (their points on one line), left out: {per_layer(lambda rs: sum(r.flat_rings for r in rs))}. Records "
        f"with no rings, so no shape: {empty or 'none'}.",
        "Countries: every ADM0 record is joined to the site's entity of its ISO_A3 code (Kosovo's KSV is KOS); the "
        "records of one entity form one multipolygon, nothing merged. Disputed areas (NDLSA) keep the World Bank's "
        "names, leading and trailing spaces removed, and carry no value.",
        f"CCKP's regions were matched to ADM1 records country by country: {method_counts['gaul']:,} by the GAUL rule "
        f"({{ISO_A3}}.{{GAUL_0}}{{GAUL_1}}), {method_counts['adm1cd']:,} by the ADM1CD rule "
        f"({{ISO_A3}}.{{GAUL_0}}{{int(ADM1CD_t[3:])}}), each only where the names agree; {method_counts['declared']:,} "
        f"by pairs declared in {DECLARED}; {method_counts['name']:,} by a name unique on both sides, only in countries "
        f"with as many CCKP regions as records. {unjoined:,} CCKP regions are listed without a shape, "
        f"{len(shapeless)} matched to a record with no shape ({', '.join(shapeless) or 'none'}), and {unmatched:,} "
        f"records are drawn without a value. CCKP's {j.disputed_left_out} DIS.* units are left out. The name rule was "
        f"withheld, the counts differing, in: {', '.join(withheld)}.",
        f"Simplified with the topojson library (topology with shared borders, Douglas-Peucker on arcs, rings kept at "
        f"least a triangle). World: coordinates snapped to {WORLD.prequantize:,} steps, epsilon {WORLD.epsilon}°, "
        f"polygons and holes under {WORLD.min_area} square degrees left out, output quantised to {WORLD.quantize:,} "
        f"steps. Each country: the same with epsilon its extent / {round(1 / COUNTRY_EPSILON)}, minimum area "
        f"(its extent / {round(1 / COUNTRY_MIN_SIDE)})², and {COUNTRY_QUANTIZE:,} steps across its extent. Each "
        f"feature's rp is a representative point of its full, unsimplified shape ({RP_DECIMALS} decimals); a "
        "feature with nothing left at its scale is a null geometry with that point. A country's file holds the "
        f"disputed areas within {TOUCH_DEGREES}° of its shapes.",
    ]
    return [
        ProcessingStep(
            script=paths.rel(module),
            transform_sha256=canonical.sha256_file(module),
            lock_sha256=lock,
            inputs=shas,
            description=t,
        )
        for t in texts
    ]


@dataclass
class Outcome:
    state: str  # built | skipped | failed
    reason: str | None = None
    source_ids: tuple[str, ...] = (SOURCE, CCKP)
    index: GeoIndex | None = None


def write(paths: Paths, registry: Registry, *, force: bool = False) -> Outcome:
    """Build and write data/v1/geo/ when the build key changed (or with force); keep the previous files on failure.
    With an unchanged key, the new bytes must equal the recorded ones."""
    if SOURCE not in registry.sources and SOURCE not in registry.source_errors:
        return Outcome("skipped", f"{SOURCE} is not in this registry")
    try:
        key, parts = build_key(paths, registry)
    except (FileNotFoundError, KeyError) as e:
        return Outcome("failed", f"{type(e).__name__}: {e}")
    rec_path = record_path(paths)
    rec = json.loads(rec_path.read_text(encoding="utf-8")) if rec_path.exists() else None
    base = paths.data
    current = rec is not None and all(
        (base / p).exists() and canonical.sha256_file(base / p) == sha for p, sha in rec.get("outputs", {}).items()
    )
    if not force and rec is not None and rec.get("key") == key and current:
        return Outcome("skipped")
    try:
        built = build(paths, registry)
    except GeoShapesError as e:
        return Outcome("failed", str(e))
    except Exception as e:  # isolation, as for a transform: a bug here never stops the indicators' export
        return Outcome("failed", f"{type(e).__name__}: {e}")
    new = {p: canonical.sha256_bytes(b) for p, b in built.files.items()}
    if rec is not None and rec.get("key") == key and rec.get("outputs") != new:
        changed = sorted(p for p in set(new) | set(rec.get("outputs", {})) if new.get(p) != rec["outputs"].get(p))
        return Outcome("failed", f"the same inputs gave different bytes than the last build: {changed[:10]}")
    for p, data in built.files.items():
        canonical.write_if_changed(base / p, data)
    adm1 = geo_dir(paths) / "adm1"
    for p in sorted(adm1.glob("*.topo.json")) if adm1.exists() else []:
        if p.relative_to(base).as_posix() not in built.files:
            p.unlink()
    canonical.write_if_changed(
        rec_path, canonical.dump_bytes({"key": key, "key_parts": parts, "outputs": dict(sorted(new.items()))})
    )
    return Outcome("built", index=built.index)
