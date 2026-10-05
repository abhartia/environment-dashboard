"""Climate Action Tracker: projected warming in 2100 under CAT's four thermometer scenarios, verbatim.

Input: the "CAT Thermometer" sheet of CAT_2025-11_PublicData_GlobalTemperatureEstimates_COP30.xlsx (with its "Info"
sheet for the publication date, unit, scenario descriptions and copyright notice).

Licence. CAT's terms are "All rights reserved" with permission to view, download, print and distribute, credited,
with the copyright notice and not for commercial use; nothing permits adaptation (class no-derivatives, see
pipeline/sources/cat-2025-thermometer.yaml). So nothing is computed here: every published number is one cell of the
sheet, unchanged. The export goes to data-private only (the build routes it by class).

Which rows. The sheet says "Values in bold are used in the CAT Thermometer." The bold value rows are Policies & action
(the "Combined" row), 2030 & 2035 Targets only, Pledges & targets and Optimistic scenario (net-zero pledges); those
four are published. The non-bold rows are not: the High and Low variants of Policies & action, the superseded
"2030 targets only" row, and an unlabelled row between Pledges & targets and Optimistic. The expected layout (labels,
cells, which rows are bold) is written out below; if any of it differs, the transform stops and a person must re-read
the file. It never searches for a row that moved.

Rounding. The sheet says "if you are reproducing the numerical values, please adhere to the displayed rounded values
in the tables". Every published cell must have the display format "+0.0" and a stored value with no more than one
decimal, so the number published is exactly the number the table displays. Otherwise the transform stops.

Ranges. Lower and upper are CAT's "Lower bound" and "Upper bound" columns of the MAGICC7 distribution, published as an
interval of type "range". The sheet tags the median column "_50" and the upper column "_83" and gives no tag for the
lower column, so no probability is attached to the range here.

Vintage. The Info sheet's "Date Published" (13 November 2025); CAT labels the release by its month ("November 2025").
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal
from pathlib import Path

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

from envdash.models import Dimension, DimensionValue, Display, Observation, Scope, Unit
from envdash.paths import Paths
from envdash.transform import Input, InputFile, PublisherCheck, Result, Spec, Transform, Validation

SOURCE = "cat-2025-thermometer"
WORKBOOK = Input(SOURCE, "global-temperature-estimates")
THERMOMETER = "CAT Thermometer"
INFO = "Info"
VINTAGE = "November 2025"


@dataclass(frozen=True)
class Scenario:
    id: str
    row: int
    """Row of the CAT Thermometer sheet holding Lower bound, Median and Upper bound in columns D, E, F."""
    label: str
    """Column B of that row, exactly."""
    variant: str | None
    """Column C of that row, exactly (None when empty)."""
    info_row: int
    """Row of the Info sheet whose column B names the scenario and column C explains it."""
    info_label: str


SCENARIOS: tuple[Scenario, ...] = (
    Scenario("policies-action", 15, "Policies & action", "Combined", 26, "Policies & action"),
    Scenario("targets-2030-2035", 19, "2030 & 2035 Targets only", None, 28, "2030 & 2035 targets only"),
    Scenario("pledges-targets", 20, "Pledges & targets", None, 29, "Pledges and Targets"),
    Scenario(
        "optimistic",
        22,
        "Optimistic scenario (net-zero pledges)",
        None,
        30,
        "Optimistic scenario (net-zero pledges)",
    ),
)
# Value rows of the table that are not in bold, with their column B and C labels as printed.
NOT_PUBLISHED: dict[int, tuple[str | None, str | None]] = {
    16: (None, "High"),
    17: (None, "Low"),
    18: ("2030 targets only", None),
    21: (None, None),
}
VALUE_ROWS = range(15, 23)
VALUE_COLUMNS = ("D", "E", "F")  # Lower bound, Median, Upper bound
DISPLAY_FORMAT = r"\+0.0"  # positive section of the cells' number format: one decimal

THERMOMETER_CELLS: dict[str, str | None] = {
    "B7": "Climate Action Tracker: The CAT Thermometer",
    "B13": "Global Mean Temperature above pre-industrial levels in 2100",
    "D13": None,
    "E13": "_50",
    "F13": "_83",
    "B14": "Temperature in °C",
    "D14": "Lower bound",
    "E14": "Median",
    "F14": "Upper bound",
    "B23": "The CAT uses the climate model MAGICC7 for its temperature estimates.",
    "B24": "Values in bold are used in the CAT Thermometer.",
    "B27": "NOTE: We include the unrounded values to assist with the preparation of graphics, however if you are "
    "reproducing the numerical values, please adhere to the displayed rounded values in the tables.",
}
INFO_CELLS: dict[str, str] = {
    "B7": "Climate Action Tracker COP30 Global Temperature Update",
    "B8": "Temperatures Unit",
    "C8": "°C about pre-industrial levels",
    "B10": "Date Published:",
    "B32": "Copyright © 2025 by Climate Analytics and NewClimate Institute. All rights reserved.",
    "B34": "The content provided by this data file is protected by copyright. You are authorised to view, download, "
    "print and distribute the copyrighted content from this data file subject to the following condition: Any "
    "reproduction, in full or in part, must credit Climate Analytics and NewClimate Institute, include a copyright "
    "notice and must not be used for commerical purposes.",
}

DEG_C = Unit(code="degC", label="degrees Celsius", short="°C")
DIMENSION = Dimension(
    id="scenario",
    label="CAT scenario",
    values=[DimensionValue(id=s.id, label=s.label) for s in SCENARIOS],
)


class CatFormatError(ValueError):
    pass


def _text(v: object) -> str | None:
    """A cell's text with runs of whitespace made single and the ends trimmed (CAT's cells carry trailing spaces)."""
    if v is None:
        return None
    if not isinstance(v, str):
        raise CatFormatError(f"expected text, found {v!r}")
    return re.sub(r"\s+", " ", v).strip()


def _expect(ws: Worksheet, cells: dict[str, str | None]) -> None:
    for ref, want in cells.items():
        got = _text(ws[ref].value)
        if got != want:
            raise CatFormatError(
                f"{ws.title}!{ref} reads {got!r}, expected {want!r}; re-read the file before trusting this layout"
            )


def _published_value(ws: Worksheet, ref: str) -> Decimal:
    c = ws[ref]
    if c.data_type != "n" or isinstance(c.value, bool) or not isinstance(c.value, int | float):
        raise CatFormatError(f"{ws.title}!{ref} is not a number as typed in the sheet: {c.value!r}")
    shown = str(c.number_format).split(";")[0]
    if shown != DISPLAY_FORMAT:
        raise CatFormatError(
            f"{ws.title}!{ref} is displayed as {c.number_format!r}, not one decimal ({DISPLAY_FORMAT})"
        )
    d = Decimal(repr(c.value))
    if d != d.quantize(Decimal("0.1")):
        raise CatFormatError(
            f"{ws.title}!{ref} stores {c.value!r}, more digits than the table displays; CAT asks that only the "
            "displayed rounded values be reproduced"
        )
    if d <= 0:
        raise CatFormatError(f"{ws.title}!{ref} is {c.value!r}; the table holds warming above pre-industrial levels")
    return d


def _bold_rows(ws: Worksheet) -> set[int]:
    rows: set[int] = set()
    for r in VALUE_ROWS:
        bold = {bool(ws[f"{col}{r}"].font.b) for col in VALUE_COLUMNS}
        if len(bold) != 1:
            raise CatFormatError(f"{ws.title} row {r}: some value cells are bold and some are not")
        if bold.pop():
            rows.add(r)
    return rows


def parse(raw: bytes) -> tuple[list[Observation], datetime]:
    """The four thermometer scenarios as observations (WLD, 2100), and the file's publication date."""
    wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=False)
    for name in (INFO, THERMOMETER):
        if name not in wb.sheetnames:
            raise CatFormatError(f"no sheet {name!r}; sheets are {wb.sheetnames}")
    info, ws = wb[INFO], wb[THERMOMETER]
    _expect(info, INFO_CELLS)
    _expect(ws, THERMOMETER_CELLS)
    published = info["C10"].value
    if not isinstance(published, datetime):
        raise CatFormatError(f"Info!C10 (Date Published) is {published!r}, not a date")

    want_bold = {s.row for s in SCENARIOS}
    if want_bold | set(NOT_PUBLISHED) != set(VALUE_ROWS):
        raise CatFormatError("the expected layout does not cover every value row")  # a mistake in this module
    bold = _bold_rows(ws)
    if bold != want_bold:
        raise CatFormatError(
            f"bold value rows are {sorted(bold)}, expected {sorted(want_bold)}; the sheet says the bold values are "
            "the ones used in the CAT Thermometer, so re-read the file"
        )
    for r, (b, c) in NOT_PUBLISHED.items():
        if (_text(ws[f"B{r}"].value), _text(ws[f"C{r}"].value)) != (b, c):
            raise CatFormatError(f"{THERMOMETER} row {r} labels changed; re-read the file")

    obs: list[Observation] = []
    for s in SCENARIOS:
        if (_text(ws[f"B{s.row}"].value), _text(ws[f"C{s.row}"].value)) != (s.label, s.variant):
            raise CatFormatError(
                f"{THERMOMETER} row {s.row} reads {_text(ws[f'B{s.row}'].value)!r} / "
                f"{_text(ws[f'C{s.row}'].value)!r}, expected {s.label!r} / {s.variant!r}"
            )
        if _text(info[f"B{s.info_row}"].value) != s.info_label:
            raise CatFormatError(
                f"Info!B{s.info_row} reads {info[f'B{s.info_row}'].value!r}, expected {s.info_label!r}"
            )
        lower, median, upper = (_published_value(ws, f"{col}{s.row}") for col in VALUE_COLUMNS)
        if not lower <= median <= upper:
            raise CatFormatError(f"{THERMOMETER} row {s.row}: median {median} is not within [{lower}, {upper}]")
        explained = _text(info[f"C{s.info_row}"].value)
        note = f"CAT's description of the scenario (Info sheet): {explained}"
        if s.variant == "Combined":
            note += (
                " This is the row CAT labels Combined; the sheet also gives High and Low variants, which the "
                "thermometer does not show."
            )
        obs.append(
            Observation(
                entity="WLD",
                period="2100",
                value=float(median),
                lower=float(lower),
                upper=float(upper),
                interval="range",
                status="projection",
                note=note,
                dims={"scenario": s.id},
            )
        )
    return obs, published


def _run(files: dict[str, InputFile]) -> Result:
    f = files[WORKBOOK.key]
    obs, published = parse(f.path.read_bytes())
    vintage = f"{published:%B %Y}"
    return Result(
        observations=obs,
        vintage=vintage,
        year=str(published.year),
        date_published=published.date().isoformat(),
        steps=[
            "Read the CAT Thermometer sheet of CAT_2025-11_PublicData_GlobalTemperatureEstimates_COP30.xlsx. The Info "
            f"sheet gives the Date Published, {published.day} {published:%B %Y}; the vintage is that month, as CAT "
            "labels the update.",
            "Kept the four rows the sheet prints in bold, which it says are the values used in the CAT Thermometer: "
            "Policies & action (Combined), 2030 & 2035 Targets only, Pledges & targets, and Optimistic scenario "
            "(net-zero pledges). Not published: the High and Low variants of Policies & action, the 2030 targets "
            "only row and an unlabelled row, none of which is in bold.",
            "Each value is the sheet's cell unchanged: Median as the value, Lower bound and Upper bound as the range. "
            "Every cell was checked to be displayed with one decimal and to store no more digits than it displays, "
            "because CAT asks that only the displayed rounded values be reproduced. Nothing was computed.",
        ],
    )


def _check(scenario: str, stated: str, quote: str, url: str) -> PublisherCheck:
    return PublisherCheck(
        source_id=SOURCE,
        vintage=VINTAGE,
        entity="WLD",
        period="2100",
        stated=stated,
        quote=quote,
        url=url,
        dims=(("scenario", scenario),),
    )


PATHWAYS = "https://climateactiontracker.org/global/emissions-pathways/"
BRIEFING = "https://climateactiontracker.org/documents/1348/CAT_2025-11-13_GlobalUpdate_COP30.pdf"

CHECKS = (
    # climateactiontracker.org/global/emissions-pathways/, "Last update: 13 November 2025" (read 2026-10-04).
    _check(
        "policies-action",
        "2.6",
        "Current policies in place around the world are projected to result in about 2.6°C",
        PATHWAYS,
    ),
    _check(
        "pledges-targets",
        "2.2",
        "When binding long-term or net-zero targets are included (our “pledges and targets” scenario), warming would "
        "be limited to about 2.2°C",
        PATHWAYS,
    ),
    _check(
        "optimistic",
        "1.9",
        "Under the optimistic assumption that governments will achieve these targets, the median warming estimate is "
        "1.9°C",
        PATHWAYS,
    ),
    # CAT's November 2025 briefing (Warming Projections Global Update), PDF page 10 (read 2026-10-04).
    _check(
        "targets-2030-2035",
        "2.6",
        "Our “2030 & 2035 targets scenario” which includes the impact of all submitted NDCs to date, remains at "
        "2.6°C, the same as last year.",
        BRIEFING,
    ),
)


def transforms(paths: Paths) -> list[Transform]:
    return [
        Transform(
            spec=Spec(
                id="warming.cat-2025.thermometer",
                title="Projected warming in 2100 under current policies and pledges (Climate Action Tracker)",
                description="Climate Action Tracker's projection of global warming in 2100 above pre-industrial "
                "levels under four sets of assumptions: the policies and action in place now, only the 2030 and 2035 "
                "national targets, all submitted pledges and targets, and an optimistic case in which every announced "
                "target is met. Each is CAT's median estimate with its lower and upper bound, as CAT publishes them "
                "in the CAT Thermometer of November 2025.",
                kind="series",
                unit=DEG_C,
                display=Display(decimals=1),
                scope=Scope(
                    geography="Global mean",
                    baseline="Pre-industrial levels, as CAT defines them",
                    basis="Warming in the year 2100 from the MAGICC7 climate model run on CAT's emissions pathways: "
                    "the median with CAT's lower and upper bounds, rounded by CAT to 0.1 °C and reproduced "
                    "unchanged. CAT's sheet tags the median column '_50' and the upper column '_83' and gives no tag "
                    "for the lower column, so no probability is attached to the range here.",
                ),
                geo_coverage="global-only",
                headline_entity="WLD",
                dimensions=(DIMENSION,),
                headline_dims=(("scenario", "policies-action"),),
            ),
            inputs=(WORKBOOK,),
            run=_run,
            module_file=Path(__file__),
            validation=Validation(min_rows=len(SCENARIOS), value_range=(0.5, 6.0)),
            checks=CHECKS,
        )
    ]
