"""UK Government greenhouse gas conversion factors 2026 (DESNZ): flight emission factors per passenger-kilometre, by
haul and seat class, with and without the radiative-forcing uplift.

Input: the flat file, version 1.2 (artifact flat-file, Open Government Licence v3.0), sheet "Factors by Category".
After a title block, the header row reads ID, Scope, Level 1, Level 2, Level 3, Level 4, Column Text, UOM,
GHG/Unit, "GHG Conversion Factor 2026"; each later row is one factor with a stable factor ID.

Rows used: Level 1 "Business travel- air", Level 2 "Flights", UOM "passenger.km", GHG/Unit "kg CO2e" (the total
of the CO2, CH4 and N2O rows that follow each one, which are not used). Level 3 is the haul (domestic, short-haul and
long-haul flights to or from the UK, and international flights between non-UK airports), Level 4 the seat class,
Column Text "With RF" or "Without RF". Every (haul, class, RF) combination must appear exactly once; each published
value is the factor as printed, with its factor ID in the note.

Not used: the well-to-tank rows (Level 1 "WTT- business travel- air"), which DESNZ publishes separately for the
upstream emissions of producing and delivering jet fuel, and the per-gas rows. Factors already include DESNZ's 8%
uplift of great-circle distance (methodology paper); "With RF" adds DESNZ's radiative-forcing multiplier for
non-CO2 effects at altitude. They are UK reporting averages for 2026, not measurements of any one flight.
"""

from __future__ import annotations

import io
from dataclasses import dataclass
from pathlib import Path

import openpyxl

from envdash.models import Dimension, DimensionValue, Display, Observation, Scope, Unit
from envdash.paths import Paths
from envdash.transform import Input, InputFile, Result, Spec, Transform, Validation

SOURCE = "desnz-ghg-factors-2026"
FLAT = Input(SOURCE, "flat-file")
INDICATOR = "travel.desnz-2026.flight-factors"
SHEET = "Factors by Category"
HEADER = (
    "ID",
    "Scope",
    "Level 1",
    "Level 2",
    "Level 3",
    "Level 4",
    "Column Text",
    "UOM",
    "GHG/Unit",
    "GHG Conversion Factor 2026",
)
LEVEL1, LEVEL2, UOM, GHG = "Business travel- air", "Flights", "passenger.km", "kg CO2e"

# Level 3 as printed -> (dimension id, label)
HAULS: dict[str, tuple[str, str]] = {
    "Domestic, to/from UK": ("domestic", "Domestic, to or from the UK"),
    "Short-haul, to/from UK": ("short-haul", "Short-haul, to or from the UK"),
    "Long-haul, to/from UK": ("long-haul", "Long-haul, to or from the UK"),
    "International, to/from non-UK": ("international-non-uk", "International, between non-UK airports"),
}
CLASSES: dict[str, tuple[str, str]] = {
    "Average passenger": ("average", "Average passenger"),
    "Economy class": ("economy", "Economy class"),
    "Premium economy class": ("premium-economy", "Premium economy class"),
    "Business class": ("business", "Business class"),
    "First class": ("first", "First class"),
}
RF: dict[str, tuple[str, str]] = {
    "With RF": ("with-rf", "With radiative forcing"),
    "Without RF": ("without-rf", "Without radiative forcing"),
}
# The (haul, class) pairs DESNZ publishes, in file order: domestic flights have only an average-passenger factor,
# short-haul no premium economy or first class.
EXPECTED: tuple[tuple[str, str], ...] = (
    ("Domestic, to/from UK", "Average passenger"),
    ("Short-haul, to/from UK", "Average passenger"),
    ("Short-haul, to/from UK", "Economy class"),
    ("Short-haul, to/from UK", "Business class"),
    *(("Long-haul, to/from UK", c) for c in CLASSES),
    *(("International, to/from non-UK", c) for c in CLASSES),
)

UNIT = Unit(code="kgCO2e/pkm", label="kilograms of CO2-equivalent per passenger-kilometre", short="kg CO₂e/pkm")


class DesnzFormatError(ValueError):
    pass


@dataclass(frozen=True)
class Factor:
    factor_id: str
    haul: str
    seat_class: str
    rf: str
    value: float


def read_factors(raw: bytes) -> list[Factor]:
    wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    if SHEET not in wb.sheetnames:
        raise DesnzFormatError(f"no sheet {SHEET!r} (sheets {wb.sheetnames})")
    rows = wb[SHEET].iter_rows(values_only=True)
    for r in rows:
        if r[: len(HEADER)] == HEADER:
            break
    else:
        raise DesnzFormatError(f"no header row {HEADER}")
    found: list[Factor] = []
    for r in rows:
        fid, _scope, l1, l2, l3, l4, text, uom, ghg, value = r[: len(HEADER)]
        if l1 != LEVEL1 or ghg != GHG:
            continue
        where = f"factor {fid}"
        if l2 != LEVEL2 or uom != UOM:
            raise DesnzFormatError(f"{where}: Level 2 {l2!r}, UOM {uom!r}; expected {LEVEL2!r} per {UOM!r}")
        if l3 not in HAULS or l4 not in CLASSES or text not in RF:
            raise DesnzFormatError(f"{where}: unknown haul {l3!r}, class {l4!r} or column text {text!r}")
        if not isinstance(value, (int, float)) or isinstance(value, bool) or not value > 0:
            raise DesnzFormatError(f"{where}: factor {value!r} is not a positive number")
        if not isinstance(fid, str) or not fid:
            raise DesnzFormatError(f"factor without an ID: {r}")
        found.append(Factor(fid, l3, l4, text, float(value)))
    keys = [(f.haul, f.seat_class, f.rf) for f in found]
    expected = [(h, c, rf) for h, c in EXPECTED for rf in RF]
    if sorted(keys) != sorted(expected):
        missing = sorted(set(expected) - set(keys))
        extra = sorted(k for k in keys if keys.count(k) > 1 or k not in expected)
        raise DesnzFormatError(f"flight factors: missing {missing}, unexpected or repeated {extra}")
    return found


def observations(factors: list[Factor]) -> list[Observation]:
    return [
        Observation(
            entity="GBR",
            period="2026",
            value=f.value,
            note=f"DESNZ factor ID {f.factor_id}: {f.haul}, {f.seat_class}, {f.rf}.",
            dims={"haul": HAULS[f.haul][0], "seat_class": CLASSES[f.seat_class][0], "radiative_forcing": RF[f.rf][0]},
        )
        for f in factors
    ]


def _run(files: dict[str, InputFile]) -> Result:
    f = files[FLAT.key]
    factors = read_factors(f.path.read_bytes())
    return Result(
        observations=observations(factors),
        vintage="2026, flat file version 1.2",
        year="2026",
        date_published="2026-06-11",
        steps=[
            f"Read sheet '{SHEET}' of the flat file (sha256 {f.snapshot.sha256[:12]}…) and kept the rows with Level 1 "
            f"'{LEVEL1}', Level 2 '{LEVEL2}', unit '{UOM}' and GHG/Unit '{GHG}' (the all-gas total): "
            f"{len(factors)} factors, one for each haul, seat class and with or without radiative forcing that DESNZ "
            "publishes. Checked that each appears exactly once.",
            "Published each factor as printed, with its DESNZ factor ID in the value's note. Well-to-tank rows and "
            "the separate CO2, CH4 and N2O rows are not used.",
        ],
    )


def transforms(paths: Paths) -> list[Transform]:
    return [
        Transform(
            spec=Spec(
                id=INDICATOR,
                title="Flight emission factors by distance and seat class (UK Government 2026)",
                description="Greenhouse gas emissions per passenger per kilometre flown, from the UK Government's "
                "2026 conversion factors: for domestic, short-haul and long-haul flights to or from the UK and for "
                "international flights elsewhere, by seat class, with and without the extra warming from aircraft "
                "emissions at altitude (radiative forcing). Premium seats take more cabin space, so each passenger "
                "in them is assigned more of the flight's emissions. Factors are averages for UK reporting, not "
                "measurements of a particular flight.",
                kind="series",
                unit=UNIT,
                display=Display(decimals=3),
                scope=Scope(
                    geography="UK Government factors for reporting by UK organisations (flights to or from the UK, "
                    "and international flights between non-UK airports)",
                    basis="Direct emissions of the flight (CO2, CH4 and N2O as CO2-equivalent, with the warming "
                    "potentials DESNZ lists in its methodology paper, Table 1), excluding the well-to-tank emissions "
                    "of producing jet fuel. The factors include DESNZ's 8% uplift on the great-circle distance "
                    "(methodology paper, paragraph 8.17). 'With radiative forcing' applies DESNZ's 1.7 multiplier to "
                    "the CO2 for the effects of emissions at altitude (paragraph 2.10); 'without' gives the "
                    "emissions alone.",
                ),
                geo_coverage="country",
                headline_entity="GBR",
                dimensions=(
                    Dimension(
                        id="haul",
                        label="Flight type",
                        values=[DimensionValue(id=i, label=label) for i, label in HAULS.values()],
                    ),
                    Dimension(
                        id="seat_class",
                        label="Seat class",
                        values=[DimensionValue(id=i, label=label) for i, label in CLASSES.values()],
                    ),
                    Dimension(
                        id="radiative_forcing",
                        label="Radiative forcing",
                        values=[DimensionValue(id=i, label=label) for i, label in RF.values()],
                    ),
                ),
                headline_dims=(("haul", "long-haul"), ("seat_class", "economy"), ("radiative_forcing", "with-rf")),
            ),
            inputs=(FLAT,),
            run=_run,
            module_file=Path(__file__),
            validation=Validation(min_rows=2 * len(EXPECTED), value_range=(0.0, 2.0)),
        )
    ]
