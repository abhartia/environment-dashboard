"""Drawdown Explorer: the climate impact of each solution at current, achievable and ceiling adoption, in billion
tonnes of CO2-equivalent per year, for the whole solution worldwide (system level).

Inputs: Project Drawdown's solution assessment spreadsheets on Zenodo (one artifact per solution, CC BY 4.0). Each
holds a results sheet ("results", or "Global results" for the protect-ecosystem solutions) with a table titled like
"Table 8: Climate impact at different levels of adoption. Unit: Gt CO2-eq (100-year basis) per year": a header row
with four adoption levels and one row of values (two rows, labelled "20-year basis" and "100-year basis", in some
transport spreadsheets). The methodology (doi:10.5281/zenodo.20865531, p. 24) defines the levels: the impact at
Current Adoption, the range between Achievable – Low and Achievable – High ("the potential climate benefit of
adopting the solution within the specified Adoption Achievable Range"), and the Adoption Ceiling ("the theoretical
maximum climate benefit").

The spreadsheets differ in layout, so each table read is declared in TABLES: the sheet, the title cell and its exact
text, the header row's four level headers (exact text), and each value row with the warming-potential basis it
states. Every declared cell is checked before a value is read; any difference stops the transform. Three header
texts name the ceiling column: "Adoption Ceiling", "Adoption ceiling" and "Potential Adoption", the last in the
fourth column of the same Table 8 layout. A value cell must be a number, or a text listed in the table's
`not_given`, which is published as a missing value with the text as the reason.

Solutions not published, and why (spreadsheets of 2026-10-04):
- Improve Cement Production: the results sheet has no table of climate impact by adoption level.
- Deploy Silvopasture: the results sheet gives sequestration by adoption level only in an untitled working table
  (rows "total adoption ha" to "total carbon sequestration GtCO2-eq/yr"), with no stated unit basis.
- Improve Annual Cropping: the achievable levels are given year by year (2025 onwards), not as one value each.
- Improve Rice Production: its table is "Global impact, Gt CO2-eq/yr in 2030", impacts in one year rather than at
  the adoption levels the other tables use.
- Improve Landfill Management, Increase Recycling, Mobilize Electric Bicycles, Protect Coastal Wetlands, Restore
  Forests, Use Smart and Programmable Thermostats: impacts are given only per sub-solution (landfill gas capture and
  biocovers; glass, metal, paper and plastic; private and shared bicycles; mangroves, salt marshes and seagrasses;
  forest types; heating and cooling), with no total for the solution. Sub-solutions are not added up here.
Where a spreadsheet gives only a 20-year basis (Distributed Solar PV, Offshore Wind, Nonmotorized Transportation,
Protect Forests), only that is published; values on the two bases are never mixed in one series.
"""

from __future__ import annotations

import io
from dataclasses import dataclass
from pathlib import Path

import openpyxl

from envdash.models import Dimension, DimensionValue, Display, Observation, Scope, Unit
from envdash.paths import Paths
from envdash.transform import Input, InputFile, Result, Spec, Transform, Validation

SOURCE = "drawdown-explorer"
INDICATOR = "solutions.drawdown.climate-impact"
PERIOD = "2026"

LEVELS: tuple[tuple[str, str], ...] = (
    ("current", "At current adoption"),
    ("achievable-low", "Achievable range, low end"),
    ("achievable-high", "Achievable range, high end"),
    ("ceiling", "At the adoption ceiling (theoretical maximum)"),
)
HEADERS: dict[str, str] = {
    "Current Adoption": "current",
    "Achieveable - Low": "achievable-low",
    "Achievable - Low": "achievable-low",
    "Achievable – Low": "achievable-low",
    "Achievable - High": "achievable-high",
    "Achievable – High": "achievable-high",
    "Adoption Ceiling": "ceiling",
    "Adoption ceiling": "ceiling",
    "Potential Adoption": "ceiling",
}
BASES: tuple[tuple[str, str], ...] = (
    ("gwp100", "100-year warming potential"),
    ("gwp20", "20-year warming potential"),
)

_T8 = "Table 8: Climate impact at different levels of adoption. Unit: Gt CO2-eq ({} basis) per year"


@dataclass(frozen=True)
class Table:
    artifact: str
    solution: str
    """The solution's name as the Drawdown Explorer gives it."""
    sheet: str
    title_cell: tuple[int, int]
    """(row, column), 1-based row, 0-based column as openpyxl's values_only rows index them."""
    title: str
    header_row: int
    first_col: int
    rows: tuple[tuple[int, str, str | None], ...]
    """(row, basis id, label printed in the column left of the values, or None)."""
    not_given: tuple[str, ...] = ()


def _std(artifact, solution, row, basis, sheet="results", title=None, col=5, not_given=()):
    """The common layout: title at (row, col), headers on the next row from col, values on the row after."""
    return Table(
        artifact,
        solution,
        sheet,
        (row, col),
        title or _T8.format("100-year" if basis == "gwp100" else "20-year"),
        row + 1,
        col,
        ((row + 2, basis, None),),
        not_given,
    )


def _two_bases(artifact, solution, title):
    """Transport layout: title at row 63, headers in columns 6-9 of row 64, rows 65 (20-year) and 66 (100-year)."""
    return Table(
        artifact,
        solution,
        "results",
        (63, 5),
        title,
        64,
        6,
        ((65, "gwp20", "20-year basis"), (66, "gwp100", "100-year basis")),
    )


TABLES: tuple[Table, ...] = (
    _std("deploy-alternative-insulation-materials", "Deploy Alternative Insulation Materials", 71, "gwp100"),
    _std("deploy-alternative-refrigerants", "Deploy Alternative Refrigerants", 71, "gwp100"),
    _std("deploy-building-automation-systems", "Deploy Building Automation Systems", 70, "gwp100"),
    _std("deploy-clean-cooking", "Deploy Clean Cooking", 70, "gwp100"),
    _std("deploy-distributed-solar-pv", "Deploy Distributed Solar PV", 70, "gwp20"),
    _std(
        "deploy-industrial-green-hydrogen",
        "Deploy Industrial Green Hydrogen",
        70,
        "gwp100",
        title="Table 8: Climate impact of green hydrogen at different levels of adoption. Unit: Gt CO2-eq (100-year "
        "basis) per year",
    ),
    _std("deploy-led-lighting", "Deploy LED Lighting", 70, "gwp100"),
    _std("deploy-offshore-wind-turbines", "Deploy Offshore Wind Turbines", 70, "gwp20"),
    _std("deploy-onshore-wind-turbines", "Deploy Onshore Wind Turbines", 53, "gwp100"),
    _std("deploy-utility-scale-solar-pv", "Deploy Utility-Scale Solar PV", 55, "gwp100"),
    _std("enhance-public-transit", "Enhance Public Transit", 48, "gwp100"),
    _std("improve-diets", "Improve Diets", 70, "gwp100", not_given=("Climate impact (100 yr GtCO2-eq/yr)",)),
    _std("improve-nonmotorized-transportation", "Improve Nonmotorized Transportation", 75, "gwp20"),
    _std(
        "improve-nutrient-management",
        "Improve Nutrient Management",
        58,
        "gwp100",
        title="Table 6: Climate impact at different levels of adoption. Unit: Gt CO2-eq (100-year basis)/yr",
    ),
    _std("improve-windows-and-glass", "Improve Windows and Glass", 70, "gwp100"),
    _two_bases("increase-carpooling", "Increase Carpooling", _T8.format("20-year")),
    _std("increase-centralized-composting", "Increase Centralized Composting", 70, "gwp100"),
    _std(
        "manage-coal-mine-methane",
        "Manage Coal Mine Methane",
        70,
        "gwp100",
        title="Table 7: Climate impact at different levels of adoption. Unit: Gt CO2-eq/yr (100-year basis)",
    ),
    _std(
        "manage-coal-mine-methane",
        "Manage Coal Mine Methane",
        70,
        "gwp20",
        col=12,
        title="Table 7b: Climate impact at different levels of adoption. Unit: Gt CO2-eq/yr (20-year basis)",
    ),
    _std(
        "manage-oil-and-gas-methane",
        "Manage Oil and Gas Methane",
        71,
        "gwp100",
        title="Table 7a: Climate impact at different levels of adoption. Unit: Gt CO2-eq/yr (100-year basis)",
    ),
    _std(
        "manage-oil-and-gas-methane",
        "Manage Oil and Gas Methane",
        71,
        "gwp20",
        col=10,
        title="Table 7b: Climate impact at different levels of adoption. Unit: Gt CO2-eq/yr (20-year basis)",
    ),
    _two_bases("mobilize-electric-cars", "Mobilize Electric Cars", _T8.format("20-year")),
    _two_bases(
        "mobilize-electric-scooters-and-motorcycles",
        "Mobilize Electric Scooters and Motorcycles",
        "Table 8: Climate impact at different levels of adoption. Unit: Gt CO2-eq per year",
    ),
    _two_bases("mobilize-hybrid-cars", "Mobilize Hybrid Cars", _T8.format("20-year")),
    _std("protect-forests", "Protect Forests", 63, "gwp20", sheet="Global results"),
    _std("protect-grasslands-and-savannas", "Protect Grasslands and Savannas", 52, "gwp100", sheet="Global results"),
    _std("protect-peatlands", "Protect Peatlands", 70, "gwp100", sheet="Global results"),
    _std(
        "protect-seaweed-ecosystems",
        "Protect Seaweed Ecosystems",
        46,
        "gwp100",
        sheet="Global results",
        title="Table 7: Climate impact at different levels of adoption. Unit: Gt CO2-eq (100-year basis) per year",
    ),
    _std("reduce-food-loss-and-waste", "Reduce Food Loss and Waste", 70, "gwp100", not_given=("Not determined",)),
    _std("use-heat-pumps", "Use Heat Pumps", 70, "gwp100"),
)
ARTIFACTS: tuple[str, ...] = tuple(dict.fromkeys(t.artifact for t in TABLES))
INPUTS = tuple(Input(SOURCE, a) for a in ARTIFACTS)

UNIT = Unit(code="GtCO2e/yr", label="billion tonnes of CO2-equivalent per year", short="Gt CO₂e/yr")


class DrawdownFormatError(ValueError):
    pass


def read_table(rows: list[tuple], t: Table) -> list[Observation]:
    def cell(r: int, c: int) -> object:
        row = rows[r - 1] if r - 1 < len(rows) else ()
        return row[c] if c < len(row) else None

    where = f"{t.artifact} '{t.sheet}'"
    if cell(*t.title_cell) != t.title:
        raise DrawdownFormatError(f"{where} row {t.title_cell[0]}: title {cell(*t.title_cell)!r} != {t.title!r}")
    levels = []
    for i in range(4):
        h = cell(t.header_row, t.first_col + i)
        if h not in HEADERS:
            raise DrawdownFormatError(f"{where} row {t.header_row}: header {h!r} is not an adoption level")
        levels.append(HEADERS[h])
    if levels != [lid for lid, _ in LEVELS]:
        raise DrawdownFormatError(f"{where} row {t.header_row}: levels {levels} out of order")
    obs = []
    for r, basis, label in t.rows:
        if label is not None and cell(r, t.first_col - 1) != label:
            raise DrawdownFormatError(f"{where} row {r}: label {cell(r, t.first_col - 1)!r} != {label!r}")
        for i, (lid, _) in enumerate(LEVELS):
            v = cell(r, t.first_col + i)
            dims = {"solution": t.artifact, "level": lid, "basis": basis}
            status = "final" if lid == "current" else "projection"
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                obs.append(Observation(entity="WLD", period=PERIOD, value=float(v), status=status, dims=dims))
            elif v in t.not_given:
                obs.append(
                    Observation(
                        entity="WLD",
                        period=PERIOD,
                        value=None,
                        status=status,
                        missing_reason=f"The spreadsheet's cell for this level reads '{v}', not a number.",
                        dims=dims,
                    )
                )
            else:
                raise DrawdownFormatError(f"{where} row {r} column {t.first_col + i}: {v!r} is not a number")
    return obs


def sheet_rows(raw: bytes, sheet: str) -> list[tuple]:
    wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    try:
        if sheet not in wb.sheetnames:
            raise DrawdownFormatError(f"no sheet {sheet!r} (sheets {wb.sheetnames})")
        return list(wb[sheet].iter_rows(values_only=True))
    finally:
        wb.close()


def _run(files: dict[str, InputFile]) -> Result:
    obs: list[Observation] = []
    cache: dict[tuple[str, str], list[tuple]] = {}
    for t in TABLES:
        f = files[f"{SOURCE}/{t.artifact}"]
        key = (t.artifact, t.sheet)
        if key not in cache:
            cache[key] = sheet_rows(f.path.read_bytes(), t.sheet)
        obs.extend(read_table(cache[key], t))
    solutions = len({o.dims["solution"] for o in obs})
    missing = sum(o.value is None for o in obs)
    return Result(
        observations=obs,
        vintage="Drawdown Explorer spreadsheets on Zenodo, June–August 2026 (methodology June 2026)",
        year="2026",
        steps=[
            f"Read the table of climate impact by adoption level from the results sheet of {len(ARTIFACTS)} solution "
            "spreadsheets, checking the title, the four level headers and any row labels of each table cell by cell "
            "against the layout declared for that spreadsheet.",
            f"Published {len(obs)} values for {solutions} solutions as the spreadsheets give them (cached cell values, "
            f"billion tonnes of CO2-equivalent per year), on the warming-potential basis each table states. {missing} "
            "cells that hold text instead of a number are published as missing values, with the text.",
            "Solutions whose spreadsheets give impacts only per sub-solution, by year, or without a stated basis are "
            "not published (see the transform's notes).",
        ],
    )


def transforms(paths: Paths) -> list[Transform]:
    solutions = dict((t.artifact, t.solution) for t in TABLES)
    return [
        Transform(
            spec=Spec(
                id=INDICATOR,
                title="Climate impact of solutions at different levels of adoption (Project Drawdown)",
                description="How much each climate solution assessed by Project Drawdown reduces or removes "
                "greenhouse gas emissions each year worldwide: at today's adoption, across the range of adoption "
                "Drawdown considers achievable in the next 5 to 25 years, and at the adoption ceiling, a theoretical "
                "maximum that probably cannot be reached. Impacts of different solutions overlap and cannot be added "
                "up. For renewable electricity the ceiling reflects technical potential only, not power demand.",
                kind="series",
                unit=UNIT,
                display=Display(decimals=2),
                scope=Scope(
                    geography="World, each solution as a whole (system level)",
                    basis="Project Drawdown's 2026 assessments against a 2023 baseline; impacts assume each "
                    "solution's effectiveness stays as it is today. Gases are combined with 100-year or 20-year "
                    "warming potentials as each spreadsheet states (dimension 'basis'); the two are not comparable.",
                ),
                geo_coverage="global-only",
                headline_entity="WLD",
                dimensions=(
                    Dimension(
                        id="solution",
                        label="Solution",
                        values=[DimensionValue(id=a, label=s) for a, s in solutions.items()],
                    ),
                    Dimension(
                        id="level", label="Adoption level", values=[DimensionValue(id=i, label=n) for i, n in LEVELS]
                    ),
                    Dimension(
                        id="basis", label="Warming potential", values=[DimensionValue(id=i, label=n) for i, n in BASES]
                    ),
                ),
                headline_dims=(("solution", "deploy-onshore-wind-turbines"), ("level", "current"), ("basis", "gwp100")),
            ),
            inputs=INPUTS,
            run=_run,
            module_file=Path(__file__),
            validation=Validation(min_rows=4 * len(TABLES), value_range=(0.0, 500.0)),
        )
    ]
