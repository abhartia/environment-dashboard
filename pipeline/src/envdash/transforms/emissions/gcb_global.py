"""Global Carbon Budget 2025: the global carbon dioxide budget, fossil emissions net of cement carbonation, and total
emissions, in billion tonnes of carbon dioxide per year.

Input: sheet "Global Carbon Budget" of Global_Carbon_Budget_2025_v1.0.xlsx (ICOS PID 11676/qSjPBsV1drZnYdH-yCJMmkGn).
Under a block of notes, a "GtC/yr" line and a "Year" header row, it has one row per year 1959-2024 with seven
columns: fossil emissions excluding carbonation, land-use change emissions, atmospheric growth, ocean sink, land sink,
cement carbonation sink, budget imbalance. Sinks are positive numbers.

"Fossil emissions excluding carbonation" is fossil fuels and industry before the cement carbonation sink is
subtracted (38.6 GtCO2 in 2024). The sheet's own definition of the budget imbalance ("the sum of emissions (fossil
fuel and industry + land-use change) minus (atmospheric growth + ocean sink + land sink + cement carbonation sink)")
treats the carbonation sink as a sink, and the transform checks that identity on every row, which pins down that
reading. The Global Carbon Budget's headline fossil value (37.8 GtCO2 in 2024, ESSD executive summary) is net of the
sink: fossil minus cement carbonation sink. Total emissions are that net value plus land-use change (42.4 GtCO2).

Units. The sheet says "All values in billion tonnes of carbon per year (GtC/yr) ... multiply the numbers below by
3.664". Values are multiplied by 3.664 with exact decimal arithmetic on the stored numbers (each read as the shortest
decimal that round-trips the stored double).

Uncertainty. The sheet states one-sigma uncertainties of ±5 % for fossil emissions and ±0.7 GtC/yr for land-use change;
those two components carry lower/upper (interval 1sigma). For the ocean sink (±0.4 GtC/yr "on average"), the land
sink (±0.5 GtC/yr "on average") and atmospheric growth ("variable uncertainty around 0.2 GtC/yr from 1980") the sheet
gives no value for a given year, so no range is published. The derived net and total series carry no range: the
sheet gives none for them, and combining the components' uncertainties would be our estimate, not the producer's.

Version. The file carries no version label; the ICOS object is pinned in the registry and is version 1.0. PINNED maps
the object URL to its version and ICOS submission date; another URL stops the build until a person adds it. If any
statement quoted in REQUIRED disappears from the sheet, the transform stops: a person must re-read the file.
"""

from __future__ import annotations

import io
from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from pathlib import Path

import openpyxl

from envdash.models import Dimension, DimensionValue, Display, Observation, Scope, Unit
from envdash.paths import Paths
from envdash.transform import Input, InputFile, PublisherCheck, Result, Spec, Transform, Validation

SOURCE = "gcb-2025-global"
BUDGET = Input(SOURCE, "global-budget")
SHEET = "Global Carbon Budget"
GTC_TO_GTCO2 = Decimal("3.664")
FOSSIL_REL_1SIGMA = Decimal("0.05")
LUC_ABS_1SIGMA_GTC = Decimal("0.7")
IMBALANCE_TOLERANCE = Decimal("1e-9")

PINNED: dict[str, tuple[str, date]] = {
    # meta.icos-cp.eu/objects/qSjPBsV1drZnYdH-yCJMmkGn (read 2026-10-04): Global_Carbon_Budget_2025_v1.0.xlsx,
    # submitted 2026-05-17, latestVersion = itself.
    "https://data.icos-cp.eu/objects/qSjPBsV1drZnYdH-yCJMmkGn": ("2025 v1.0", date(2026, 5, 17)),
}

COLUMNS = (
    "Year",
    "fossil emissions excluding carbonation",
    "land-use change emissions",
    "atmospheric growth",
    "ocean sink",
    "land sink",
    "cement carbonation sink",
    "budget imbalance",
)

REQUIRED = (
    "All values in billion tonnes of carbon per year (GtC/yr), for the globe. For values in billion tonnes of carbon "
    "dioxide per year (GtCO2/yr) , multiply the numbers below by 3.664.",
    "All uncertainties represent ± 1 sigma error (68 % chance of being in the range provided)",
    "Emissions from fossil fuel combustion and industrial processes (uncertainty of ±5% for a ± 1 sigma confidence "
    "level):",
    "Emissions from land-use change (uncertainty of ±0.7 GtC/yr).",
    "The budget imbalance is the sum of emissions (fossil fuel and industry + land-use change) minus (atmospheric "
    "growth + ocean sink + land sink + cement carbonation sink)",
)

COMPONENTS: tuple[tuple[str, str, str], ...] = (
    # (dimension value id, label, sheet column)
    ("fossil", "Fossil fuels and industry (before the cement carbonation sink)", COLUMNS[1]),
    ("land-use-change", "Land-use change", COLUMNS[2]),
    ("atmospheric-growth", "Growth in the atmosphere", COLUMNS[3]),
    ("ocean-sink", "Ocean sink", COLUMNS[4]),
    ("land-sink", "Land sink", COLUMNS[5]),
    ("cement-carbonation-sink", "Cement carbonation sink", COLUMNS[6]),
    ("budget-imbalance", "Budget imbalance", COLUMNS[7]),
)

GTCO2_YR = Unit(code="GtCO2/yr", label="billion tonnes of carbon dioxide per year", short="Gt CO₂/yr")
ESSD_URL = "https://essd.copernicus.org/articles/18/3211/2026/"


class GcbFormatError(ValueError):
    pass


@dataclass(frozen=True)
class Row:
    year: int
    values: dict[str, Decimal]
    """Sheet column name -> GtC/yr, as stored."""


def _decimal(v: object, where: str) -> Decimal:
    # bool is an int subclass; a formula would come back as a string ("=...") and is refused too.
    if isinstance(v, bool) or not isinstance(v, int | float):
        raise GcbFormatError(f"{where}: expected a number, found {v!r}")
    return Decimal(repr(v))


def read_budget(raw: bytes) -> list[Row]:
    """The yearly rows of the "Global Carbon Budget" sheet, in GtC/yr, after checking its notes and layout."""
    wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
    try:
        if SHEET not in wb.sheetnames:
            raise GcbFormatError(f"no sheet {SHEET!r} (sheets: {wb.sheetnames})")
        rows = [tuple(r) for r in wb[SHEET].iter_rows(values_only=True)]
    finally:
        wb.close()
    header_at = next((i for i, r in enumerate(rows) if r and r[0] == "Year"), None)
    if header_at is None:
        raise GcbFormatError("no 'Year' header row")
    notes = " ".join(str(c) for r in rows[:header_at] for c in r if c is not None)
    for s in REQUIRED:
        if s not in notes:
            raise GcbFormatError(f"the sheet no longer says {s!r}; re-read the file before trusting these rules")
    if rows[header_at - 1][0] != "GtC/yr":
        raise GcbFormatError(f"the line above the header is {rows[header_at - 1][0]!r}, not 'GtC/yr'")
    header = rows[header_at]
    if tuple(header[: len(COLUMNS)]) != COLUMNS or any(c is not None for c in header[len(COLUMNS) :]):
        raise GcbFormatError(f"header {header[: len(COLUMNS) + 2]} != expected {COLUMNS}")
    out: list[Row] = []
    for n, r in enumerate(rows[header_at + 1 :], start=header_at + 2):
        if all(c is None for c in r):
            continue
        if any(c is not None for c in r[len(COLUMNS) :]):
            raise GcbFormatError(f"row {n} has values beyond column {len(COLUMNS)}")
        year = r[0]
        if isinstance(year, bool) or not isinstance(year, int):
            raise GcbFormatError(f"row {n}: Year {year!r} is not an integer")
        if out and year != out[-1].year + 1:
            raise GcbFormatError(f"row {n}: year {year} does not follow {out[-1].year}")
        values = {col: _decimal(r[i], f"row {n} {col}") for i, col in enumerate(COLUMNS[1:], start=1)}
        out.append(Row(year, values))
    if not out:
        raise GcbFormatError("no data rows under the header")
    return out


def largest_imbalance_residual(rows: list[Row]) -> Decimal:
    """Checks the sheet's own definition of the budget imbalance on every row; returns the largest residual."""
    worst = Decimal(0)
    for row in rows:
        v = row.values
        sources = v[COLUMNS[1]] + v[COLUMNS[2]]
        sinks = v[COLUMNS[3]] + v[COLUMNS[4]] + v[COLUMNS[5]] + v[COLUMNS[6]]
        residual = abs(sources - sinks - v[COLUMNS[7]])
        if residual > IMBALANCE_TOLERANCE:
            raise GcbFormatError(
                f"{row.year}: budget imbalance {v[COLUMNS[7]]} != emissions minus sinks {sources - sinks}; the columns "
                "no longer mean what this transform assumes"
            )
        worst = max(worst, residual)
    return worst


def to_gtco2(gtc: Decimal) -> float:
    return float(gtc * GTC_TO_GTCO2)


def budget_observations(rows: list[Row]) -> list[Observation]:
    obs: list[Observation] = []
    for comp_id, _, col in COMPONENTS:
        for row in rows:
            gtc = row.values[col]
            if comp_id == "fossil":
                half = gtc * FOSSIL_REL_1SIGMA
            elif comp_id == "land-use-change":
                half = LUC_ABS_1SIGMA_GTC
            else:
                half = None
            obs.append(
                Observation(
                    entity="WLD",
                    period=f"{row.year:04d}",
                    value=to_gtco2(gtc),
                    lower=None if half is None else to_gtco2(gtc - half),
                    upper=None if half is None else to_gtco2(gtc + half),
                    interval=None if half is None else "1sigma",
                    dims={"component": comp_id},
                )
            )
    return obs


def fossil_net_gtc(row: Row) -> Decimal:
    return row.values[COLUMNS[1]] - row.values[COLUMNS[6]]


def fossil_net_observations(rows: list[Row]) -> list[Observation]:
    return [Observation(entity="WLD", period=f"{r.year:04d}", value=to_gtco2(fossil_net_gtc(r))) for r in rows]


def total_observations(rows: list[Row]) -> list[Observation]:
    return [
        Observation(entity="WLD", period=f"{r.year:04d}", value=to_gtco2(fossil_net_gtc(r) + r.values[COLUMNS[2]]))
        for r in rows
    ]


def _pinned(f: InputFile) -> tuple[str, date]:
    url = str(f.snapshot.url) if f.snapshot.url else None
    if url not in PINNED:
        raise GcbFormatError(f"{url!r} is not a pinned Global Carbon Budget object; add its version to PINNED")
    return PINNED[url]


def _read_step(rows: list[Row], version: str, residual: Decimal) -> list[str]:
    return [
        f"Read the sheet '{SHEET}' of Global_Carbon_Budget_2025_v1.0.xlsx (Global Carbon Budget {version}, ICOS "
        f"object qSjPBsV1drZnYdH-yCJMmkGn): one row per year {rows[0].year}–{rows[-1].year}, in billion tonnes of "
        "carbon per year. Sinks are positive numbers in the sheet and stay positive here.",
        "Checked the sheet's definition of the budget imbalance on every year: fossil emissions plus land-use change "
        "emissions, minus atmospheric growth, ocean sink, land sink and cement carbonation sink, equals the imbalance "
        f"column (largest difference {float(residual):.1e} GtC/yr). This confirms that the fossil column is before "
        "the cement carbonation sink is subtracted.",
    ]


CONVERT_STEP = (
    "Converted from billion tonnes of carbon to billion tonnes of carbon dioxide by multiplying by 3.664, the factor "
    'the sheet states ("multiply the numbers below by 3.664"), with exact decimal arithmetic on the stored values.'
)


def _run_budget(files: dict[str, InputFile]) -> Result:
    f = files[BUDGET.key]
    version, published = _pinned(f)
    rows = read_budget(f.path.read_bytes())
    residual = largest_imbalance_residual(rows)
    return Result(
        observations=budget_observations(rows),
        vintage=version,
        date_published=published.isoformat(),
        steps=[
            *_read_step(rows, version, residual),
            CONVERT_STEP,
            "Lower and upper are given for two components only, from the one-sigma uncertainties the sheet states: "
            "±5 % of fossil emissions, and ±0.7 billion tonnes of carbon (2.5648 billion tonnes of carbon dioxide) "
            "per year for land-use change. The sheet gives the ocean sink (±0.4) and land sink (±0.5 billion tonnes "
            "of carbon per year) uncertainties only 'on average' and atmospheric growth's as 'variable', so those "
            "components and the imbalance have no range.",
        ],
        changes="converted from billion tonnes of carbon to billion tonnes of carbon dioxide (× 3.664); one-sigma "
        "ranges added for fossil and land-use change emissions from the uncertainties stated in the sheet.",
    )


def _run_fossil_net(files: dict[str, InputFile]) -> Result:
    f = files[BUDGET.key]
    version, published = _pinned(f)
    rows = read_budget(f.path.read_bytes())
    residual = largest_imbalance_residual(rows)
    return Result(
        observations=fossil_net_observations(rows),
        vintage=version,
        date_published=published.isoformat(),
        steps=[
            *_read_step(rows, version, residual),
            "Subtracted the cement carbonation sink from fossil emissions (excluding carbonation) for each year, "
            "which gives the Global Carbon Budget's headline fossil value.",
            CONVERT_STEP,
        ],
        changes="cement carbonation sink subtracted from fossil emissions; converted from billion tonnes of carbon "
        "to billion tonnes of carbon dioxide (× 3.664).",
    )


def _run_total(files: dict[str, InputFile]) -> Result:
    f = files[BUDGET.key]
    version, published = _pinned(f)
    rows = read_budget(f.path.read_bytes())
    residual = largest_imbalance_residual(rows)
    return Result(
        observations=total_observations(rows),
        vintage=version,
        date_published=published.isoformat(),
        steps=[
            *_read_step(rows, version, residual),
            "For each year: fossil emissions (excluding carbonation) minus the cement carbonation sink, plus land-use "
            "change emissions.",
            CONVERT_STEP,
        ],
        changes="cement carbonation sink subtracted from fossil emissions and land-use change emissions added; "
        "converted from billion tonnes of carbon to billion tonnes of carbon dioxide (× 3.664).",
    )


def transforms(paths: Paths) -> list[Transform]:
    here = Path(__file__)
    dims = (
        Dimension(
            id="component",
            label="Budget component",
            values=[DimensionValue(id=i, label=label) for i, label, _ in COMPONENTS],
        ),
    )
    return [
        Transform(
            spec=Spec(
                id="emissions.gcb-2025.budget-global",
                title="Global carbon dioxide budget: emissions and where they go",
                description="Each year since 1959, the carbon dioxide released by fossil fuels and industry and by "
                "land-use change, and where it went: growth in the atmosphere, uptake by the ocean, by land "
                "ecosystems and by cement as it ages (the cement carbonation sink). The budget imbalance is what "
                "the estimates of emissions and sinks do not account for.",
                kind="series",
                unit=GTCO2_YR,
                display=Display(decimals=2),
                scope=Scope(
                    geography="World",
                    lulucf="included",
                    bunkers="included",
                    basis="Carbon dioxide only. Fossil emissions are before the cement carbonation sink is "
                    "subtracted; sinks are positive numbers. Land-use change is the average of three bookkeeping "
                    "models (BLUE, OSCAR, LUCE); the ocean and land sinks are model and data-product averages.",
                ),
                geo_coverage="global-only",
                headline_entity="WLD",
                dimensions=dims,
                headline_dims=(("component", "fossil"),),
            ),
            inputs=(BUDGET,),
            run=_run_budget,
            module_file=here,
            validation=Validation(min_rows=7 * 66, value_range=(-15.0, 50.0)),
        ),
        Transform(
            spec=Spec(
                id="emissions.gcb-2025.fossil-net-global",
                title="Fossil carbon dioxide emissions, world (Global Carbon Budget)",
                description="Carbon dioxide released each year since 1959 by burning fossil fuels, making cement "
                "and other industrial processes, minus the carbon dioxide that cement takes back up as it ages. "
                "This is the Global Carbon Budget's headline fossil emissions value.",
                kind="series",
                unit=GTCO2_YR,
                display=Display(decimals=1),
                scope=Scope(
                    geography="World",
                    lulucf="excluded",
                    bunkers="included",
                    basis="Carbon dioxide from fossil fuels and industry, net of the cement carbonation sink.",
                ),
                geo_coverage="global-only",
                headline_entity="WLD",
            ),
            inputs=(BUDGET,),
            run=_run_fossil_net,
            module_file=here,
            validation=Validation(min_rows=66, value_range=(5.0, 45.0)),
            checks=(
                PublisherCheck(
                    source_id=SOURCE,
                    vintage="2025 v1.0",
                    entity="WLD",
                    period="2024",
                    stated="37.8",
                    quote="Consolidated data confirm a growth of 1.1 % in 2024 relative to 2023, with fossil CO2 "
                    "emissions of 10.3 ± 0.5 GtC yr−1 (37.8 ± 1.8 GtCO2 yr−1) in 2024.",
                    url=ESSD_URL,
                ),
            ),
        ),
        Transform(
            spec=Spec(
                id="emissions.gcb-2025.total-co2-global",
                title="Total carbon dioxide emissions from human activity, world (Global Carbon Budget)",
                description="Carbon dioxide released each year since 1959 by fossil fuels and industry (net of the "
                "cement carbonation sink) and by land-use change such as deforestation.",
                kind="series",
                unit=GTCO2_YR,
                display=Display(decimals=1),
                scope=Scope(
                    geography="World",
                    lulucf="included",
                    bunkers="included",
                    basis="Carbon dioxide only: fossil fuels and industry net of the cement carbonation sink, plus "
                    "land-use change (average of three bookkeeping models).",
                ),
                geo_coverage="global-only",
                headline_entity="WLD",
            ),
            inputs=(BUDGET,),
            run=_run_total,
            module_file=here,
            validation=Validation(min_rows=66, value_range=(10.0, 55.0)),
            checks=(
                PublisherCheck(
                    source_id=SOURCE,
                    vintage="2025 v1.0",
                    entity="WLD",
                    period="2024",
                    stated="42.4",
                    quote="Total anthropogenic emissions (fossil and land use, including the carbonation sink) were "
                    "11.6 GtC yr−1 (42.4 GtCO2 yr−1) in 2024",
                    url=ESSD_URL,
                ),
            ),
        ),
    ]
